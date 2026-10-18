"""The report of an analysis (v0.9, step 6, catalogue § 8): Markdown, its tables in Excel,
and the equivalent R script with the command file and the preparation script it runs. The
same files for the Analyses page and `ddict analyser`.

Only the last analysis run is kept, in memory (D-078): the project stays a codebook and
holds no result. The team keeps the reports it writes or downloads.
"""

from __future__ import annotations

import io
import platform
import re
import zipfile
from dataclasses import dataclass
from importlib.metadata import version
from pathlib import Path

from openpyxl import Workbook

from ddict import __version__
from ddict.analysis.apa import ApaDocument
from ddict.analysis.book import Codebook
from ddict.analysis.catalogue import CATALOGUE, CHOICES, OPTIONS, ROLES, TITLES
from ddict.analysis.data import PreparedData, level_codes
from ddict.analysis.results import AnalysisResult
from ddict.analysis.rscript import r_script
from ddict.export.preparation import needed as preparation_needed
from ddict.export.preparation import render_preparation
from ddict.export.syntax import render_syntax
from ddict.export.xlsx import _sheet
from ddict.i18n import _, colon, in_language, semicolon
from ddict.project import fingerprint
from ddict.project.model import Project

SHEET_FORBIDDEN = re.compile(r"[\[\]:*?/\\]")
SHEET_LENGTH = 31


@dataclass(frozen=True)
class AnalysisFiles:
    """Where the files of an analysis go: "etude-analyse-1.md", ".xlsx", ".R", and the
    command file and preparation script the R script runs."""

    markdown: Path
    xlsx: Path
    r: Path
    syntax: Path
    preparation: Path


def versions() -> dict[str, str]:
    """The versions the result depends on, written in every report."""
    return {
        "dictionnaire-donnees": __version__,
        "Python": platform.python_version(),
        "SciPy": version("scipy"),
        "statsmodels": version("statsmodels"),
    }


def excluded_reason(result: AnalysisResult) -> str:
    """Why observations were left out of the analysis."""
    if result.options.get("levels"):
        return _("valeurs manquantes, ou hors des deux groupes comparés")
    return _("valeurs manquantes")


def _versions_text() -> str:
    return ", ".join(f"{name} {number}" for name, number in versions().items())


def _variables(result: AnalysisResult, book: Codebook) -> list[tuple[str, str]]:
    """(role, "name (label) ; name (label)") for each role with variables chosen."""
    found = []
    for role in CATALOGUE[result.analysis].roles:
        shown = []
        for name in result.chosen.get(role.key, []):
            label = book.label(name)
            shown.append(label if label.startswith(name) else f"{name} ({label})")
        if shown:
            found.append((ROLES[role.key], semicolon().join(shown)))
    return found


def _options(result: AnalysisResult, book: Codebook) -> list[tuple[str, str]]:
    found = []
    for key, value in result.options.items():
        if key == "levels":
            group = result.chosen["group"][0]
            shown = ", ".join(book.category(group, c) for c in level_codes(value))
        else:
            shown = CHOICES.get(value, str(value))
        found.append((OPTIONS.get(key, key), shown))
    return found


def _about(
    project: Project, result: AnalysisResult, book: Codebook, files: AnalysisFiles, today: str
) -> list[tuple[str, str]]:
    return [
        (_("Analyse"), TITLES[result.analysis]),
        (_("Base"), project.source.name),
        (_("Empreinte du codebook"), fingerprint.short(project)),
        (_("Date"), today),
        (_("Versions"), _versions_text()),
        *_variables(result, book),
        *_options(result, book),
        (_("Observations retenues"), str(result.n_used)),
        (_("Observations écartées"), f"{result.n_excluded} ({excluded_reason(result)})"),
        (_("Script R équivalent"), files.r.name),
    ]


def _cell(text: str) -> str:
    """A table cell; its leading spaces (a category under its variable) kept as no-break
    spaces, which Markdown does not drop."""
    text = str(text).replace("|", "\\|").replace("\n", " ")
    stripped = text.lstrip(" ")
    return "\u00a0" * (2 * (len(text) - len(stripped))) + stripped


def _markdown_table(columns: list[str], rows: list[list[str]]) -> list[str]:
    lines = ["| " + " | ".join(_cell(c) for c in columns) + " |"]
    lines.append("|" + "---|" * len(columns))
    lines += ["| " + " | ".join(_cell(v) for v in row) + " |" for row in rows]
    return lines


@in_language
def render_report(
    project: Project,
    prepared: PreparedData,
    result: AnalysisResult,
    files: AnalysisFiles,
    today: str,
) -> str:
    """The Markdown report of an analysis."""
    book = Codebook(project, prepared)
    lines = [f"# {TITLES[result.analysis]}", ""]
    for name, value in _about(project, result, book, files, today)[1:]:  # after the title
        lines.append(f"- **{name}**{colon()}{value}")
    lines += ["", f"## {_('Résultat')}", ""] + ([result.sentence, ""] if result.sentence else [])
    for table in result.tables:
        lines += [f"### {table.title}", "", *_markdown_table(table.columns, table.rows), ""]
    lines += [f"## {_('Avertissements')}", ""]
    lines += [f"- {x}" for x in result.warnings] or [_("Aucun avertissement.")]
    needed = [files.syntax.name] + ([files.preparation.name] if preparation_needed(project) else [])
    lines += [
        "",
        "## " + _("Refaire l'analyse dans R"),
        "",
        _(
            "Le script {script} refait cette analyse en R de base, avec haven pour lire la "
            "base. Placez-le dans le dossier de la base {base}, avec {needed}, puis exécutez-le "
            "dans ce dossier."
        ).format(script=files.r.name, base=project.source.name, needed=", ".join(needed)),
        "",
        _(
            "Ces résultats viennent des données de l'étude : à l'équipe de décider de ce qui "
            "peut être publié, en particulier les effectifs de moins de 5."
        ),
        "",
    ]
    return "\n".join(lines)


def _sheet_title(title: str, taken: set[str]) -> str:
    base = SHEET_FORBIDDEN.sub("-", title)[:SHEET_LENGTH] or _("Tableau")
    found, index = base, 2
    while found.lower() in taken:
        suffix = f" ({index})"
        found, index = base[: SHEET_LENGTH - len(suffix)] + suffix, index + 1
    taken.add(found.lower())
    return found


@in_language
def report_xlsx(
    project: Project,
    prepared: PreparedData,
    result: AnalysisResult,
    files: AnalysisFiles,
    today: str,
) -> bytes:
    """The tables of an analysis in Excel: the information first, then one sheet per table."""
    book = Codebook(project, prepared)
    workbook = Workbook()
    workbook.remove(workbook.active)
    about = [list(x) for x in _about(project, result, book, files, today)]
    about.append([_("Résultat"), result.sentence])
    about += [[_("Avertissement"), x] for x in result.warnings]
    taken = {_("Informations").lower()}
    _sheet(workbook, _("Informations"), [_("Élément"), _("Valeur")], about)
    for table in result.tables:
        _sheet(workbook, _sheet_title(table.title, taken), table.columns, table.rows)
    out = io.BytesIO()
    workbook.save(out)
    return out.getvalue()


def _r_files(
    project: Project, files: AnalysisFiles, stem: str, result: AnalysisResult, today: str
) -> dict[str, str]:
    found = {
        files.r.name: r_script(
            project, stem, result.analysis, result.chosen, result.options, today
        ),
        files.syntax.name: render_syntax(project, "r", today),
    }
    if preparation_needed(project):
        found[files.preparation.name] = render_preparation(project, "r", today)
    return found


@in_language
def r_bundle(
    project: Project, files: AnalysisFiles, stem: str, result: AnalysisResult, today: str
) -> bytes:
    """The R script and the files it runs, in a zip: what redoes the analysis in R."""
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as archive:
        for name, text in _r_files(project, files, stem, result, today).items():
            info = zipfile.ZipInfo(name, date_time=(1980, 1, 1, 0, 0, 0))
            info.compress_type = zipfile.ZIP_DEFLATED
            archive.writestr(info, text.encode("utf-8"))
    return out.getvalue()


@in_language
def write_report(
    project: Project,
    prepared: PreparedData,
    result: AnalysisResult,
    files: AnalysisFiles,
    stem: str,
    today: str,
) -> list[Path]:
    """Write the report, its Excel tables and the R files next to the project."""
    files.markdown.write_text(
        render_report(project, prepared, result, files, today), encoding="utf-8", newline="\n"
    )
    files.xlsx.write_bytes(report_xlsx(project, prepared, result, files, today))
    written = [files.markdown, files.xlsx]
    for name, text in _r_files(project, files, stem, result, today).items():
        path = files.r.with_name(name)
        path.write_text(text, encoding="utf-8", newline="\n")
        written.append(path)
    return written


def apa_paths(files: AnalysisFiles, document: ApaDocument, keys: list[str] | None = None):
    """Where each APA format goes, next to the report: "etude-analyse-1.docx",
    "etude-analyse-1-graphiques.xlsx"; only the formats asked (all by default) that have
    something to render."""
    from ddict.analysis.formats import available

    stem = files.markdown.stem
    return {
        found.key: (found, files.markdown.with_name(stem + found.suffix))
        for found in available(document)
        if keys is None or found.key in keys
    }


def write_apa(
    files: AnalysisFiles, document: ApaDocument, keys: list[str] | None = None
) -> list[Path]:
    """Write the APA tables and figures (v0.10) in each format asked."""
    written = []
    for found, path in apa_paths(files, document, keys).values():
        path.write_bytes(found.render(document))
        written.append(path)
    return written
