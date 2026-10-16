"""The APA tables and figures in a Word file (v0.10, step 3, D-084, D-085).

Each table is a real Word table, in APA 7 style: its number in bold and its title in italic
above it, horizontal rules only (above and below the header, a short one under a grouped
header, at the bottom), the first column left, the others centred, statistical symbols in
italic, its note below. Each figure is an image, its number, title and note in text that
Word edits like any other. At the end, in small print, the result's sentence and the
codebook's fingerprint, for the record.
"""

from __future__ import annotations

import io

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from ddict import i18n
from ddict.analysis.apa import ApaDocument, ApaFigure, ApaTable, Text, plain
from ddict.analysis.formats import Format, register
from ddict.analysis.formats.figures import WIDTH_CM, image
from ddict.i18n import N_, _

FONT = "Times New Roman"
SIZE = Pt(12)
RULE = "8"  # eighths of a point: a 1-point rule
TEXT_WIDTH_CM = 16.5  # a letter page less its 2.54 cm margins
CHAR_CM = 0.21  # a character of Times New Roman 12, on average
PADDING_CM = 0.4  # the cell's margins
FIRST_LEAST_CM = 3.5


def _rule(cell, *sides: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    borders = properties.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        properties.append(borders)
    for side in sides:
        element = OxmlElement(f"w:{side}")
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), RULE)
        element.set(qn("w:color"), "000000")
        borders.append(element)


def _needs(texts: list[str], whole: bool = True) -> float:
    """The width a column needs for these texts on one line (whole) or word by word."""
    pieces = texts if whole else [w for x in texts for w in x.split()]
    return max(len(x) for x in [*pieces, "    "]) * CHAR_CM + PADDING_CM


def _widths(found: ApaTable) -> list[float]:
    """The numbers never wrap: each column of numbers first gets the width its numbers (and
    the longest word of its header) need; the first column, its labels, takes the rest; what
    is left lets the headers stand on one line, then widens the first column."""
    cells = [[plain(row.cells[j]) for row in found.rows] for j in range(len(found.columns))]
    headers = [plain(c.header) for c in found.columns]
    base = [_needs(cells[j] + [headers[j]], whole=False) for j in range(len(headers))]
    for j in range(1, len(headers)):
        base[j] = max(base[j], _needs(cells[j]))
    full = [max(base[j], _needs(cells[j] + [headers[j]])) for j in range(len(headers))]
    first_need = _needs(cells[0] + [headers[0]])
    others = sum(base[1:])
    first = min(first_need, max(TEXT_WIDTH_CM - others, FIRST_LEAST_CM))
    widths = [first, *base[1:]]
    slack = TEXT_WIDTH_CM - sum(widths)
    if slack > 0:
        deficits = [full[j] - widths[j] for j in range(1, len(widths))]
        wanted = sum(deficits)
        for j, deficit in enumerate(deficits, start=1):
            widths[j] += deficit * min(1.0, slack / wanted) if wanted else 0.0
        slack = TEXT_WIDTH_CM - sum(widths)
        widths[0] += max(0.0, min(slack, first_need - widths[0]))
        slack = TEXT_WIDTH_CM - sum(widths)
        widths = [w + slack / len(widths) for w in widths]
    scale = TEXT_WIDTH_CM / sum(widths)
    return [w * scale for w in widths]


def _no_borders(table) -> None:
    """A table without the grid Word draws by default: APA tables have rules only."""
    properties = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = OxmlElement(f"w:{side}")
        element.set(qn("w:val"), "nil")
        borders.append(element)
    properties.append(borders)


def _runs(paragraph, text: Text, bold: bool = False, italic: bool = False) -> None:
    for run in text:
        added = paragraph.add_run(run.text)
        added.italic = run.italic or italic
        added.bold = bold


def _cell(cell, text: Text, align: str, indent: int = 0) -> None:
    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT if align == "left" else WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.0
    if indent:
        paragraph.paragraph_format.left_indent = Cm(0.5 * indent)
    _runs(paragraph, text)


def _caption(document, label: str, title: str) -> None:
    number = document.add_paragraph()
    number.add_run(label).bold = True
    number.paragraph_format.space_after = Pt(0)
    number.paragraph_format.keep_with_next = True
    heading = document.add_paragraph()
    heading.add_run(title).italic = True
    heading.paragraph_format.keep_with_next = True


def _note(document, note: Text) -> None:
    paragraph = document.add_paragraph()
    if note:
        paragraph.add_run(_("Note.")).italic = True
        paragraph.add_run(" ")
        _runs(paragraph, note)
    paragraph.paragraph_format.space_after = Pt(18)


def _table(document, found: ApaTable) -> None:
    _caption(document, f"{_('Tableau')} {found.number}", found.title)
    above = 1 if found.spanners else 0
    table = document.add_table(rows=above + 1 + len(found.rows), cols=len(found.columns))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    _no_borders(table)
    for j, width in enumerate(_widths(found)):
        for row in table.rows:
            row.cells[j].width = Cm(width)
    if found.spanners:
        for j in range(len(found.columns)):
            _rule(table.cell(0, j), "top")
        for spanner in found.spanners:
            cell = table.cell(0, spanner.first)
            if spanner.last != spanner.first:
                cell = cell.merge(table.cell(0, spanner.last))
            _cell(cell, spanner.label, "center")
            _rule(cell, "bottom")
    for j, column in enumerate(found.columns):
        cell = table.cell(above, j)
        _cell(cell, column.header, column.align)
        _rule(cell, "bottom", *([] if found.spanners else ["top"]))
    for i, row in enumerate(found.rows):
        for j, (text, column) in enumerate(zip(row.cells, found.columns, strict=True)):
            cell = table.cell(above + 1 + i, j)
            _cell(cell, text, column.align, row.indent if j == 0 else 0)
            if i == len(found.rows) - 1:
                _rule(cell, "bottom")
    _note(document, found.note)


def _figure(document, found: ApaFigure, language: str) -> None:
    _caption(document, f"{_('Figure')} {found.number}", found.title)
    picture = document.add_paragraph()
    picture.paragraph_format.keep_with_next = True
    picture.add_run().add_picture(io.BytesIO(image(found, language)), width=Cm(WIDTH_CM))
    _note(document, found.note)


def _apa_heading(document) -> None:
    """An APA level-2 heading: bold, flush left, the document's font, black (not Word's
    blue)."""
    style = document.styles["Heading 2"]
    style.font.name = FONT
    style.font.size = SIZE
    style.font.bold = True
    style.font.italic = False
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.element.rPr.rFonts.set(qn("w:asciiTheme"), "")
    for attribute in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        style.element.rPr.rFonts.attrib.pop(qn(attribute), None)
    style.element.rPr.rFonts.set(qn("w:ascii"), FONT)
    style.element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    style.paragraph_format.space_before = Pt(12)
    style.paragraph_format.space_after = Pt(6)


def _small(paragraph, text: str) -> None:
    """The record at the end, in small grey print."""
    run = paragraph.add_run(text)
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor(0x60, 0x60, 0x60)


def render(found: ApaDocument) -> bytes:
    """The Word file of an analysis's tables and figures."""
    with i18n.using(found.language):
        document = Document()
        style = document.styles["Normal"]
        style.font.name = FONT
        style.font.size = SIZE
        style.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        for section in document.sections:
            section.left_margin = section.right_margin = Cm(2.54)
            section.top_margin = section.bottom_margin = Cm(2.54)
        if found.sections:  # v0.11: the results of several analyses
            _apa_heading(document)
            for section in found.sections:
                document.add_heading(section.heading, level=2)
                if section.sentence:
                    document.add_paragraph(section.sentence)
                for table in section.tables:
                    _table(document, table)
                for figure in section.figures:
                    _figure(document, figure, found.language)
            for line in found.closing:
                _small(document.add_paragraph(), line)
        else:
            for table in found.tables:
                _table(document, table)
            for figure in found.figures:
                _figure(document, figure, found.language)
            text = found.sentence
            if found.fingerprint:
                text += (" " if text else "") + _("Empreinte du codebook : {fingerprint}.").format(
                    fingerprint=found.fingerprint
                )
            _small(document.add_paragraph(), text)
        document.core_properties.author = "dictionnaire-donnees"
        out = io.BytesIO()
        document.save(out)
    return out.getvalue()


def read_text(content: bytes) -> str:
    """Every paragraph and every cell, in order: for the test common to every format."""
    document = Document(io.BytesIO(content))
    parts = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts += [c.text for c in row.cells]
    return "\n".join(parts)


DOCX = register(
    Format(
        key="docx",
        label=N_("Word (tableaux et figures APA)"),
        suffix=".docx",
        mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        tables=True,
        figures=True,
        render=render,
        read_text=read_text,
    )
)
