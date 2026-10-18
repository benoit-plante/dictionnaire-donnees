"""The APA figures as native Excel charts (v0.10, step 4, D-084): one sheet per figure, its
number, title and note in cells, its data in a small table, and the chart drawn from that
table. Everything stays editable: the data, the colours, the axes, the labels; the chart is
then copied into Word, where it stays an editable chart.

Bars of means: one bar per group, with the error bars of the 95 % interval taken from the
table (plus and minus). Grouped bars: one series per row category. Scatter: the points,
and the regression line as a second series of two points. Coefficients: horizontal bars of
β with the error bars of their interval (Excel has no dot plot that keeps the predictors'
names on the axis). The flow of participants (v0.12, D-098): the table of its counts, with
no chart (Excel has no chart for a flow diagram).
"""

from __future__ import annotations

import io
import math

from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, LineChart, Reference, ScatterChart, Series
from openpyxl.chart.data_source import NumDataSource, NumRef
from openpyxl.chart.error_bar import ErrorBars
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.drawing.line import LineProperties
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter, quote_sheetname

from ddict import i18n
from ddict.analysis.apa import ApaDocument, ApaFigure, plain
from ddict.analysis.formats import Format, register
from ddict.i18n import N_, _

FIRST_ROW = 5  # the data table starts here, under the number, the title and a blank line
GREYS = ["737373", "D9D9D9", "A6A6A6", "404040", "F2F2F2", "8C8C8C"]
CHART_COLUMN = "F"
NUMBER_FORMAT = "0.00"


def _sheet_title(figure: ApaFigure) -> str:
    return f"{_('Figure')} {figure.number}"[:31]


def _grey(series, color: str, outline: bool = True) -> None:
    series.graphicalProperties = GraphicalProperties(solidFill=color)
    if outline:
        series.graphicalProperties.line = LineProperties(solidFill="000000", w=9525)


def _errors(sheet, column_plus: int, column_minus: int, first: int, last: int):
    """Error bars from two columns of the table; along the values (a bar chart has no other
    direction: errDir is for scatter charts only)."""

    def ref(column: int) -> NumDataSource:
        letter = get_column_letter(column)
        return NumDataSource(
            numRef=NumRef(f=f"{quote_sheetname(sheet.title)}!${letter}${first}:${letter}${last}")
        )

    return ErrorBars(
        errBarType="both",
        errValType="cust",
        plus=ref(column_plus),
        minus=ref(column_minus),
        noEndCap=False,
    )


def _axes(chart, figure: ApaFigure, x_title: str, y_title: str) -> None:
    chart.x_axis.title = x_title or None
    chart.y_axis.title = y_title or None
    for axis in (chart.x_axis, chart.y_axis):
        axis.delete = False  # without it, recent Excel hides the axes of an openpyxl chart
        axis.majorGridlines = None
        axis.number_format = "General"  # "20", not "20,00"
        if axis.title is not None:
            axis.title.overlay = False  # beside the tick labels, not over them
    if chart.legend is not None:
        chart.legend.overlay = False
    if figure.y_min is not None:
        chart.y_axis.scaling.min = figure.y_min
    if figure.y_max is not None:
        chart.y_axis.scaling.max = figure.y_max
    chart.title = None  # the title is in the cells, as APA puts it above the figure
    chart.width, chart.height = 16, 9.5


def _write_rows(sheet, header: list[str], rows: list[list]) -> int:
    sheet.append([])
    for column, text in enumerate(header, start=1):
        cell = sheet.cell(row=FIRST_ROW - 1, column=column, value=text)
        cell.font = Font(bold=True)
    for i, row in enumerate(rows):
        for column, value in enumerate(row, start=1):
            cell = sheet.cell(row=FIRST_ROW + i, column=column, value=value)
            if isinstance(value, float):
                cell.number_format = NUMBER_FORMAT
    return FIRST_ROW + len(rows) - 1


def _bars(sheet, figure: ApaFigure):
    series = figure.series[0]
    rows = [
        [name, mean, high - mean, mean - low]
        for name, mean, low, high in zip(
            figure.categories, series.values, series.low, series.high, strict=True
        )
    ]
    last = _write_rows(
        sheet, [figure.x_label, figure.y_label, _("IC 95 % (+)"), _("IC 95 % (−)")], rows
    )
    chart = BarChart()
    chart.type, chart.grouping, chart.legend = "col", "clustered", None
    chart.add_data(
        Reference(sheet, min_col=2, min_row=FIRST_ROW - 1, max_row=last), titles_from_data=True
    )
    chart.set_categories(Reference(sheet, min_col=1, min_row=FIRST_ROW, max_row=last))
    _grey(chart.series[0], "B3B3B3")
    chart.series[0].errBars = _errors(sheet, 3, 4, FIRST_ROW, last)
    return chart, last


def _line(sheet, figure: ApaFigure):
    """v0.14: the means of the measures in their order, joined, with their intervals; or
    several series (a moderation's simple slopes), with a legend."""
    single = len(figure.series) == 1 and bool(figure.series[0].low)
    rows = []
    for i, name in enumerate(figure.categories):
        row = [name, *[s.values[i] for s in figure.series]]
        if single:
            series = figure.series[0]
            row += [series.high[i] - series.values[i], series.values[i] - series.low[i]]
        rows.append(row)
    header = [figure.x_label, *[s.name for s in figure.series]]
    if single:
        header = [figure.x_label, figure.y_label, _("IC 95 % (+)"), _("IC 95 % (−)")]
    last = _write_rows(sheet, header, rows)
    chart = LineChart()
    if len(figure.series) == 1:
        chart.legend = None
    chart.add_data(
        Reference(
            sheet, min_col=2, max_col=1 + len(figure.series), min_row=FIRST_ROW - 1, max_row=last
        ),
        titles_from_data=True,
    )
    chart.set_categories(Reference(sheet, min_col=1, min_row=FIRST_ROW, max_row=last))
    styles = [("circle", "000000", None), ("square", "595959", "dash"),
              ("triangle", "8C8C8C", "sysDot")]  # fmt: skip
    for line, (symbol, color, dash) in zip(chart.series, styles, strict=False):
        line.graphicalProperties.line.solidFill = color
        if dash:
            line.graphicalProperties.line.dashStyle = dash
        line.marker.symbol = symbol
        line.marker.graphicalProperties = GraphicalProperties(solidFill=color)
        line.smooth = False
    if single:
        chart.series[0].errBars = _errors(sheet, 3, 4, FIRST_ROW, last)
    return chart, last


def _grouped(sheet, figure: ApaFigure):
    rows = [
        [category, *[s.values[i] for s in figure.series]]
        for i, category in enumerate(figure.categories)
    ]
    last = _write_rows(sheet, [figure.x_label, *[s.name for s in figure.series]], rows)
    chart = BarChart()
    chart.type, chart.grouping = "col", "clustered"
    chart.add_data(
        Reference(
            sheet, min_col=2, max_col=1 + len(figure.series), min_row=FIRST_ROW - 1, max_row=last
        ),
        titles_from_data=True,
    )
    chart.set_categories(Reference(sheet, min_col=1, min_row=FIRST_ROW, max_row=last))
    for i, series in enumerate(chart.series):
        _grey(series, GREYS[i % len(GREYS)])
    chart.legend.position = "b"
    return chart, last


def _scatter(sheet, figure: ApaFigure):
    rows = [[x, y] for x, y in figure.points]
    last = _write_rows(sheet, [figure.x_label, figure.y_label], rows)
    chart = ScatterChart()
    chart.style, chart.legend = 13, None
    xs = [p[0] for p in figure.points]
    chart.x_axis.scaling.min = _floor(min(xs), max(xs))  # not from 0: where the points are
    points = Series(
        Reference(sheet, min_col=2, min_row=FIRST_ROW, max_row=last),
        Reference(sheet, min_col=1, min_row=FIRST_ROW, max_row=last),
    )
    points.marker.symbol, points.marker.size = "circle", 4
    points.marker.graphicalProperties = GraphicalProperties(solidFill="808080")
    points.marker.graphicalProperties.line = LineProperties(noFill=True)
    points.graphicalProperties = GraphicalProperties()
    points.graphicalProperties.line = LineProperties(noFill=True)
    chart.series.append(points)
    if figure.line is not None:
        intercept, slope = figure.line
        xs = [min(p[0] for p in figure.points), max(p[0] for p in figure.points)]
        column = 4
        sheet.cell(row=FIRST_ROW - 1, column=column, value=_("Droite (x)")).font = Font(bold=True)
        sheet.cell(row=FIRST_ROW - 1, column=column + 1, value=_("Droite (y)")).font = Font(
            bold=True
        )
        for i, x in enumerate(xs):
            sheet.cell(row=FIRST_ROW + i, column=column, value=x)
            sheet.cell(row=FIRST_ROW + i, column=column + 1, value=intercept + slope * x)
        line = Series(
            Reference(sheet, min_col=column + 1, min_row=FIRST_ROW, max_row=FIRST_ROW + 1),
            Reference(sheet, min_col=column, min_row=FIRST_ROW, max_row=FIRST_ROW + 1),
        )
        line.marker.symbol = "none"
        line.graphicalProperties = GraphicalProperties()
        line.graphicalProperties.line = LineProperties(solidFill="000000", w=19050)
        chart.series.append(line)
    return chart, last


def _coefficients(sheet, figure: ApaFigure):
    series = figure.series[0]
    rows = [
        [name, beta, high - beta, beta - low]
        for name, beta, low, high in zip(
            figure.categories, series.values, series.low, series.high, strict=True
        )
    ][::-1]  # Excel draws the first category at the bottom: the first predictor ends on top
    header = [_("Prédicteur"), series.name, _("IC 95 % (+)"), _("IC 95 % (−)")]
    last = _write_rows(sheet, header, rows)
    chart = BarChart()
    chart.type, chart.grouping, chart.legend = "bar", "clustered", None
    chart.add_data(
        Reference(sheet, min_col=2, min_row=FIRST_ROW - 1, max_row=last), titles_from_data=True
    )
    chart.set_categories(Reference(sheet, min_col=1, min_row=FIRST_ROW, max_row=last))
    chart.x_axis.tickLblPos = "low"  # the names on the left, not over the negative bars
    _grey(chart.series[0], "B3B3B3")
    chart.series[0].invertIfNegative = False  # a negative bar stays grey, not white
    chart.series[0].errBars = _errors(sheet, 3, 4, FIRST_ROW, last)
    if figure.log_scale:  # v0.14: odds ratios, the bars growing from 1 either way
        chart.y_axis.scaling.logBase = 10
        chart.x_axis.crossesAt = 1.0
    return chart, last


def _floor(low: float, high: float) -> float:
    """A round value just below `low`, for the start of an axis: 25 for 25.3 to 54.5."""
    span = max(high - low, 1e-9)
    step = 10 ** math.floor(math.log10(span))
    step = step / 2 if span / step < 4 else step
    return math.floor(low / step) * step


def _flow(sheet, figure: ApaFigure):
    """The counts of the flow diagram, editable: the answers received, each rule, the
    sample analysed."""
    rows = [["", figure.x_label, None, figure.flow_start]]
    rows += [[s.rule, s.reason, s.excluded, s.remaining] for s in figure.flow]
    excluded = sum(s.excluded for s in figure.flow)
    analysed = figure.flow[-1].remaining if figure.flow else figure.flow_start
    rows.append(["", figure.y_label, excluded, analysed])
    last = _write_rows(sheet, [_("Étape"), _("Raison"), _("Exclues"), _("Restantes")], rows)
    for column in (1, 2, 3, 4):
        sheet.cell(row=last, column=column).font = Font(bold=True)
    return None, last


DRAW = {
    "bars": _bars,
    "grouped_bars": _grouped,
    "scatter": _scatter,
    "coefficients": _coefficients,
    "flow": _flow,
    "line": _line,
}


def _figure(book: Workbook, figure: ApaFigure) -> None:
    sheet = book.create_sheet(_sheet_title(figure))
    sheet["A1"] = f"{_('Figure')} {figure.number}"
    sheet["A1"].font = Font(bold=True)
    sheet["A2"] = figure.title
    sheet["A2"].font = Font(italic=True)
    chart, last = DRAW[figure.kind](sheet, figure)
    if chart is not None:  # the flow has its table only
        if figure.kind == "coefficients":  # horizontal bars: Excel's axes the other way round
            _axes(chart, figure, figure.y_label, figure.x_label)
        else:
            _axes(chart, figure, figure.x_label, figure.y_label)
        sheet.add_chart(chart, f"{CHART_COLUMN}{FIRST_ROW - 1}")
    note = sheet.cell(row=last + 2, column=1, value=f"{_('Note.')} {plain(figure.note)}".strip())
    note.font = Font(italic=False)
    sheet.column_dimensions["A"].width = 28
    for letter in "BCDE":
        sheet.column_dimensions[letter].width = 14
    if figure.kind == "flow":  # the reasons
        sheet.column_dimensions["A"].width = 10
        sheet.column_dimensions["B"].width = 48


def render(found: ApaDocument) -> bytes:
    """A workbook of the analysis's figures, one sheet each, with its native chart."""
    if not found.figures:
        raise ValueError("No figure to render: see formats.available()")
    with i18n.using(found.language):
        book = Workbook()
        book.remove(book.active)
        for figure in found.figures:
            _figure(book, figure)
        out = io.BytesIO()
        book.save(out)
    return out.getvalue()


def read_text(content: bytes) -> str:
    """The text of every cell, sheet by sheet."""
    book = load_workbook(io.BytesIO(content))
    return "\n".join(
        str(value)
        for sheet in book.worksheets
        for row in sheet.iter_rows(values_only=True)
        for value in row
        if value is not None
    )


XLSX_CHARTS = register(
    Format(
        key="xlsx",
        label=N_("Excel (graphiques modifiables)"),
        suffix="-graphiques.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        tables=False,
        figures=True,
        render=render,
        read_text=read_text,
    )
)
