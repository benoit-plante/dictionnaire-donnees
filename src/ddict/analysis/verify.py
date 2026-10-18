"""The checker of the figures of a manuscript (v0.15, step 5, D-117, D-118;
docs/verification-depot.md, § 2.2 to 2.5): each statistic read in the manuscript is compared
with what the tool would write itself, the draft of the article recomputed on the database
with the decisions in force, section by section; without the database, with the sentences
the notebook keeps. Each value is classed: it agrees, differs (with the value expected),
has its sign reversed, has no origin, or was not read. The report and a copy of the
manuscript with Word comments say where; the manuscript itself is never changed.
"""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from pathlib import Path

from ddict.analysis import manuscript as m
from ddict.analysis.apa import plain
from ddict.i18n import N_, Texts, _

CLASSES = Texts(
    {
        "agrees": N_("concorde"),
        "differs": N_("diffère"),
        "sign": N_("signe inversé"),
        "no_origin": N_("sans origine"),
        "unread": N_("non reconnue"),
    }
)
ORDER = ("differs", "sign", "no_origin", "unread", "agrees")
# The same statistic under two names: d_z is a d, a total N a count.
SAME = {"dz": "d", "g": "d", "N": "n"}
# A test statistic: in a sentence or a table row, it says which analysis the others are of.
ANCHORS = {"t", "F", "chi2", "H", "U", "V", "z", "W"}
# Whose sign depends on the order of the groups, which an author may have reversed.
SIGNED = {"t", "z", "d", "dz", "g", "diff"}
SECTION_ID = re.compile(r"^(A\d+|H\d+)\b")
NOT_READ = N_(
    "Non lus : les figures et les images, les zones de texte, les en-têtes et pieds de page, "
    "les notes de bas de page, les commentaires ; les pourcentages."
)


@dataclass(frozen=True)
class Reference:
    """What the tool writes in a section of its draft: an analysis ("A2"), a hypothesis
    ("H1"), a part of the Method ("participants"…)."""

    key: str
    heading: str
    values: tuple[m.Value, ...]


@dataclass(frozen=True)
class Checked:
    block: m.Block
    kind: str  # a key of CLASSES
    value: m.Value | None = None  # None for a writing not read
    unread: str = ""
    expected: m.Value | None = None
    reference: str = ""  # the section's key: "A1", "H2", "participants"
    others: tuple[str, ...] = ()  # other sections the value could come from
    pending: bool = False  # the analysis waits for a confirmation


@dataclass
class Verification:
    manuscript: str
    recomputed: bool  # on the database; otherwise against the notebook
    checked: list[Checked] = field(default_factory=list)
    thresholds: int = 0  # "p < ,05" of a rule, not compared

    def counts(self) -> dict[str, int]:
        found = dict.fromkeys(ORDER, 0)
        for c in self.checked:
            found[c.kind] += 1
        return found

    def problems(self) -> list[Checked]:
        return [c for c in self.checked if c.kind != "agrees"]


# --- The references: the tool's own draft. ---


def _values_of(text: str) -> list[m.Value]:
    found: list[m.Value] = []
    for sentence in m.sentences(text):
        found += m.statistics(sentence)[0]
    return found


def _section_values(section) -> list[m.Value]:
    found = _values_of(section.sentence)
    for paragraph in section.paragraphs:
        found += _values_of(plain(paragraph))
    for table in section.tables:
        header = [plain(c.header) for c in table.columns]
        for row in table.rows:
            found += m.row_values(header, [plain(cell) for cell in row.cells])
        found += _values_of(plain(table.note))
    for figure in section.figures:
        found += _values_of(plain(figure.note))
    return found


def references(project, prepared, notebook, *, language=None, today=None):
    """The references, and the analyses waiting for a confirmation."""
    from ddict.analysis import notebook as nb
    from ddict.analysis import notebook_export as export

    pending: list[str] = []
    if prepared is None:
        found = [
            Reference(e.id, export.heading(e), tuple(_values_of(e.computed.sentence)))
            for e in notebook.entries
            if e.computed is not None
        ]
        return found, pending
    document, _analyses = export._results(project, prepared, notebook, language=language,
                                          today=today, pending=pending)  # fmt: skip
    found = []
    for section in document.sections:
        values = _section_values(section)
        if values:
            found.append(Reference(section.key, section.heading, tuple(values)))
    # A hypothesis waits when its analysis does.
    for entry in notebook.entries:
        if entry.id in pending and entry.hypothesis:
            pending.append(entry.hypothesis)
    pending += [nb.METHOD] if nb.METHOD in pending else []
    return found, pending


# --- The comparison. ---


def _stat(value: m.Value) -> str:
    return SAME.get(value.stat, value.stat)


def _close(a: float, b: float, decimals: int) -> bool:
    return abs(a - b) <= 0.5 * 10.0**-decimals + 1e-9


def _same_kind(value: m.Value, other: m.Value) -> bool:
    if _stat(value) != _stat(other):
        return False
    if value.stat in ("ci_low", "ci_high") and value.of and other.of:
        return SAME.get(value.of, value.of) == SAME.get(other.of, other.of)
    return True


def agrees(value: m.Value, other: m.Value) -> bool:
    """A written value agrees with the tool's: the same statistic, the same degrees of
    freedom, and within half a unit of the coarser of their last decimals; "p < ,001"
    agrees with a p under ,001."""
    if not _same_kind(value, other):
        return False
    if value.df and other.df:
        if len(value.df) != len(other.df):
            return False
        if not all(_close(a.value, b.value, min(a.decimals, b.decimals))
                   for a, b in zip(value.df, other.df, strict=True)):  # fmt: skip
            return False
    decimals = min(value.decimals, other.decimals)
    if value.relation == "<":
        return other.value <= value.value + 1e-12 if other.relation in ("<", "=") else False
    if value.relation == ">":
        return other.value >= value.value - 1e-12 if other.relation in (">", "=") else False
    if other.relation != "=":
        return False
    return _close(value.value, other.value, decimals)


def _sign_only(value: m.Value, other: m.Value) -> bool:
    return (
        value.stat in SIGNED
        and value.relation == "="
        and value.value != 0
        and agrees(m.Value(value.stat, -value.value, value.decimals, "=", value.text,
                           value.df, value.of), other)
    )  # fmt: skip


def _compare(value: m.Value, ref: Reference, pending: set[str]) -> tuple | None:
    """(kind, expected) against one section; None when it has no such statistic."""
    same = [x for x in ref.values if _same_kind(value, x)]
    if not same:
        return None
    for other in same:
        if agrees(value, other):
            return "agrees", other
    for other in same:
        if _sign_only(value, other):
            return "sign", other
    with_df = [x for x in same if not value.df or not x.df or len(x.df) == len(value.df)]
    closest = min(with_df or same, key=lambda x: abs(x.value - value.value))
    return "differs", closest


class _Sections:
    def __init__(self, refs: list[Reference]):
        self.refs = refs
        self.by_heading = {r.heading: r for r in refs}
        self.by_key = {r.key: r for r in refs}

    def of(self, heading: str) -> Reference | None:
        if heading in self.by_heading:
            return self.by_heading[heading]
        found = SECTION_ID.match(heading.strip())
        return self.by_key.get(found.group(1)) if found else None


def _is_analysis(ref: Reference) -> bool:
    return bool(SECTION_ID.match(ref.key))


def _analyses_first(refs: list[Reference]) -> list[Reference]:
    """The sections of an analysis or a hypothesis before the others (the Discussion, the
    abstract repeat their figures)."""
    return sorted(refs, key=lambda r: not _is_analysis(r))


def _same_df(anchor: m.Value, other: m.Value) -> bool:
    return (
        _same_kind(anchor, other)
        and bool(anchor.df)
        and len(anchor.df) == len(other.df)
        and all(_close(d.value, e.value, min(d.decimals, e.decimals))
                for d, e in zip(anchor.df, other.df, strict=True))
    )  # fmt: skip


def _of_anchor(anchor: m.Value, refs: list[Reference], own) -> Reference | None:
    """The section a test statistic designates: where it agrees (its own section, else an
    analysis, else the only one); else the only analysis with the same statistic and the
    same degrees of freedom ("t(818) = 2,28" for −2,28: a sign reversed)."""
    agreeing = _analyses_first([r for r in refs if any(agrees(anchor, x) for x in r.values)])
    if agreeing:
        return own if own in agreeing else agreeing[0]
    same = _analyses_first([r for r in refs if any(_same_df(anchor, x) for x in r.values)])
    analyses = [r for r in same if _is_analysis(r)]
    if own is not None and own in same:
        return own
    if len(analyses) == 1 or (len(same) == 1):
        return (analyses or same)[0]
    return None


def _check_block(block: m.Block, sections: _Sections, pending: set[str]) -> list[Checked]:
    """Each value compared with the section of the test statistic before it in the block
    (a sentence may report two analyses); without one, with the block's own section; a
    value beside a test statistic of no known origin has no origin either."""
    refs = sections.refs
    own = sections.of(block.section)
    found = []
    current = own
    anchored = False
    for value in block.values:
        if value.stat in ANCHORS or (value.stat == "r" and value.df):
            anchored = True
            current = _of_anchor(value, refs, own)
        if value.threshold and current is None:
            found.append(Checked(block, "threshold", value))
            continue
        result = _compare(value, current, pending) if current is not None else None
        if result is not None:
            kind, expected = result
            found.append(Checked(block, kind, value, expected=expected, reference=current.key,
                                 pending=current.key in pending))  # fmt: skip
            continue
        if anchored and current is None:
            found.append(Checked(block, "no_origin", value))
            continue
        matching = _analyses_first(
            [r for r in refs if any(agrees(value, x) and not x.threshold for x in r.values)]
        )
        if matching:
            first = matching[0]
            expected = next(x for x in first.values if agrees(value, x) and not x.threshold)
            found.append(Checked(block, "agrees", value, expected=expected, reference=first.key,
                                 others=tuple(r.key for r in matching[1:]),
                                 pending=first.key in pending))  # fmt: skip
        else:
            found.append(Checked(block, "no_origin", value))
    for text in block.unread:
        found.append(Checked(block, "unread", unread=text))
    return found


def verify(project, prepared, notebook, path: Path | str, *, language=None, today=None):
    """The manuscript checked against the tool's draft (D-118)."""
    refs, pending = references(project, prepared, notebook, language=language, today=today)
    sections = _Sections(refs)
    result = Verification(Path(path).name, prepared is not None)
    for block in m.read(path):
        for checked in _check_block(block, sections, set(pending)):
            if checked.kind == "threshold":
                result.thresholds += 1
            else:
                result.checked.append(checked)
    return result


# --- What it gives: the report and the commented copy. ---


def shown(value: m.Value | None, like: m.Value | None = None) -> str:
    """A value as the tool writes it, with the decimal mark of the manuscript's value."""
    if value is None:
        return "—"
    text = f"{value.value:.{value.decimals}f}"
    if like is not None and "," in like.text:
        text = text.replace(".", ",")
    if like is not None and re.search(r"(?<!\d)[.,]\d", like.text):  # ",25": no leading zero
        text = re.sub(r"^(-?)0(?=[.,])", r"\1", text)
    prefix = "< " if value.relation == "<" else "> " if value.relation == ">" else ""
    return prefix + text.replace("-", "−")


def _statistic(value: m.Value | None) -> str:
    if value is None:
        return ""
    name = {"ci_low": _("borne basse de l'IC"), "ci_high": _("borne haute de l'IC")}.get(
        value.stat, value.text.split("=")[0].split("<")[0].strip() or value.stat
    )
    if value.stat in ("ci_low", "ci_high") and value.of:
        name += f" ({value.of})"
    return name


def comment(checked: Checked) -> str:
    """The comment a value gets in the copy of the manuscript."""
    where = checked.reference + (_(", analyse à confirmer") if checked.pending else "")
    if checked.kind == "differs":
        return _("Diffère : l'outil écrit {expected} ({where}).").format(
            expected=shown(checked.expected, checked.value), where=where
        )
    if checked.kind == "sign":
        return _(
            "Signe inversé : l'outil écrit {expected} ({where}) ; l'ordre des groupes est "
            "peut-être inversé."
        ).format(expected=shown(checked.expected, checked.value), where=where)
    if checked.kind == "no_origin":
        return _("Sans origine : aucune analyse du carnet ni la Méthode ne donne cette valeur.")
    return _("Non reconnue : une écriture que l'outil ne sait pas lire sans deviner.")


def rows(result: Verification) -> list[list[str]]:
    """One row per value, in the order of the document."""
    return [
        [
            c.block.where,
            _statistic(c.value) if c.value else c.unread,
            c.value.text if c.value else c.unread,
            shown(c.expected, c.value),
            c.reference + (" *" if c.pending else ""),
            CLASSES[c.kind],
        ]
        for c in result.checked
    ]


HEADER = (N_("Emplacement"), N_("Statistique"), N_("Écrit"), N_("Attendu"), N_("Analyse"),
          N_("Classe"))  # fmt: skip


def summary_lines(result: Verification) -> list[str]:
    counts = result.counts()
    lines = [
        _("Manuscrit : {name}.").format(name=result.manuscript),
        _("Comparé à l'ébauche que l'outil écrit, recalculée sur la base.")
        if result.recomputed
        else _("Base introuvable : comparaison aux valeurs gardées dans le carnet, sans recalcul."),
        " ; ".join(f"{CLASSES[k]} : {counts[k]}" for k in ORDER) + ".",
    ]
    if result.thresholds:
        lines.append(_("Seuils de règle, non comparés (« p < ,05 ») : {n}.").format(
            n=result.thresholds))  # fmt: skip
    if any(c.pending for c in result.checked):
        lines.append(_("* analyse à confirmer dans le carnet : comparée à son nouveau calcul."))
    lines.append(_(NOT_READ))
    return lines


def render_report(result: Verification) -> str:
    lines = [f"# {_('Vérification des chiffres')}", ""]
    lines += [f"- {line}" for line in summary_lines(result)] + [""]
    header = [_(h) for h in HEADER]
    lines += ["| " + " | ".join(header) + " |", "|" + "---|" * len(header)]
    for row in rows(result):
        lines.append("| " + " | ".join(cell.replace("|", "\\|") for cell in row) + " |")
    return "\n".join(lines) + "\n"


def report_xlsx(result: Verification) -> bytes:
    import io

    from openpyxl import Workbook

    book = Workbook()
    sheet = book.active
    sheet.title = _("Vérification")
    for line in summary_lines(result):
        sheet.append([line])
    sheet.append([])
    sheet.append([_(h) for h in HEADER])
    for row in rows(result):
        sheet.append(row)
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()


def _runs_of(paragraph, text: str):
    """The runs that hold a written number, or the whole paragraph's."""
    numbers = m.NUMBER.findall(text)
    target = numbers[-1].strip() if numbers else text
    for run in paragraph.runs:
        if target and target in run.text:
            return [run]
    return list(paragraph.runs)


def commented_copy(source: Path | str, target: Path | str, result: Verification) -> int:
    """A copy of the manuscript with a comment on each value that does not agree; the
    manuscript itself is never changed. Returns the number of comments."""
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    document = Document(str(source))
    paragraphs, tables = {}, {}
    rank = table_rank = 0
    for child in document.element.body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            paragraph = Paragraph(child, document)
            if paragraph.text.strip():
                rank += 1
                paragraphs[rank] = paragraph
        elif tag == "tbl":
            table_rank += 1
            tables[table_rank] = Table(child, document)
    count = 0
    for checked in result.problems():
        block = checked.block
        text = checked.value.text if checked.value else checked.unread
        if block.table:
            row = tables[block.table].rows[block.row]
            cell = row.cells[max(checked.value.cell, 0) if checked.value else 0]
            runs = [r for p in cell.paragraphs for r in p.runs]
        else:
            runs = _runs_of(paragraphs[block.paragraph], text)
        if not runs:
            continue
        document.add_comment(runs, text=comment(checked), author="dictionnaire-donnees",
                             initials="dd")  # fmt: skip
        count += 1
    document.save(target if hasattr(target, "write") else str(target))
    return count


def paths(manuscript: Path | str) -> dict[str, Path]:
    """ "manuscrit-verification.md", ".xlsx" and "manuscrit-verifie.docx", next to it."""
    path = Path(manuscript)
    return {
        "md": path.with_name(f"{path.stem}-verification.md"),
        "xlsx": path.with_name(f"{path.stem}-verification.xlsx"),
        "docx": path.with_name(f"{path.stem}-verifie.docx"),
    }


def write(result: Verification, manuscript: Path | str) -> list[Path]:
    found = paths(manuscript)
    with found["md"].open("w", encoding="utf-8", newline="\n") as stream:
        stream.write(render_report(result))
    found["xlsx"].write_bytes(report_xlsx(result))
    commented_copy(manuscript, found["docx"], result)
    return [found["md"], found["xlsx"], found["docx"]]
