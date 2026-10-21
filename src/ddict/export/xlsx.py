"""The Excel codebook: one sheet per kind of information."""

from __future__ import annotations

from collections.abc import Sequence
from pathlib import Path
from typing import BinaryIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from ddict.export import wording as w
from ddict.i18n import N_, TextList, _, colon, in_language, semicolon
from ddict.project import aside, fingerprint
from ddict.project.model import Profile, Project, TrackedValue

MAX_WIDTH = 60


def _cell_value(value: object) -> object:
    return "" if value is None else value


def _sheet(book: Workbook, title: str, header: Sequence[str], rows: list[list]) -> Worksheet:
    sheet = book.create_sheet(title)
    sheet.append(list(header))
    for row in rows:
        sheet.append([_cell_value(v) for v in row])
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            # A label starting with "=" must stay text, never become a formula.
            if cell.data_type == "f":
                cell.data_type = "s"
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    sheet.freeze_panes = "B2"
    if rows:
        sheet.auto_filter.ref = sheet.dimensions
    for index, title_text in enumerate(header, start=1):
        values = [str(title_text)] + [str(_cell_value(r[index - 1])) for r in rows]
        width = max(len(line) for v in values for line in v.split("\n"))
        sheet.column_dimensions[get_column_letter(index)].width = min(width + 2, MAX_WIDTH)
    return sheet


def _about(project: Project, kind: w.ExportKind, omit_groups: tuple[str, ...]) -> list[list]:
    source = project.source
    kind_text = (
        _("finale (valeurs acceptées seulement)")
        if kind == "final"
        else _("de travail (propositions de l'IA en attente incluses)")
    )
    if source is None:
        data = [[_("Fichier de données"), _(w.PLANNED_DICTIONARY)]]
    else:
        data = [
            [_("Fichier de données"), source.name],
            [_("Empreinte SHA-256"), source.sha256],
            [_("Lignes"), source.n_rows],
            [_("Variables"), source.n_columns],
            [_("Étiquette du fichier"), source.label],
        ]
    return [
        *data,
        [_("Version"), kind_text],
        [
            _("Confidentialité"),
            _(
                "Effectifs inférieurs à {n} masqués, ainsi qu'un second effectif quand le "
                "premier pourrait être recalculé par soustraction ; variables sensibles "
                "décrites sans statistiques."
            ).format(n=project.settings.min_cell_count),
        ],
        [_("Format du projet"), project.schema_version],
        [_("Empreinte du codebook"), fingerprint.short(project)],
        *(
            [[_("Codes des vides dans les scripts SPSS et R"), w.blank_codes_text(project)]]
            if w.blank_codes_text(project)
            else []
        ),
        *[[name, value] for name, value in w.dataset_facts(project)],
        *[
            [
                _("Complétude : {element}").format(element=name),
                count + (f"{semicolon()}{_('à documenter')}{colon()}{missing}" if missing else ""),
            ]
            for name, count, missing in w.completeness_rows(project)
        ],
        *_questionnaire_rows(project, omit_groups),
    ]


def _questionnaire_rows(project: Project, omit_groups: tuple[str, ...]) -> list[list]:
    rows = [[name, value] for name, value in w.questionnaire_facts(project)]
    if not rows:
        return []
    omitted = sorted(
        {g.group_id for g in project.groups if g.omit_question_text} | set(omit_groups)
    )
    rows.append(
        [
            _("Texte des questions"),
            _("omis pour les groupes {groups} (instrument protégé, D-014)").format(
                groups=", ".join(omitted)
            )
            if omitted
            else _("exporté pour toutes les variables rattachées"),
        ]
    )
    return rows


def _statistics(p: Profile | None) -> list:
    """The statistics columns; blank for a variable planned before the database (v0.20)."""
    if p is None:
        return [None] * 12
    return [
        p.n_valid,
        p.n_system_missing,
        p.n_declared_missing,
        w.pct_missing(p),
        p.n_distinct,
        p.min,
        p.max,
        p.mean,
        p.sd,
        p.median,
        p.date_min,
        p.date_max,
    ]


def _variables(project: Project, kind: w.ExportKind, omitted: set[str]) -> list[list]:
    rows = []
    renamed = w.renamed(project, kind)
    presented = w.has_presentation(project, kind)  # v0.21: two columns, only when needed
    for e in aside.kept(project):
        reference = w.question_reference(project, e, kind)
        p = e.profile
        sourced_range = ""
        if e.theoretical_range:
            r = e.theoretical_range
            sourced_range = f"{w.number(r.min)} – {w.number(r.max)} ({r.source})"
        thresholds = "\n".join(
            f"{w.number(t.value)}{colon()}{t.meaning} ({t.source})" for t in e.thresholds
        )
        rows.append(
            [
                e.position,
                w.name_text(e, kind),
                *([e.name] if renamed else []),  # v0.7: the original name, when renamed
                w.label_text(e, kind),
                w.text(e.description, kind),
                ""
                if reference is None
                else reference.value + (f" [{_('proposé')}]" if reference.pending else ""),
                w.question_text(project, e, kind, omitted),
                *(
                    [w.asked_if_text(project, e, kind), w.presentation_text(project, e)]
                    if presented
                    else []
                ),
                w.TYPES.get(e.data_type, e.data_type),
                w.level_text(e, kind),
                e.file.display_format,
                w.text(e.unit, kind),
                *_statistics(p),
                w.missing_codes(e),
                w.privacy_text(e),
                w.text(e.group_id, kind),
                None if p is None else p.item_rest_r,
                sourced_range,
                thresholds,
                "\n".join(e.references),
                "\n".join(w.issues(e, kind)),
                e.notes,
            ]
        )
    return rows


VARIABLE_HEADER = TextList(
    [
        N_("Position"),
        N_("Nom"),
        N_("Libellé"),
        N_("Description"),
        N_("Question"),
        N_("Texte de la question"),
        N_("Type"),
        N_("Niveau de mesure"),
        N_("Format"),
        N_("Unité"),
        N_("Valides"),
        N_("Cases vides"),
        N_("Manquants déclarés"),
        N_("% manquants"),
        N_("Valeurs distinctes"),
        N_("Min"),
        N_("Max"),
        N_("Moyenne"),
        N_("Écart-type"),
        N_("Médiane"),
        N_("Date min"),
        N_("Date max"),
        N_("Codes manquants déclarés"),
        N_("Confidentialité"),
        N_("Groupe"),
        N_("Corrélation avec le reste du groupe"),
        N_("Étendue théorique (source)"),
        N_("Seuils (source)"),
        N_("Références"),
        N_("Signalements"),
        N_("Notes"),
    ]
)


def _values(project: Project, kind: w.ExportKind) -> list[list]:
    rows = []
    for e in aside.kept(project):
        source = w.value_labels_source(e, kind)
        for code, label, count, is_missing in w.value_rows(e):
            rows.append(
                [
                    w.final_name(e),
                    code,
                    label,
                    count,
                    _("oui") if is_missing else "",
                    source if label else "",
                ]
            )
        if source == _(w.LABELS_PROPOSED):  # proposals are not mixed with the counts
            for item in e.value_labels.value or []:
                rows.append([w.final_name(e), w.code_text(item.code), item.label, "", "", source])
    return rows


def _groups(project: Project, kind: w.ExportKind) -> list[list]:
    rows = []
    names = w.names_of(project)
    for g in project.groups:
        members = w.shown(g.members, kind)
        suspects = [
            e.name
            for e in project.variables
            if e.profile is not None
            and e.profile.reverse_suspect
            and members
            and e.name in members.value
        ]
        rows.append(
            [
                g.group_id,
                w.group_description(g, kind),
                ", ".join(names.get(m, m) for m in members.value) if members else "",
                w.basis_text(g.basis),
                ", ".join(names[s] for s in suspects),
                ", ".join(names.get(r, r) for r in g.reverse_items),
                w.text(g.ai_hypothesis, kind),
                f"{g.instrument_name.value} ({g.instrument_name.source})"
                if g.instrument_name
                else "",
                w.scoring_text(project, g),
                g.n_complete_rows,
                g.cronbach_alpha,
                g.cronbach_alpha_reversed,
            ]
        )
    return rows


GROUP_HEADER = TextList(
    [
        N_("Groupe"),
        N_("Description"),
        N_("Variables"),
        N_("Indices du regroupement"),
        N_("Items inversés probables"),
        N_("Items inversés confirmés"),
        N_("Hypothèse de l'IA"),
        N_("Instrument (source)"),
        N_("Cotation (source)"),
        N_("Lignes complètes"),
        N_("Alpha de Cronbach observé"),
        N_("Alpha, items probables inversés"),
    ]
)


def _provenance(project: Project, kind: w.ExportKind) -> list[list]:
    rows = []

    def as_text(value: object) -> str:
        if value is None:
            return ""
        return ", ".join(value) if isinstance(value, list) else str(value)

    def add(name: str, field: str, tracked: TrackedValue, text: str | None = None) -> None:
        """`text`: the value as this codebook shows it, when it differs from the stored one."""
        value = w.shown(tracked, kind)
        if value is None:
            return
        rows.append(
            [
                name,
                field,
                text if text is not None else as_text(value.value),
                w.ORIGINS[tracked.origin],
                w.status_text(tracked),
                as_text(tracked.proposed_value),
                tracked.ai_run,
            ]
        )

    for e in aside.kept(project):
        add(e.name, _("Nouveau nom"), e.new_name)
        add(e.name, _("Question"), e.question)
        if e.value_labels.value is not None:
            value = w.shown(e.value_labels, kind)
            if value is not None:
                rows.append(
                    [
                        e.name,
                        _("Étiquettes de valeurs"),
                        "; ".join(f"{w.code_text(v.code)} = {v.label}" for v in value.value),
                        w.ORIGINS[e.value_labels.origin],
                        w.STATUSES[e.value_labels.status],
                        "",
                        None,
                    ]
                )
        for decision in e.sentinel_decisions:
            rows.append(
                [
                    e.name,
                    _("Code suspect {code}").format(code=w.code_text(decision.code)),
                    _("code manquant") if decision.missing else _("vraie valeur"),
                    w.ORIGINS["human"],
                    "",
                    "",
                    None,
                ]
            )
        add(e.name, _("Libellé"), e.label)
        add(e.name, _("Description"), e.description)
        add(e.name, _("Unité"), e.unit)
        add(e.name, _("Groupe"), e.group_id)
        add(e.name, _("Confidentialité"), e.privacy.privacy_class)
        for issue in e.issues:
            add(e.name, _("Signalement"), issue)
    for g in project.groups:
        add(g.group_id, _("Variables"), g.members)
        add(g.group_id, _("Description"), g.descriptive_label, w.group_description(g, kind))
        add(g.group_id, _("Hypothèse de l'IA"), g.ai_hypothesis)
    return rows


@in_language
def write_xlsx(
    project: Project,
    path: Path | str | BinaryIO,
    *,
    kind: w.ExportKind = "final",
    omit_groups: tuple[str, ...] = (),
) -> None:
    """`omit_groups`: groups whose question text is left out of this export (D-014)."""
    omitted = w.omitted_variables(project, omit_groups)
    book = Workbook()
    book.remove(book.active)
    _sheet(book, _("Informations"), [_("Élément"), _("Valeur")], _about(project, kind, omit_groups))
    header = list(VARIABLE_HEADER)
    if w.has_presentation(project, kind):  # v0.21: after "Texte de la question"
        at = header.index(_("Texte de la question")) + 1
        header[at:at] = [_("Posée si"), _("Présentation")]
    if w.renamed(project, kind):
        header.insert(2, _("Nom d'origine"))
    _sheet(book, _("Variables"), header, _variables(project, kind, omitted))
    _sheet(
        book,
        _("Valeurs"),
        [
            _("Variable"),
            _("Code"),
            _("Étiquette"),
            _("Effectif"),
            _("Manquant déclaré"),
            _("Source de l'étiquette"),
        ],
        _values(project, kind),
    )
    _sheet(book, _("Groupes"), GROUP_HEADER, _groups(project, kind))
    if project.recodes.rules:  # v0.15 (D-120): the derived variables
        _sheet(
            book,
            _("Recodages"),
            list(w.RECODE_HEADER),
            [list(row) for row in w.recode_rows(project)],
        )
    if project.glossary:
        _sheet(
            book,
            _("Glossaire"),
            [_("Sigle"), _("Définition"), _("Origine"), _("Statut"), _("Source")],
            [list(row) for row in w.glossary_rows(project, kind)],
        )
    if aside.aside(project):  # D-158: what the codebook leaves out, and why
        _sheet(
            book,
            _("Variables écartées"),
            [_("Variable"), _("Raison"), _("Date")],
            [[e.name, e.set_aside.reason, e.set_aside.date] for e in aside.aside(project)],
        )
    if w.renamed(project, kind):
        _sheet(
            book,
            _("Correspondance"),
            [_("Nom d'origine"), _("Nouveau nom"), _("Statut"), _("Origine")],
            [list(row) for row in w.correspondence_rows(project, kind)],
        )
    _sheet(
        book,
        _("Provenance"),
        [
            _("Variable"),
            _("Champ"),
            _("Valeur"),
            _("Origine"),
            _("Statut"),
            _("Valeur d'origine"),
            _("Appel IA"),
        ],
        _provenance(project, kind),
    )
    book.save(path if hasattr(path, "write") else Path(path))  # a file, or bytes in memory
