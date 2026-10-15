"""Read an Excel workbook (.xlsx) as a team types it: names in the first row of the first
sheet, one row per participant (v0.2).

The export of a survey platform has two header rows (v0.7): the question once above its
first column, its items or options below. It is recognized (ddict.platforms), its columns
get the names of the platform's .sav file, and the second row is never read as data.

Each cell is turned into the text a CSV export would hold ("170", "2026-01-05 08:00"), and
the columns are then typed by the CSV rules: a workbook and its CSV export give the same
profile. Formulas are read through their last computed value. An Excel file has no
variable labels nor value labels.
"""

from __future__ import annotations

import datetime as dt
from pathlib import Path
from zipfile import BadZipFile

import pandas as pd
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from ddict.io.csv_reader import _blank_tokens, _to_column
from ddict.io.model import LoadedData, ReadError, SourceFile, VariableMetadata
from ddict.io.sav import file_sha256
from ddict.platforms import split_header


def _texts(values: list[object]) -> list[str]:
    """The cells of a column as CSV text. Dates without time are written as dates when the
    whole column has no time of day, as a CSV export would show them."""
    stamps = [v for v in values if isinstance(v, dt.datetime)]
    dates_only = all(v.time() == dt.time(0) for v in stamps)
    texts = []
    for value in values:
        if value is None:
            texts.append("")
        elif isinstance(value, bool):
            texts.append(str(int(value)))
        elif isinstance(value, dt.datetime):
            if dates_only:
                texts.append(value.date().isoformat())
            else:
                texts.append(
                    value.strftime("%Y-%m-%d %H:%M" if not value.second else "%Y-%m-%d %H:%M:%S")
                )
        elif isinstance(value, dt.date):
            texts.append(value.isoformat())
        elif isinstance(value, float) and value.is_integer():
            texts.append(str(int(value)))
        else:
            texts.append(str(value))
    return texts


def read_xlsx(path: Path | str) -> LoadedData:
    path = Path(path)
    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
    except (InvalidFileException, BadZipFile, OSError, KeyError, ValueError) as error:
        raise ReadError(f"Cannot read {path.name} as an Excel file: {error}") from error
    try:
        rows = [list(row) for row in workbook.worksheets[0].iter_rows(values_only=True)]
    finally:
        workbook.close()
    rows = [row for row in rows if any(cell is not None and cell != "" for cell in row)]
    if not rows:
        raise ReadError(f"{path.name}: the first sheet is empty")

    split = split_header(rows)
    if split is not None:
        platform, names = split
        width = len(names)
        records = [[*row[:width], *[None] * (width - len(row))] for row in rows[2:]]
        return _loaded(path, names, records, platform.name)

    width = max(len(row) for row in rows)
    header = [*rows[0], *[None] * (width - len(rows[0]))]
    while header and header[-1] in (None, ""):  # columns without a name nor any value
        if any(
            len(row) >= len(header) and row[len(header) - 1] not in (None, "") for row in rows[1:]
        ):
            break
        header.pop()
    width = len(header)
    records = [[*row[:width], *[None] * (width - len(row))] for row in rows[1:]]
    names = [
        str(name).strip() if name not in (None, "") else f"colonne_{position}"
        for position, name in enumerate(header, start=1)
    ]
    return _loaded(path, names, records, None)


def _loaded(
    path: Path, names: list[str], records: list[list[object]], platform: str | None
) -> LoadedData:
    width = len(names)
    duplicates = sorted({name for name in names if names.count(name) > 1})
    if duplicates:
        raise ReadError(f"{path.name}: duplicate column names: {', '.join(duplicates)}")

    blanked = [_blank_tokens(_texts([record[i] for record in records])) for i in range(width)]
    table = pd.DataFrame(
        {
            name: _to_column(values, ".")
            for name, (values, _counts) in zip(names, blanked, strict=True)
        }
    )
    variables = tuple(
        VariableMetadata(
            name=name,
            position=position,
            storage_type="numeric" if table[name].dtype == "float64" else "string",
            missing_tokens=counts,
        )
        for position, (name, (_values, counts)) in enumerate(zip(names, blanked, strict=True), 1)
    )
    source = SourceFile(
        name=path.name,
        format="xlsx",
        sha256=file_sha256(path),
        n_rows=len(table),
        n_columns=width,
        platform=platform,
    )
    return LoadedData(source=source, variables=variables, table=table)
