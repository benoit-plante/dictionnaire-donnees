"""The designed questionnaire in Word (v0.20, step E1, D-152), written by computation, never
by the AI, in two versions.

- For the ethics committee: what a person answering would read (title, sections and their
  instructions, numbered questions, choices with ○ or ☐, grids as tables, lines to fill,
  displayed texts, "*" for a required question), in the language of the study; since v0.21,
  the estimated duration, the conditions in words ("Si vous avez répondu « Oui » à la
  question 3 :") and the random orders. No name, no code, no provenance; an attention check
  shows as an ordinary question, as a person answering sees it.
- Working: the same, with under each question its code, its kind and its LimeSurvey type, the
  codes of its choices and items, the reversed items, the provenance, the group and its
  instrument, its variables, its condition for the team and its attention checks (v0.21); the
  formula of the duration; at the end, the review points open and kept. The tool's own words
  in the language of the interface.

The footer gives the version, the day the questionnaire last changed and its fingerprint,
never the day of its production: the same questionnaire gives the same bytes.
"""

from __future__ import annotations

import contextlib
import io
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Literal

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

from ddict import design, design_duration, design_points, i18n
from ddict.analysis.deposit import steady_xlsx
from ddict.i18n import N_, Texts, _
from ddict.project import skip
from ddict.project.model import Group, Project, Provenance, Question, Section

Version = Literal["ethics", "working"]
VERSIONS: tuple[Version, ...] = ("ethics", "working")
VERSION_NAMES = Texts({"ethics": N_("pour le comité d'éthique"), "working": N_("de travail")})
KINDS = Texts(
    {
        "open": N_("ouverte"),
        "single": N_("choix unique"),
        "checkbox": N_("à cocher"),
        "grid": N_("grille"),
        "display": N_("texte d'affichage"),
    }
)
ANSWERS = Texts(
    {
        "short": N_("texte court"),
        "long": N_("texte long"),
        "number": N_("nombre"),
        "date": N_("date"),
    }
)
PROVENANCES = Texts({"team": N_("tapé par l'équipe"), "source": N_("tiré d'une source")})
ROUND, SQUARE = "○", "☐"  # one choice, several
# v0.21 (D-166): what a person answering is told of a random order, in the language of the
# study; the page "Publier" says the same.
SHUFFLED = Texts(
    {
        "single": N_("(Les choix sont présentés dans un ordre aléatoire.)"),
        "checkbox": N_("(Les options sont présentées dans un ordre aléatoire.)"),
        "grid": N_("(Les énoncés de cette question sont présentés dans un ordre aléatoire.)"),
        "section": N_("(Les questions de cette section sont présentées dans un ordre aléatoire.)"),
    }
)
LINE = "_" * 40
FONT = "Calibri"
MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


class WordOpen(Exception):
    """The file is open in Word, which holds it: nothing was written."""

    def __init__(self, path: Path) -> None:
        super().__init__(path.name)
        self.path = path


@dataclass(frozen=True)
class Numbered:
    """A question as a person answering sees it: its number, or none for a displayed text."""

    number: int | None
    question: Question


def provenance_text(provenance: Provenance | None) -> str:
    if provenance is None:
        return ""
    if provenance.kind == "source":
        page = _(", p. {n}").format(n=provenance.page) if provenance.page else ""
        return _("source : {source}{page}").format(source=provenance.source, page=page)
    return _(PROVENANCES.get(provenance.kind, provenance.kind))


def numbered(project: Project) -> list[tuple[Section, list[Numbered]]]:
    """The sections in order, each with its questions numbered through the questionnaire; a
    displayed text, which asks nothing, has no number."""
    found = design.questionnaire(project)
    number = 0
    out = []
    for section in found.sections:
        questions = []
        for question in (q for q in found.questions if q.section == section.section_id):
            if question.kind != "display":
                number += 1
            questions.append(Numbered(number if question.kind != "display" else None, question))
        out.append((section, questions))
    return out


def condition_line(project: Project, question: Question) -> str | None:
    """A question's own condition, as a person answering reads it, before the question:
    "Si vous avez répondu « Oui » à la question 3 :" (its section's is said under its title)."""
    if question.condition is None:
        return None
    return skip.committee_text(project, question.condition) + i18n.colon().rstrip()


def section_lines(project: Project, section: Section) -> list[str]:
    """Under a section's title: whom it is for, and the random order of its questions."""
    found = []
    if section.condition is not None:
        found.append(skip.committee_text(project, section.condition, section=True) + ".")
    if section.shuffle:
        found.append(_(SHUFFLED["section"]))
    return found


def shuffle_line(question: Question) -> str | None:
    return _(SHUFFLED[question.kind]) if question.shuffle and question.kind in SHUFFLED else None


def attention_lines(question: Question) -> list[str]:
    """The attention checks of a question, for the team: "Item de contrôle : bonne réponse 4
    (Plutôt d'accord)"."""
    found = []
    if question.attention:
        found.append(
            _("Item de contrôle : {answers}").format(
                answers=skip.attention_text(question, question.attention)
            )
        )
    for row in question.rows:
        if row.attention:
            found.append(
                _("Item de contrôle « {row} » : {answers}").format(
                    row=row.text, answers=skip.attention_text(question, row.attention)
                )
            )
    return found


def duration_line(project: Project) -> str | None:
    if not design.questionnaire(project).questions:
        return None
    found = design_duration.estimate(project)
    return _("Durée estimée : {duration}.").format(duration=design_duration.text(found))


def version(project: Project) -> tuple[int, str]:
    """The questionnaire's version and the day it last changed (D-152). Each content sent
    (D-150) is a version, in the order sent; what changed since is the next one."""
    found = design.questionnaire(project)
    sent: list[str] = []
    for publication in project.design.publications:
        if publication.fingerprint not in sent:
            sent.append(publication.fingerprint)
    number = sent.index(found.sha256) + 1 if found.sha256 in sent else len(sent) + 1
    return number, project.design.changed or project.design.created


def footer_text(project: Project) -> str:
    number, changed = version(project)
    return _("Questionnaire « {title} », version {n} du {date} — empreinte {fingerprint}").format(
        title=design.questionnaire(project).file_name,
        n=number,
        date=changed,
        fingerprint=design.questionnaire(project).sha256[:12],
    )


def stem(project: Project) -> str:
    """The questionnaire's name as a file name: a character Windows refuses becomes "-"."""
    found = design.questionnaire(project).file_name
    return "".join("-" if c in '<>:"/\\|?*' else c for c in found).strip(" .") or "questionnaire"


def file_name(project: Project, which: Version) -> str:
    """ "Étude du sommeil - comité.docx", "Étude du sommeil - travail.docx"."""
    words = {"ethics": _("comité"), "working": _("travail")}
    return f"{stem(project)} - {words[which]}.docx"


# --- Writing. ---


def _grey(paragraph, text: str) -> None:
    run = paragraph.add_run(text)
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(0x60, 0x60, 0x60)


def _plain(document, text: str, *, italic: bool = False, indent: bool = False):
    paragraph = document.add_paragraph()
    run = paragraph.add_run(text)
    run.italic = italic
    if indent:
        paragraph.paragraph_format.left_indent = Pt(18)
    return paragraph


def _group(project: Project, question: Question) -> Group | None:
    return next((g for g in project.groups if g.question == question.question_id), None)


def _reversed(project: Project, question: Question) -> set[str]:
    """The ids of the items of a grid that its group says are reversed."""
    group = _group(project, question)
    if group is None:
        return set()
    names = set(group.reverse_items)
    return {
        e.question.value
        for e in project.variables
        if e.name in names and e.question.value is not None
    }


def _variables(project: Project, question: Question) -> list[str]:
    return [
        e.name
        for e in project.variables
        if (e.question.value or "").partition(".")[0] == question.question_id
    ]


def _coded(code: str | None, text: str, working: bool) -> str:
    return f"[{code}] {text}" if working and code else text


def _choices(document, question: Question, working: bool) -> None:
    mark = SQUARE if question.kind == "checkbox" else ROUND
    for option in question.options:
        if option.other:  # LimeSurvey keeps what is typed, under no code
            text = option.text + " : " + "_" * 20
        else:
            text = _coded(option.code, option.text, working)
        _plain(document, f"{mark} {text}", indent=True)


def _grid(document, project: Project, question: Question, working: bool) -> None:
    reversed_ids = _reversed(project, question) if working else set()
    table = document.add_table(rows=1 + len(question.rows), cols=1 + len(question.options))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for n, option in enumerate(question.options, start=1):
        cell = table.rows[0].cells[n]
        cell.text = _coded(option.code, option.text, working)
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    for r, row in enumerate(question.rows, start=1):
        text = _coded(row.code, row.text, working)
        if row.row_id in reversed_ids:
            text += " " + _("(inversé)")
        table.rows[r].cells[0].text = text
        for n in range(1, len(question.options) + 1):
            cell = table.rows[r].cells[n]
            cell.text = ROUND
            cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER


def _open(document, question: Question) -> None:
    if question.answer == "long":
        for _line in range(3):
            _plain(document, LINE + "_" * 30, indent=True)
        return
    hint = {"number": _("(nombre)"), "date": _("(date : aaaa-mm-jj)")}.get(question.answer)
    _plain(document, LINE[:20] + (f" {hint}" if hint else ""), indent=True)


def _working_lines(project: Project, question: Question, tool: str) -> list[str]:
    """What the working version says under a question, in the language of the interface."""
    with i18n.using(tool):
        kind = _(KINDS.get(question.kind, question.kind))
        if question.kind == "open" and question.answer:
            kind += f" ({_(ANSWERS.get(question.answer, question.answer))})"
        if question.kind == "single" and question.layout == "dropdown":
            kind += " " + _("(liste déroulante)")
        parts = [
            _("Code : {code}").format(code=design.code(project, question)),
            kind,
            _("type LimeSurvey {type}").format(type=design.limesurvey_type(question)),
        ]
        if question.required:
            parts.append(_("obligatoire"))
        lines = [f"{question.question_id} — " + " — ".join(parts)]
        names = _variables(project, question)
        if names and question.kind == "display":  # a column LimeSurvey exports empty
            lines.append(_("Variable : {name} (toujours vide)").format(name=names[0]))
        elif names:
            lines.append(_("Variables : {names}").format(names=", ".join(names)))
        provenance = provenance_text(question.provenance)
        if provenance:
            lines.append(_("Provenance : {text}").format(text=provenance))
        if question.condition is not None:  # v0.21
            lines.append(skip.team_text(project, question.condition))
        lines += attention_lines(question)
        group = _group(project, question)
        if group is not None:
            label = group.descriptive_label.value or group.group_id
            line = _("Groupe : {label}").format(label=label)
            if group.instrument_name is not None:
                line += " — " + _("instrument : {name} ({source})").format(
                    name=group.instrument_name.value, source=group.instrument_name.source
                )
            lines.append(line)
    return lines


def _question(document, project: Project, item: Numbered, working: bool, tool: str) -> None:
    question = item.question
    lead = condition_line(project, question)
    if lead is not None:
        paragraph = _plain(document, lead, italic=True)
        paragraph.paragraph_format.space_before = Pt(10)
        paragraph.paragraph_format.keep_with_next = True
    if question.kind == "display":
        _plain(document, question.text, italic=True)
    else:
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.space_before = Pt(10)
        paragraph.paragraph_format.keep_with_next = True
        paragraph.add_run(f"{item.number}. ").bold = True
        paragraph.add_run(question.text).bold = True
        if question.required:
            paragraph.add_run(" *").bold = True
    if question.instruction:
        _plain(document, question.instruction, italic=True)
    shuffled = shuffle_line(question)
    if shuffled is not None:
        _plain(document, shuffled, italic=True)
    if question.kind in ("single", "checkbox"):
        _choices(document, question, working)
    elif question.kind == "grid":
        _grid(document, project, question, working)
    elif question.kind == "open":
        _open(document, question)
    if working:
        for line in _working_lines(project, question, tool):
            _grey(document.add_paragraph(), line)


def _points(document, project: Project, tool: str) -> None:
    with i18n.using(tool):
        found = design_points.points(project, kept=True)
        document.add_heading(_("Points de relecture"), level=1)
        if not found:
            _plain(document, _("Aucun point de relecture."))
            return
        number = {n.question.question_id: n.number for _s, q in numbered(project) for n in q}
        for point in found:
            where = number.get(point.question)
            label = _("question {n}").format(n=where) if where else point.question
            kept = " " + _("(gardé tel quel)") if design_points.is_kept(project, point) else ""
            _plain(document, f"• {label} : {point.text}{kept}")


def render(project: Project, which: Version) -> bytes:
    """The Word file of the designed questionnaire, in one of its two versions."""
    tool = i18n.language()
    working = which == "working"
    found = design.questionnaire(project)
    _number, changed = version(project)
    with i18n.using(project.settings.language):
        document = Document()
        style = document.styles["Normal"]
        style.font.name = FONT
        style.font.size = Pt(11)
        document.add_heading(found.file_name, level=0)
        if working:
            with i18n.using(tool):
                _grey(
                    document.add_paragraph(),
                    _(
                        "Version de travail : avec les codes, la provenance et les points de "
                        "relecture. Pour le comité d'éthique, produisez l'autre version."
                    ),
                )
        duration = duration_line(project)
        if duration is not None:
            _plain(document, duration)
            if working:
                with i18n.using(tool):
                    for line in design_duration.formula(design_duration.estimate(project)):
                        _grey(document.add_paragraph(), line)
        if any(q.required for q in found.questions):
            _plain(document, _("* Réponse obligatoire."), italic=True)
        for section, questions in numbered(project):
            if section.title:
                document.add_heading(section.title, level=1)
            if section.instruction:
                _plain(document, section.instruction, italic=True)
            for line in section_lines(project, section):
                _plain(document, line, italic=True)
            if working and section.condition is not None:
                with i18n.using(tool):
                    _grey(document.add_paragraph(), skip.team_text(project, section.condition))
            for item in questions:
                _question(document, project, item, working, tool)
        if working:
            _points(document, project, tool)
        footer = document.sections[0].footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _grey(footer, footer_text(project))
        properties = document.core_properties
        stamp = datetime.fromisoformat(changed)
        properties.author = properties.last_modified_by = "dictionnaire-donnees"
        properties.title = found.file_name
        properties.created = properties.modified = stamp
        properties.revision = 1
        out = io.BytesIO()
        document.save(out)
    # The zip entries carry the time of the saving: fixed to the day of the last change.
    return steady_xlsx(out.getvalue(), changed)


def write(project: Project, which: Version, folder: Path) -> Path:
    """Write one version next to the project, through a temporary file: a file open in Word
    is refused, and never left half written."""
    path = folder / file_name(project, which)
    content = render(project, which)
    temporary = path.with_name(path.name + ".tmp")
    try:
        temporary.write_bytes(content)
        temporary.replace(path)
    except PermissionError as error:
        with contextlib.suppress(OSError):
            temporary.unlink()
        raise WordOpen(path) from error
    return path


def read_text(content: bytes) -> str:
    """Every paragraph, every cell and the footer, in order: for the tests."""
    document = Document(io.BytesIO(content))
    parts = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts += [c.text for c in row.cells]
    parts += [p.text for p in document.sections[0].footer.paragraphs]
    return "\n".join(parts)
