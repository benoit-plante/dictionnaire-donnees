"""The header of a survey platform's Excel export as the questionnaire (v0.7, D-045).

Its two header rows hold the text of each question once, above its first column, and below
each column its grid item or checkbox option: the questionnaire without the uncertain cutting
of a PDF. Only these two rows are read, never the answers below them.

Question identifiers follow the platform's ranks: `q0003_0004` is `Q3.r4`. The header does
not say whether a question with several columns is a grid or a set of checkboxes; the
structure of the project's database does (`kinds`), else the wording of the question, else
an "other" column (a set of checkboxes). It does not hold the answer options of a choice
question either: those stay the value labels of the .sav file.
"""

from __future__ import annotations

from pathlib import Path
from zipfile import BadZipFile

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from ddict.io.model import ReadError
from ddict.io.sav import file_sha256
from ddict.platforms import split_header
from ddict.project.model import GridRow, Question, Questionnaire, QuestionOption
from ddict.questionnaire.segment import _CHECKBOX

OPEN_CELLS = ("Open-Ended Response",)


def read_header_questionnaire(
    path: Path | str, kinds: dict[int, str] | None = None
) -> Questionnaire:
    """The questions of an export's header. `kinds`: "grid" or "checkbox" by question rank,
    from the database of the project when it is the same survey."""
    path = Path(path)
    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
    except (InvalidFileException, BadZipFile, OSError, KeyError, ValueError) as error:
        raise ReadError(f"Cannot read {path.name} as an Excel file: {error}") from error
    try:
        rows = [
            list(row)
            for row in workbook.worksheets[0].iter_rows(min_row=1, max_row=2, values_only=True)
        ]
    finally:
        workbook.close()
    split = split_header(rows)
    if split is None:
        raise ReadError(f"{path.name}: no survey platform header in its first two rows")
    platform, names = split
    first, second = (
        [("" if c is None else str(c).strip()) for c in row] + [""] * len(names) for row in rows
    )

    blocks: dict[int, dict] = {}
    for column, name in enumerate(names):
        place = platform.place(name)
        if place is None:
            continue
        block = blocks.setdefault(place.question, {"text": first[column], "cells": []})
        block["cells"].append((place, second[column]))

    questions = []
    for rank, block in sorted(blocks.items()):
        question_id = f"Q{rank}"
        items = [(p, text) for p, text in block["cells"] if p.item is not None]
        others = [text for p, text in block["cells"] if p.other]
        whole = [text for p, text in block["cells"] if p.item is None and not p.other]
        if not items:
            kind = "open" if whole and whole[0] in OPEN_CELLS else "single"
            # Its choices are not in the header; its "other" option is, by its own column.
            options = [
                QuestionOption(option_id=f"{question_id}.o{k}", text=text)
                for k, text in enumerate(others, start=1)
            ]
            questions.append(
                Question(question_id=question_id, text=block["text"], kind=kind, options=options)
            )
            continue
        kind = (kinds or {}).get(rank) or _guess(block["text"], others)
        if kind == "grid":
            rows_ = [
                GridRow(row_id=f"{question_id}.r{p.item}", number=None, text=text)
                for p, text in items
            ]
            questions.append(
                Question(question_id=question_id, text=block["text"], kind="grid", rows=rows_)
            )
            continue
        options = [
            QuestionOption(option_id=f"{question_id}.o{p.item}", text=text) for p, text in items
        ]
        options += [
            QuestionOption(option_id=f"{question_id}.o{len(items) + k}", text=text)
            for k, text in enumerate(others, start=1)
        ]
        questions.append(
            Question(question_id=question_id, text=block["text"], kind="checkbox", options=options)
        )
    return Questionnaire(
        file_name=path.name,
        sha256=file_sha256(path),
        format="xlsx",
        n_pages=None,
        questions=questions,
        texts=[],
    )


def _guess(text: str, others: list[str]) -> str:
    return "checkbox" if others or _CHECKBOX.search(text) else "grid"
