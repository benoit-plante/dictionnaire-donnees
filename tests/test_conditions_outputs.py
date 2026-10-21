"""v0.21, step 5 (§ 5.6 of the plan): the conditions, the random orders, the attention checks
and the estimated duration in the two Words, the preview of the page "Publier" and the planned
codebook (Excel, Markdown, analysis context). Nothing real: the questionnaire of `test_skip`."""

import io
from pathlib import Path

import pytest
from openpyxl import load_workbook
from streamlit.testing.v1 import AppTest

from ddict import design, design_word, i18n
from ddict.export import render_markdown, write_xlsx
from ddict.export.context import render_context
from ddict.project import save_project
from projet_concu import project as designed
from test_skip import famille, term, when

APP = str(Path(__file__).resolve().parent.parent / "src" / "ddict" / "review" / "app.py")

# Each line of the table of § 5.1, on the question that carries it.
CONDITIONS = {
    "Q4": when(term("Q3", "eq", "1")),
    "Q6": when(term("Q5", "in", ["1", "2"])),
    "Q8": when(term("Q7.r2", "eq", "4")),
    "Q9": when(term("Q2", "ge", 18)),
    "Q10": when(term("Q4", "answered")),
}
COMMITTEE_FR = [
    "Si vous avez répondu « Oui » à la question 3 :",
    "Si vous avez répondu « Temps plein » ou « Temps partiel » à la question 5 :",
    "Si vous avez répondu « Plutôt d'accord » à l'énoncé « est sûr. » de la question 7 :",
    "Si votre réponse à la question 2 est 18 ou plus :",
    "Si vous avez répondu à la question 4 :",
]
COMMITTEE_EN = [
    "If you answered “Oui” to question 3:",
    "If you answered “Temps plein” or “Temps partiel” to question 5:",
    "If you answered “Plutôt d'accord” to the statement “est sûr.” of question 7:",
    "If your answer to question 2 is 18 or more:",
    "If you gave an answer to question 4:",
]
TEAM_FR = [
    "Posée si Q3 (enfants) = 1 (Oui)",
    "Posée si Q5 (emploi) ∈ {1 (Temps plein), 2 (Temps partiel)}",
    "Posée si Q7 (climat_SQ002) = 4 (Plutôt d'accord)",
    "Posée si Q2 (age) ≥ 18",
    "Posée si Q4 (nbenfants) a une réponse",
]


def filtered(language="fr"):
    project = famille()
    project.settings.language = language
    for owner, condition in CONDITIONS.items():
        design.set_condition(project, owner, condition)
    design.set_condition(
        project, "S2", when(term("Q10.o2", "checked"), term("Q1", "eq", "2"), combine="any")
    )
    design.set_shuffle(project, "S2", True)
    design.set_shuffle(project, "Q7", True)
    design.set_shuffle(project, "Q5", True)
    design.set_attention(project, "Q7.r2", ["4"])
    return project


def words(project, which, tool="fr"):
    with i18n.using(tool):
        content = design_word.render(project, which)
        assert content == design_word.render(project, which)  # the same bytes twice
    return content, design_word.read_text(content)


# --- The Word (D-152). ---


def test_the_word_for_the_committee():
    _content, text = words(filtered(), "ethics")
    lines = text.split("\n")
    for line in COMMITTEE_FR:
        assert line in lines
    assert (
        "Cette section s'adresse aux personnes qui ont coché « Vélo » à la question 10 ou qui "
        "ont répondu « Homme » à la question 1." in lines
    )
    assert "(Les questions de cette section sont présentées dans un ordre aléatoire.)" in lines
    assert "(Les énoncés de cette question sont présentés dans un ordre aléatoire.)" in lines
    assert "(Les choix sont présentés dans un ordre aléatoire.)" in lines
    assert any(line.startswith("Durée estimée : environ ") for line in lines)
    # No code, no name, no team's words, and the attention check as an ordinary statement.
    for word in ("Posée si", "(enfants)", "climat_SQ002", "Item de contrôle", "IReST", "Q3"):
        assert word not in text
    assert "est sûr." in lines


def test_the_working_word():
    _content, text = words(filtered(), "working")
    lines = text.split("\n")
    for line in COMMITTEE_FR + TEAM_FR:
        assert line in lines
    assert "Posée si Q10 (transport) : Vélo coché ou Q1 (genre) = 2 (Homme)" in lines
    assert "Item de contrôle « est sûr. » : bonne réponse 4 (Plutôt d'accord)" in lines
    assert any("195 mots par minute" in line for line in lines)
    assert any(line.startswith("Chemin le plus court") for line in lines)


def test_the_word_in_english():
    _content, text = words(filtered("en"), "ethics", tool="en")
    lines = text.split("\n")
    for line in COMMITTEE_EN:
        assert line in lines
    assert "(The statements of this question are presented in random order.)" in lines
    assert any(line.startswith("Estimated duration: about ") for line in lines)
    _content, text = words(filtered("en"), "working", tool="en")
    assert "Asked if Q3 (enfants) = 1 (Oui)" in text.split("\n")


def test_the_word_of_a_questionnaire_without_conditions_only_gains_its_duration():
    _content, text = words(designed(), "ethics")
    assert "Durée estimée : environ une minute." in text.split("\n")
    assert "Si vous" not in text and "aléatoire" not in text


# --- The preview of the page "Publier". ---


@pytest.mark.allow_hosts(["127.0.0.1"])
def test_the_preview_says_the_same(tmp_path):
    path = tmp_path / "Familles" / "familles.dd.json"
    path.parent.mkdir()
    save_project(filtered(), path)
    app = AppTest.from_file(APP, default_timeout=60)
    app.session_state["project_path"] = str(path)
    app.session_state["mode"] = "publish"
    app.run()
    assert not app.exception, app.exception
    shown = [m.value for m in app.markdown]
    assert "*Si vous avez répondu « Oui » à la question 3 :*" in shown
    assert any(m.startswith("Durée estimée : environ") for m in shown)
    assert any("Cette section s'adresse aux personnes" in m for m in shown)
    captions = [x.value for x in app.caption]
    assert "(Les énoncés de cette question sont présentés dans un ordre aléatoire.)" in captions
    assert "Posée si Q3 (enfants) = 1 (Oui)" not in captions
    app.toggle(key="publish-names").set_value(True).run()
    captions = [x.value for x in app.caption]
    assert "Posée si Q3 (enfants) = 1 (Oui)" in captions
    assert "Item de contrôle « est sûr. » : bonne réponse 4 (Plutôt d'accord)" in captions


# --- The planned codebook. ---


def _sheet(project):
    out = io.BytesIO()
    write_xlsx(project, out)
    book = load_workbook(io.BytesIO(out.getvalue()))
    rows = list(book["Variables"].iter_rows(values_only=True))
    header = list(rows[0])
    return header, {
        row[header.index("Nom")]: dict(zip(header, row, strict=True)) for row in rows[1:]
    }


def test_the_codebook_says_when_a_variable_is_asked_and_how():
    header, rows = _sheet(filtered())
    at = header.index("Texte de la question")
    assert header[at + 1 : at + 3] == ["Posée si", "Présentation"]
    assert rows["nbenfants"]["Posée si"] == "enfants = 1 (Oui)"
    assert rows["genre"]["Posée si"] in (None, "")
    assert rows["climat_SQ002"]["Posée si"] == "transport_2 coché ou genre = 2 (Homme)"
    assert rows["transport_2"]["Posée si"] == (  # its section's, then its own
        "(transport_2 coché ou genre = 2 (Homme)) et nbenfants a une réponse"
    )
    assert rows["climat_SQ002"]["Présentation"] == (
        "Ordre de présentation : aléatoire ; Item de contrôle (bonne réponse 4 (Plutôt d'accord))"
    )
    assert rows["emploi"]["Présentation"] == "Ordre de présentation : aléatoire"
    markdown = render_markdown(filtered())
    assert "- Posée si : enfants = 1 (Oui)" in markdown
    context = render_context(filtered())
    assert "## Questions filtrées et présentation" in context
    assert "- `nbenfants` : posée si enfants = 1 (Oui)" in context


def test_a_codebook_without_conditions_keeps_its_columns():
    header, _rows = _sheet(designed())
    assert "Posée si" not in header and "Présentation" not in header
    assert "Posée si" not in render_markdown(designed())
    assert "Questions filtrées" not in render_context(designed())
