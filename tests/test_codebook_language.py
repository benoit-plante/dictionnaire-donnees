"""v0.5, step 3: the language of the codebook, chosen at export (D-032).

The two reports of the fictional base are kept in tests/reference; to write them again after
a deliberate change: DDICT_REGENERATE=1 uv run pytest tests/test_codebook_language.py
"""

import os
import shutil
from pathlib import Path

import pytest
from openpyxl import load_workbook
from streamlit.testing.v1 import AppTest

from ddict import i18n
from ddict.cli import main
from ddict.export import render_markdown, write_xlsx
from ddict.export.syntax import render_syntax
from ddict.project import load_project, save_project
from ddict.review import logic
from ddict.workflow import add_questionnaire, profile_database

NOW = "2026-10-14T18:00:00"
REFERENCE = Path(__file__).resolve().parent / "reference"
APP = str(Path(__file__).resolve().parent.parent / "src" / "ddict" / "review" / "app.py")


@pytest.fixture
def project_path(fictional_files, tmp_path):
    for name in ("etude-sommeil.sav", "questionnaire-sommeil.txt"):
        shutil.copy(fictional_files[name], tmp_path / name)
    project, path = profile_database(tmp_path / "etude-sommeil.sav")
    add_questionnaire(project, tmp_path / "questionnaire-sommeil.txt")
    logic.link_question(project, "age", "Q3", now=NOW)
    logic.decide_sentinel(project, "age", 999, True, now=NOW)
    logic.glossary_decide(project, "TCC", "accept", now=NOW)
    save_project(project, path)
    return path


@pytest.mark.parametrize("language", ["fr", "en"])
def test_reference_reports(project_path, language):
    text = render_markdown(load_project(project_path), language=language)
    reference = REFERENCE / f"etude-sommeil.{language}.md"
    if os.environ.get("DDICT_REGENERATE"):
        reference.write_text(text, encoding="utf-8", newline="\n")
    assert text == reference.read_text(encoding="utf-8")
    assert i18n.language() == "fr"  # the interface keeps its language


def test_english_codebook(project_path, tmp_path):
    project = load_project(project_path)
    report = render_markdown(project, language="en")
    assert "# Data dictionary: etude-sommeil.sav" in report
    assert "Code 999 confirmed as missing by the team" in report
    assert "0.363" in report and "0,363" not in report  # decimal point in English
    assert "Genre" in report  # the file's own label is never translated
    path = tmp_path / "c.xlsx"
    write_xlsx(project, path, language="en")
    book = load_workbook(path)
    assert book.sheetnames == [
        "Information",
        "Variables",
        "Values",
        "Groups",
        "Glossary",
        "Provenance",
    ]
    assert [c.value for c in book["Variables"][1]][:4] == [
        "Position",
        "Name",
        "Label",
        "Description",
    ]
    assert "0,363" in render_markdown(project, language="fr")
    assert i18n.language() == "fr"


def test_command_files_in_english(project_path):
    project = load_project(project_path)
    spss = render_syntax(project, "spss", "2026-10-14", language="en")
    assert "* Only the accepted decisions:" in spss
    assert "MISSING VALUES age (999)." in spss  # the commands themselves do not change
    r = render_syntax(project, "r", "2026-10-14", language="en")
    assert "# Before: load your data into an object named" in r


def test_command(project_path):
    assert main(["export", str(project_path), "--language", "en", "--syntax"]) == 0
    book = load_workbook(project_path.with_name("etude-sommeil.xlsx"))
    assert book.sheetnames[0] == "Information"
    text = project_path.with_name("etude-sommeil-metadonnees.sps").read_text(encoding="utf-8")
    assert "Only the accepted decisions" in text
    assert main(["exporter", str(project_path)]) == 0  # the interface's language: French
    assert load_workbook(project_path.with_name("etude-sommeil.xlsx")).sheetnames[0] == (
        "Informations"
    )


@pytest.mark.allow_hosts(["127.0.0.1"])
def test_export_page(project_path):
    app = AppTest.from_file(APP, default_timeout=30)
    app.session_state["project_path"] = str(project_path)
    app.session_state["mode"] = "export"
    app.run()
    assert app.button_group(key="export-language").value == "fr"  # the interface's
    app.button_group(key="export-language").set_value("en").run()
    app.button(key="export-write").click().run()
    assert not app.exception, app.exception
    assert app.header[0].value == "Export"  # the interface stays in French
    assert "Fichiers écrits à côté du projet" in app.success[0].value
    book = load_workbook(project_path.with_name("etude-sommeil.xlsx"))
    assert book.sheetnames[0] == "Information"


def test_group_description_follows_the_language(project_path):
    project = load_project(project_path)
    group = project.groups[0]
    assert group.descriptive_label.value == "8 items, codes de 1 à 5, préfixe « sommeil_q »"
    from ddict.export import wording as w

    assert w.group_description(group, "final") == group.descriptive_label.value
    with i18n.using("en"):
        assert (
            w.group_description(group, "final") == '8 items, codes from 1 to 5, prefix "sommeil_q"'
        )
    group.descriptive_label.value = "Échelle du sommeil, version courte"  # not the tool's
    with i18n.using("en"):
        assert w.group_description(group, "final") == "Échelle du sommeil, version courte"
