"""v0.12, step 8: the Method page, `ddict exclusions`, `ddict carnet echantillon`, the
Method's state in the notebook and "Tout confirmer"; the same files from the page and from
the command line, on the fictional database of administration."""

import io
import shutil
import zipfile

import pytest
from docx import Document
from streamlit.testing.v1 import AppTest

from ddict.analysis import notebook as nb
from ddict.cli import main
from ddict.project import load_project, save_project
from test_exclusions import EXPECTED, KEY, OTHER_ORDER, SAV, _project
from test_method import INSTRUMENT
from test_notebook_page import APP

RULES = KEY["rules"]
ANALYSIS = ["correlation", "--variables", "stress", "anx_total"]


def _folder(tmp_path, name="projet"):
    folder = tmp_path / name
    folder.mkdir()
    shutil.copy(SAV, folder / SAV.name)
    project, _table = _project(rules=False)
    project.groups[0].instrument_name = INSTRUMENT
    path = folder / "etude-passation.dd.json"
    save_project(project, path)
    return path


def _arguments(rule: dict) -> list[str]:
    """A rule of the answer key, as the command line takes it."""
    kind = rule["kind"]
    if kind == "condition":
        found = ["--condition", rule["variable"], rule["operator"], str(rule["value"])]
    elif kind == "completeness":
        found = ["--completude", rule["group"], str(rule["min_answered"])]
    elif kind == "duration":
        found = ["--duree", rule["start"], rule["end"], str(rule["min_seconds"])]
    elif kind == "straightlining":
        found = ["--ligne-droite", rule["group"]]
    else:
        found = ["--doublons", *rule["variables"]]
    found += ["--raison", rule["reason"]]
    return found + (["--source", rule["source"]] if "source" in rule else [])


def _rules_by_command(path):
    for rule in RULES:
        assert main(["exclusions", "ajouter", str(path), *_arguments(rule)]) == 0


def _counts(out: str) -> dict[str, int]:
    found = {}
    for line in out.splitlines():
        if " exclues" in line and line[0].isdigit():
            rule = line.split(". ", 1)[1].split(" — ", 1)[0]
            found[rule] = int(line.rsplit(" — ", 1)[1].split()[0])
    return found


# --- The command line. ---


def test_redirected_output_is_written_in_utf8(tmp_path, monkeypatch):
    """Redirected to a file, Windows gives cp1252, which has no "−" of the results."""
    path = _folder(tmp_path)
    written = io.BytesIO()
    monkeypatch.setattr("sys.stdout", io.TextIOWrapper(written, encoding="cp1252"))
    assert main(["exclusions", "lister", str(path)]) == 0
    import sys

    sys.stdout.write("r(195) = −,07\n")
    sys.stdout.flush()
    assert sys.stdout.encoding == "utf-8"
    out = written.getvalue().decode("utf-8")
    assert "Aucune règle d'exclusion" in out and "−,07" in out


def test_the_exclusion_commands(tmp_path, capsys):
    path = _folder(tmp_path)
    assert main(["exclusions", "lister", str(path)]) == 0
    assert "Aucune règle d'exclusion" in capsys.readouterr().out
    _rules_by_command(path)
    out = capsys.readouterr().out
    assert "Règle ajoutée : E1 — consentement = 2" in out
    assert "Règle ajoutée : E4 — passation de moins de 180 s, de debut à fin" in out
    assert main(["exclusions", "lister", str(path)]) == 0
    out = capsys.readouterr().out
    assert _counts(out) == EXPECTED
    assert "Réponses reçues : 400 ; exclues : 53 ; échantillon analysé : 347." in out
    assert f"   Source : {RULES[3]['source']}" in out
    assert main(["exclusions", "deplacer", str(path), "E5", "4"]) == 0
    assert "E5 est maintenant au rang 4." in capsys.readouterr().out
    assert main(["exclusions", "lister", str(path)]) == 0
    assert _counts(capsys.readouterr().out) == OTHER_ORDER
    assert main(["exclusions", "retirer", str(path), "E6"]) == 0
    assert [r.id for r in load_project(path).exclusions.rules] == ["E1", "E2", "E3", "E5", "E4"]
    assert main(["exclusions", "retirer", str(path), "E6"]) == 2
    assert "E6 : aucune règle ne porte ce numéro." in capsys.readouterr().err


def test_the_exclusion_commands_refuse(tmp_path, capsys):
    path = _folder(tmp_path)
    for arguments, reason in (
        (["--condition", "courriel", "eq", "x@y", "--raison", "x"], "nommerait une personne"),
        (["--duree", "debut", "fin", "180", "--raison", "Trop court"], "donnez sa source"),
        (["--condition", "age", "lt", "dix-huit", "--raison", "x"], "n'est pas un nombre"),
        (["--condition", "age", "environ", "18", "--raison", "x"], "Opérateur inconnu"),
        (["--completude", "G1", "9", "--raison", "x"], "le groupe en compte 6"),
        (["--ligne-droite", "G9", "--raison", "x"], "groupe inconnu"),
    ):
        assert main(["exclusions", "ajouter", str(path), *arguments]) == 2
        err = capsys.readouterr().err
        assert "Règle refusée" in err and reason in err
    assert load_project(path).exclusions.rules == []


def test_the_quality_signals_by_command(tmp_path, capsys):
    path = _folder(tmp_path)
    assert main(["exclusions", "signaux", str(path)]) == 0
    out = capsys.readouterr().out
    assert "parmi les 400 personnes" in out
    assert "Durée : choisissez le début et la fin parmi debut, fin." in out
    assert "Ligne droite, G2 (8 items) :" in out
    assert main(["exclusions", "signaux", str(path), "--debut", "debut", "--fin", "fin",
                 "--duree-min", "180", "--doublons", "courriel"]) == 0  # fmt: skip
    out = capsys.readouterr().out
    assert "Durée, de debut à fin :" in out and "Moins de 180 s (3 min) :" in out
    assert "Doublons sur courriel :" in out
    assert load_project(path).exclusions.rules == []  # a signal is never applied


def test_the_method_in_the_notebook_by_command(tmp_path, capsys):
    path = _folder(tmp_path)
    assert main(["carnet", "echantillon", str(path), "age", "genre"]) == 0
    assert "Variables qui décrivent l'échantillon : age, genre" in capsys.readouterr().out
    assert main(["carnet", "ajouter", str(path), *ANALYSIS]) == 0
    assert main(["carnet", "lister", str(path)]) == 0
    assert "M. Méthode — à jour" in capsys.readouterr().out
    assert main(["exclusions", "ajouter", str(path), *_arguments(RULES[0])]) == 0
    assert "seront recalculées" in capsys.readouterr().out
    assert main(["resultats", str(path)]) == 2
    err = capsys.readouterr().err
    assert "A1, la Méthode (M)" in err
    assert main(["carnet", "lister", str(path)]) == 0
    out = capsys.readouterr().out
    assert "M. Méthode — changée, à confirmer" in out
    assert "n_analysed : 400 → 388" in out and "excluded_E1" not in out  # new: kept at once
    assert main(["carnet", "confirmer", str(path), "--tout"]) == 0
    out = capsys.readouterr().out
    assert "Confirmée : A1" in out and "Confirmée : la Méthode (M)" in out
    assert main(["resultats", str(path)]) == 0
    texts = [p.text for p in Document(path.with_name("etude-passation-resultats.docx")).paragraphs]
    assert "Méthode" in texts and "Participants" in texts
    assert any(t.startswith("Au total, 400 personnes") for t in texts)
    assert main(["carnet", "echantillon", str(path), "--vider"]) == 0
    assert "Variables qui décrivent l'échantillon : aucune" in capsys.readouterr().out


# --- The page. ---


def _open(path, mode="method"):
    app = AppTest.from_file(APP, default_timeout=60)
    app.session_state["project_path"] = str(path)
    app.session_state["mode"] = mode
    app.run()
    assert not app.exception, app.exception
    return app


def _add(app, kind=None, reason="", source="", **fields):
    if kind is not None:
        app.selectbox(key="method-kind").set_value(kind).run()
    for name, value in fields.items():
        widget = app.multiselect if name == "duplicates" else None
        key = f"method-new-{name}"
        if widget is None:
            found = [w for w in (*app.selectbox, *app.text_input, *app.number_input)
                     if w.key == key]  # fmt: skip
            widget = found[0]
            widget.set_value(value).run()
        else:
            widget(key=key).set_value(value).run()
    app.text_input(key="method-new-reason").set_value(reason).run()
    app.text_input(key="method-new-source").set_value(source).run()
    app.button(key="method-add").click().run()
    assert not app.exception, app.exception


def _rules_on_the_page(app):
    """The rules of the answer key, by the form and by the signals' buttons."""
    first, second, third, fourth, fifth, sixth = RULES
    _add(app, "condition", first["reason"], variable="consentement", operator="eq", value="2")
    assert "Règle ajoutée : E1." in app.success[0].value
    _add(app, "condition", second["reason"], variable="age", operator="lt", value="18")
    _add(app, "completeness", third["reason"], group="G1", min=5)
    _add(app, "duration", fourth["reason"], start="debut", end="fin", seconds=180.0)
    assert "donnez sa source" in app.error[0].value  # refused, kept in the form
    _add(app, None, fourth["reason"], fourth["source"])
    app.button(key="method-signal-grid-G2").click().run()
    assert app.selectbox(key="method-kind").value == "straightlining"
    _add(app, None, fifth["reason"])
    app.button(key="method-signal-duplicate-courriel").click().run()
    _add(app, None, sixth["reason"])


@pytest.mark.allow_hosts(["127.0.0.1"])
def test_the_method_page(tmp_path):
    path = _folder(tmp_path)
    app = _open(path)
    assert app.header[0].value == "Méthode"
    assert "Aucune règle d'exclusion" in app.info[0].value
    _rules_on_the_page(app)
    project = load_project(path)
    assert [r.id for r in project.exclusions.rules] == ["E1", "E2", "E3", "E4", "E5", "E6"]
    assert [m.value for m in app.metric] == ["400", "53", "347"]
    assert "12 exclues" in " ".join(m.value for m in app.markdown)
    app.button(key="method-rule-E5-up").click().run()
    assert [r.id for r in load_project(path).exclusions.rules][3] == "E5"
    page = " ".join(m.value for m in app.markdown)
    assert "8 exclues" in page and "Au total, 400 personnes" in page
    assert "à compléter : ce que mesure le groupe G1" in page and "background:#fff59d" in page
    app.multiselect(key="method-sample").set_value(["age", "genre"]).run()
    assert not app.exception, app.exception
    notebook = nb.load_notebook(nb.notebook_path(path))
    assert nb.sample_names(notebook, load_project(path)) == ["age", "genre"]
    assert "M = 45,07" in " ".join(m.value for m in app.markdown).replace("<i>", "").replace(
        "</i>", ""
    )


@pytest.mark.allow_hosts(["127.0.0.1"])
def test_confirm_everything_on_the_notebook_page(tmp_path):
    path = _folder(tmp_path)
    assert main(["carnet", "ajouter", str(path), *ANALYSIS]) == 0
    app = _open(path, "notebook")
    assert "M. Méthode" in " ".join(m.value for m in app.markdown)
    assert main(["exclusions", "ajouter", str(path), *_arguments(RULES[1])]) == 0
    app = _open(path, "notebook")  # the page keeps its project for the session
    assert "A1, la Méthode (M)" in app.warning[0].value
    assert app.button(key="notebook-write").disabled
    app.button(key="notebook-confirm-all").click().run()
    assert not app.exception, app.exception
    assert not app.warning
    assert not app.button(key="notebook-write").disabled
    notebook = nb.load_notebook(nb.notebook_path(path))
    assert notebook.method.pending is None and notebook.method.computed.n_used == 391  # 9 minors


@pytest.mark.allow_hosts(["127.0.0.1"])
def test_the_page_gives_the_command_files(tmp_path):
    page_path = _folder(tmp_path, "page")
    app = _open(page_path)
    _rules_on_the_page(app)
    app.button(key="method-rule-E5-up").click().run()
    app.multiselect(key="method-sample").set_value(["age", "genre"]).run()
    assert main(["carnet", "ajouter", str(page_path), *ANALYSIS]) == 0
    app.session_state["mode"] = "notebook"
    app.run()
    app.button(key="notebook-write").click().run()
    assert "Fichiers écrits à côté du projet" in app.success[0].value

    command_path = _folder(tmp_path, "commande")
    _rules_by_command(command_path)
    assert main(["exclusions", "deplacer", str(command_path), "E5", "4"]) == 0
    assert main(["carnet", "echantillon", str(command_path), "age", "genre"]) == 0
    assert main(["carnet", "ajouter", str(command_path), *ANALYSIS]) == 0
    assert main(["resultats", str(command_path)]) == 0
    assert load_project(page_path).exclusions == load_project(command_path).exclusions
    name = "etude-passation-resultats-R.zip"
    assert page_path.with_name(name).read_bytes() == command_path.with_name(name).read_bytes()
    name = "etude-passation-resultats.docx"
    page, command = Document(page_path.with_name(name)), Document(command_path.with_name(name))
    assert [p.text for p in page.paragraphs] == [p.text for p in command.paragraphs]
    assert [[c.text for c in t._cells] for t in page.tables] == [
        [c.text for c in t._cells] for t in command.tables
    ]
    name = "etude-passation-resultats-graphiques.xlsx"
    with zipfile.ZipFile(io.BytesIO(page_path.with_name(name).read_bytes())) as book:
        assert "xl/worksheets/sheet1.xml" in book.namelist()


def test_the_method_differences_are_written_as_apa():
    from ddict.analysis.notebook_check import _difference

    assert _difference("alpha_n_G1", 370.0, 362.0).shown_after == "362"
    assert _difference("excluded_E1", 12.0, 20.0).shown_after == "20"
    assert _difference("v1_n", 374.0, 380.0).shown_after == "380"
    assert _difference("alpha_G1", 0.8904, 0.8911).shown_after == ",891"
    assert _difference("v1_m", 44.12, 45.07).shown_after == "45,07"
