"""v0.20.2 (D-158): a variable set aside from the codebook by the team, with its reason —
format 0.10.0; refused while something names it, each use said; undone by "Remettre".
The fictional base only."""

import json
from types import SimpleNamespace

import pytest

from ddict.io import read_sav
from ddict.project import SCHEMA_VERSION, aside, build_project, load_project, save_project
from ddict.project import exclusions as rules
from ddict.project.model import DuplicateRule, TrackedValue, VariableRef
from ddict.project.recodes import add as add_recode
from ddict.project.recodes import validated

TODAY = "2026-10-20"


@pytest.fixture
def project(fictional_files):
    return build_project(read_sav(fictional_files["etude-sommeil.sav"]))


def test_set_aside_and_back(project, tmp_path):
    assert SCHEMA_VERSION == "0.10.0"
    aside.set_aside(project, "q_retiree", "Question retirée de l'étude.", TODAY)
    entry = project.variable("q_retiree")
    assert aside.is_aside(entry) and entry.set_aside.reason == "Question retirée de l'étude."
    assert "q_retiree" not in aside.kept_names(project)
    assert [e.name for e in aside.aside(project)] == ["q_retiree"]
    path = tmp_path / "etude.dd.json"
    save_project(project, path)
    data = json.loads(path.read_text(encoding="utf-8"))
    saved = next(v for v in data["variables"] if v["name"] == "q_retiree")
    assert saved["set_aside"] == {"reason": "Question retirée de l'étude.", "date": TODAY}
    assert data["schema_version"] == "0.10.0"
    again = load_project(path)
    assert aside.restore(again, "q_retiree") is True
    assert aside.restore(again, "q_retiree") is False  # not set aside any more
    assert "q_retiree" in aside.kept_names(again)


def test_the_fingerprint_changes_only_for_a_variable_set_aside(project):
    from ddict.project import fingerprint

    before = fingerprint.fingerprint(project)
    aside.set_aside(project, "q_retiree", "Question retirée.", TODAY)
    assert fingerprint.fingerprint(project) != before
    aside.restore(project, "q_retiree")
    assert fingerprint.fingerprint(project) == before


def test_a_reason_is_needed(project):
    with pytest.raises(ValueError):
        aside.set_aside(project, "q_retiree", "  ", TODAY)
    assert project.variable("q_retiree").set_aside is None


def test_a_0_9_0_project_reads_as_it_is(project, tmp_path):
    path = tmp_path / "etude.dd.json"
    save_project(project, path)
    data = json.loads(path.read_text(encoding="utf-8"))
    data["schema_version"] = "0.9.0"
    for entry in data["variables"]:
        entry.pop("set_aside", None)
    path.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")
    old = load_project(path)
    assert old.schema_version == "0.10.0" and not aside.aside(old)


def test_a_variable_in_use_is_refused_each_use_said(project):
    group = project.groups[0]
    item = group.members.value[0]
    rule = rules.add(project, {"kind": "duplicate", "variables": ["courriel"],
                               "reason": "Une seule réponse par personne."})  # fmt: skip
    recode = add_recode(project, validated(project, {
        "kind": "classes", "source": {"variable": "age"}, "name": "age_classes",
        "label": "Âge en classes", "cuts": [30], "labels": ["Moins de 30", "30 et plus"],
        "reason": "Deux classes."}))  # fmt: skip
    project.dataset.design.wave = "vague"
    found = {
        item: [f"le groupe {group.group_id}"],
        "courriel": [f"la règle d'exclusion {rule.id}"],
        "age": [f"le recodage {recode.id}"],
        "vague": ["la structure de la base (page « Base »)"],
    }
    for name, uses in found.items():
        with pytest.raises(aside.InUse) as refused:
            aside.set_aside(project, name, "Pas utile.", TODAY)
        assert refused.value.uses == uses
        assert project.variable(name).set_aside is None
    assert isinstance(rule, DuplicateRule)


def test_a_variable_of_the_notebook_is_refused(project):
    from ddict.analysis.notebook import MethodPart

    class Entry:  # what the notebook keeps of an analysis: its variables by their key
        id = "A1"

        def model_dump(self, mode):
            return {"variables": {"variables": [{"variable": "heures_sommeil"},
                                                {"variable": "nb_reveils"}]}}  # fmt: skip

    notebook = SimpleNamespace(
        entries=[Entry()], method=MethodPart(variables=[VariableRef(variable="age")])
    )
    assert aside.uses(project, "nb_reveils", notebook) == ["l'analyse A1 du carnet"]
    assert aside.uses(project, "age", notebook) == [
        "la description de l'échantillon (page « Méthode »)"
    ]
    assert aside.uses(project, "nb_reveils") == []


# --- The outputs (step 2) ----------------------------------------------------------------------


def _aside(project):
    project.variable("q_retiree").label = TrackedValue[str](value="Retirée", origin="human")
    aside.set_aside(project, "q_retiree", "Question retirée de l'étude.", TODAY)
    return project


def test_the_codebook_lists_it_apart(project, tmp_path):
    from openpyxl import load_workbook

    from ddict.export.markdown import render_markdown
    from ddict.export.xlsx import write_xlsx

    _aside(project)
    path = tmp_path / "codebook.xlsx"
    write_xlsx(project, path)
    book = load_workbook(path)
    names = {
        cell for row in book["Variables"].iter_rows(min_row=2, values_only=True) for cell in row
    }
    assert "q_retiree" not in names and "age" in names
    assert list(book["Variables écartées"].iter_rows(values_only=True)) == [
        ("Variable", "Raison", "Date"),
        ("q_retiree", "Question retirée de l'étude.", TODAY),
    ]
    text = render_markdown(project)
    assert (
        "## Variables écartées" in text and "- `q_retiree` : Question retirée de l'étude." in text
    )
    assert "### q_retiree" not in text and "Retirée" not in text


def test_a_codebook_without_one_has_no_such_part(project, tmp_path):
    from openpyxl import load_workbook

    from ddict.export.markdown import render_markdown
    from ddict.export.xlsx import write_xlsx

    path = tmp_path / "codebook.xlsx"
    write_xlsx(project, path)
    assert "Variables écartées" not in load_workbook(path).sheetnames
    assert "Variables écartées" not in render_markdown(project)


def test_the_context_leaves_it_out_and_its_package_keeps_the_column(project):
    from ddict.export.context import context_package, render_context

    _aside(project)
    assert "q_retiree" not in render_context(project)
    fields = context_package(project)["resources"][0]["schema"]["fields"]
    field = next(f for f in fields if f["name"] == "q_retiree")
    assert field["_ddict"] == {"set_aside": "Question retirée de l'étude."}
    assert "title" not in field
    assert len(fields) == len(project.variables)  # the columns of the file, all of them


def test_the_command_files_leave_it_alone(project):
    from ddict.export.syntax import render_syntax

    _aside(project)
    for program in ("spss", "r"):
        assert "q_retiree" not in render_syntax(project, program, "2026-10-20")


def test_the_documented_copy_keeps_its_column_as_in_the_file(project, fictional_files, tmp_path):
    import shutil

    import pyreadstat

    from ddict.export.documented import write_documented_copy

    database = tmp_path / "etude-sommeil.sav"
    shutil.copy(fictional_files["etude-sommeil.sav"], database)
    _aside(project)
    result = write_documented_copy(project, database)
    _data, meta = pyreadstat.read_sav(str(result.path), metadataonly=True)
    labels = dict(zip(meta.column_names, meta.column_labels, strict=True))
    assert "q_retiree" in labels and labels["q_retiree"] != "Retirée"
    assert labels["q_retiree"] == project.variable("q_retiree").file.label


# --- The counters and the choices (step 3) -----------------------------------------------------


def test_it_leaves_the_counters_and_the_filters(project):
    from ddict.review import logic

    _done, total = logic.progress(project)
    _aside(project)
    assert logic.progress(project)[1] == total - 1
    assert [e.name for e in logic.select(project, "aside")] == ["q_retiree"]
    assert "q_retiree" not in [e.name for e in logic.select(project, "all")]


def test_no_analysis_no_rule_no_recoding_takes_it(project):
    from ddict.analysis import conditions
    from ddict.project.exclusions import ExclusionRefused
    from ddict.project.recodes import RecodeRefused

    aside.set_aside(project, "nb_reveils", "Pas utile ici.", TODAY)
    assert "nb_reveils" in conditions.excluded(project)
    with pytest.raises(ExclusionRefused) as refused:
        rules.add(project, {"kind": "condition", "variable": "nb_reveils", "operator": "missing",
                            "reason": "Sans réponse."})  # fmt: skip
    assert "nb_reveils : variable écartée du codebook." in str(refused.value)
    with pytest.raises(RecodeRefused) as refused:
        validated(project, {
            "kind": "classes", "source": {"variable": "nb_reveils"}, "name": "reveils_classes",
            "label": "Réveils en classes", "cuts": [2], "labels": ["Peu", "Beaucoup"],
            "reason": "Deux classes."})  # fmt: skip
    assert "nb_reveils : variable écartée du codebook." in str(refused.value)


def test_no_new_name_is_proposed_for_it(project):
    from ddict import naming

    aside.set_aside(project, "q_retiree", "Question retirée.", TODAY)
    naming.propose(project)
    assert project.variable("q_retiree").new_name.value is None
    assert "q_retiree" not in naming.to_name(project)


# --- The page (step 4) -------------------------------------------------------------------------


@pytest.mark.allow_hosts(["127.0.0.1"])
def test_set_aside_on_the_page_then_back(project, tmp_path):
    from pathlib import Path

    from streamlit.testing.v1 import AppTest

    path = tmp_path / "etude.dd.json"
    save_project(project, path)
    app_path = Path(__file__).resolve().parent.parent / "src" / "ddict" / "review" / "app.py"
    app = AppTest.from_file(str(app_path), default_timeout=60)
    app.session_state["project_path"] = str(path)
    app.session_state["mode"] = "variables"
    app.run()
    assert "aside" not in app.sidebar.radio(key="filter").options  # none set aside yet
    app.sidebar.radio(key="filter").set_value("all").run()
    app.selectbox(key="variable").set_value("q_retiree").run()
    app.button(key="aside-q_retiree").click().run()  # no reason: refused
    assert any("Donnez la raison" in e.value for e in app.error)
    assert load_project(path).variable("q_retiree").set_aside is None
    app.text_input(key="aside-reason-q_retiree").set_value("Question retirée de l'étude.").run()
    app.button(key="aside-q_retiree").click().run()
    assert not app.exception, app.exception
    assert load_project(path).variable("q_retiree").set_aside.reason == (
        "Question retirée de l'étude."
    )
    app.sidebar.radio(key="filter").set_value("aside").run()
    app.selectbox(key="variable").set_value("q_retiree").run()
    assert any("Écartée du codebook le" in w.value for w in app.warning)
    app.button(key="restore-q_retiree").click().run()
    assert not app.exception, app.exception
    assert load_project(path).variable("q_retiree").set_aside is None


@pytest.mark.allow_hosts(["127.0.0.1"])
def test_a_variable_in_use_is_refused_on_the_page(project, tmp_path):
    from pathlib import Path

    from streamlit.testing.v1 import AppTest

    item = project.groups[0].members.value[0]
    path = tmp_path / "etude.dd.json"
    save_project(project, path)
    app_path = Path(__file__).resolve().parent.parent / "src" / "ddict" / "review" / "app.py"
    app = AppTest.from_file(str(app_path), default_timeout=60)
    app.session_state["project_path"] = str(path)
    app.session_state["mode"] = "variables"
    app.run()
    app.sidebar.radio(key="filter").set_value("all").run()
    app.selectbox(key="variable").set_value(item).run()
    app.text_input(key=f"aside-reason-{item}").set_value("Pas utile.").run()
    app.button(key=f"aside-{item}").click().run()
    group = project.groups[0].group_id
    assert any(f"sert encore : le groupe {group}" in e.value for e in app.error)
    assert load_project(path).variable(item).set_aside is None
