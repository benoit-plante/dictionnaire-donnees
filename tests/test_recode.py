"""v0.15, step 2: the recodings (D-120; docs/verification-depot.md, § 1): a derived variable
kept as a decision of the team, computed in memory after the scores and before the
exclusion rules, refused when a value has no place, and usable in the analyses, the
exclusion rules, the subgroups, the notebook and the plan."""

import json

import pandas as pd
import pytest

from ddict.analysis import conditions, recode
from ddict.analysis import notebook as nb
from ddict.analysis.book import Codebook
from ddict.analysis.catalogue import CATALOGUE
from ddict.analysis.compute import run
from ddict.analysis.data import apply_decisions, decided
from ddict.analysis.results import Subset
from ddict.io import read_database
from ddict.project import build_project, fingerprint, load_project, recodes, save_project
from ddict.project import exclusions as rules
from ddict.project.model import VariableRef
from test_paired import EXAMPLES, KEY
from test_text_categories import _base

SEXE_2 = {
    "kind": "groups",
    "source": {"variable": "sexe"},
    "name": "sexe_2",
    "label": "Sexe, en deux groupes",
    "groups": [
        {"code": 1, "label": "Femme", "values": [1]},
        {"code": 2, "label": "Homme ou autre identité", "values": [2, 3]},
    ],
    "reason": "Trop peu de personnes d'une autre identité pour un groupe à part.",
}
AGE_CLASSE = {
    "kind": "classes",
    "source": {"variable": "age"},
    "name": "age_classe",
    "label": "Âge, en trois classes",
    "cuts": [30, 45],
    "labels": ["moins de 30", "30 à moins de 45", "45 et plus"],
    "reason": "Classes d'âge du rapport de l'organisme.",
}


@pytest.fixture
def base():
    data = read_database(EXAMPLES / KEY["sav"])
    return build_project(data), data.table


def _added(project, *fields):
    for data in fields:
        recodes.add(project, recodes.validated(project, data))
    return project


def _checked(project, table, data):
    rule = recodes.validated(project, data)
    before, _rev, _scores = decided(project, table)
    book = Codebook(project, apply_decisions(project, table))
    source = recodes.source_name(project, rule.source)
    return recode.check(project, rule, before, book.value_labels(source))


# --- The examples of the sheet. ---


def test_the_examples_of_the_sheet(base):
    project, table = base
    sexe = _checked(project, table, SEXE_2)
    assert sexe.ok and sexe.counts == [("Femme", 133), ("Homme ou autre identité", 107)]
    age = _checked(project, table, AGE_CLASSE)
    assert age.ok and age.counts == [
        ("moins de 30", 49),
        ("30 à moins de 45", 79),
        ("45 et plus", 112),
    ]
    assert sexe.warnings == age.warnings == []


def test_in_memory_after_the_scores(base):
    project, table = base
    _added(project, SEXE_2, AGE_CLASSE)
    data = apply_decisions(project, table).table
    assert list(data.columns[-2:]) == ["sexe_2", "age_classe"]
    assert data["sexe_2"].value_counts().to_dict() == {1.0: 133, 2.0: 107}
    # A cut belongs to the class it starts, as R's findInterval.
    ages = pd.Series([29.9, 30.0, 44.99, 45.0, None])
    rule = project.recodes.rules[1]
    assert recode.column(rule, ages).tolist()[:4] == [1.0, 2.0, 2.0, 3.0]
    assert pd.isna(recode.column(rule, ages).iloc[4])


def test_the_project_keeps_the_rule_never_the_column(base, tmp_path):
    project, _table = base
    before = fingerprint.fingerprint(project)
    _added(project, SEXE_2, AGE_CLASSE)
    assert fingerprint.fingerprint(project) != before  # a recoding is a decision
    path = tmp_path / "etude-suivi.dd.json"
    save_project(project, path)
    saved = json.loads(path.read_text(encoding="utf-8"))
    assert saved["schema_version"] == "0.9.0"
    assert saved["recodes"]["next_id"] == 3
    assert saved["recodes"]["rules"][0] == {"id": "R1", **SEXE_2, "level": "nominal",
                                            "missing_values": []}  # fmt: skip
    assert load_project(path) == project
    # A project without recodings keeps its fingerprint.
    assert fingerprint.fingerprint(build_project(read_database(EXAMPLES / KEY["sav"]))) == before


# --- Refusals. ---


@pytest.mark.parametrize(
    ("change", "reason"),
    [
        ({"reason": " "}, "Donnez la raison du recodage, en clair."),
        ({"groups": [{"code": 1, "label": "Tous", "values": [1, 2, 3]}]},
         "Il faut au moins deux groupes."),
        ({"groups": [{"code": 1, "label": "A", "values": [1]},
                     {"code": 1, "label": "B", "values": [2, 3]}]},
         "Deux groupes ont le même code."),
        ({"groups": [{"code": 1, "label": "A", "values": [1, 2]},
                     {"code": 2, "label": "B", "values": [2.0, 3]}]},
         "Une valeur est dans deux groupes."),
        ({"source": {"recode": "R1"}}, "Une variable recodée se recode depuis sa propre source."),
        ({"source": {"variable": "inconnue"}}, "inconnue : variable inconnue du projet."),
        ({"source": {"variable": "id_participant"}},
         "id_participant : identifiant, texte libre ou date ; une variable ne se recode pas "
         "ainsi."),
        ({"source": {"variable": "age"}},
         "age : seules des catégories (nominales ou ordinales) se regroupent ; un nombre se "
         "découpe en classes."),
        ({"name": "age"}, "age : nom déjà pris."),
        ({"name": "anxiete_t1"}, "anxiete_t1 : nom déjà pris."),
        ({"name": "Sexe 2"}, "Sexe 2 : nom invalide (lettres minuscules, chiffres et _, au plus "
                             "32 caractères, pas un mot réservé)."),
    ],
)  # fmt: skip
def test_refused_groups(base, change, reason):
    project, _table = base
    with pytest.raises(recodes.RecodeRefused) as refused:
        recodes.validated(project, {**SEXE_2, **change})
    assert reason in refused.value.reasons


@pytest.mark.parametrize(
    ("change", "reason"),
    [
        ({"cuts": [45, 30]}, "Les seuils doivent aller croissant."),
        ({"cuts": []}, "Il faut au moins un seuil."),
        ({"labels": ["moins de 30", "30 et plus"]},
         "Une étiquette par classe : une de plus que les seuils."),
        ({"source": {"variable": "sexe"}},
         "sexe : seul un nombre (une échelle ou un score) se découpe en classes ; des catégories "
         "se regroupent."),
    ],
)  # fmt: skip
def test_refused_classes(base, change, reason):
    project, _table = base
    with pytest.raises(recodes.RecodeRefused) as refused:
        recodes.validated(project, {**AGE_CLASSE, **change})
    assert reason in refused.value.reasons


def test_every_observed_value_has_a_place(base):
    project, table = base
    two = {**SEXE_2, "groups": [{"code": 1, "label": "Femme", "values": [1]},
                                {"code": 2, "label": "Homme", "values": [2]}]}  # fmt: skip
    found = _checked(project, table, two)
    assert found.refusals == [
        "Valeurs sans groupe : 3 (Autre identité, 14 personnes). Mettez-les dans un groupe, ou "
        "déclarez-les manquantes."
    ]
    declared = _checked(project, table, {**two, "missing_values": [3]})
    assert declared.ok and declared.counts == [("Femme", 133), ("Homme", 93)]


def test_warnings_empty_class_and_small_category(base):
    project, table = base
    found = _checked(project, table, {**AGE_CLASSE, "cuts": [23, 70],
                                      "labels": ["jeunes", "adultes", "aînés"]})  # fmt: skip
    assert found.ok
    assert found.warnings == ["Moins de 5 personnes : jeunes (1).", "Classe vide : aînés."]


def test_text_categories_regroup(tmp_path):
    project, _path = _base(tmp_path)
    table = read_database(_path.with_name("etude-textes.csv")).table
    data = {"kind": "groups", "source": {"variable": "region"}, "name": "region_2",
            "label": "Région, en deux", "reason": "Comparer le centre et le reste.",
            "groups": [{"code": 1, "label": "Centre-ville", "values": ["centre-ville"]},
                       {"code": 2, "label": "Ailleurs",
                        "values": ["nord-est", "sud", "ouest"]}]}  # fmt: skip
    found = _checked(project, table, data)
    assert found.ok and sum(n for _label, n in found.counts) == len(table)


# --- Where a recoded variable goes. ---


def test_in_an_analysis_with_its_labels(base):
    project, table = base
    _added(project, SEXE_2)
    prepared = apply_decisions(project, table)
    assert conditions.levels(project)["sexe_2"] == "nominal"
    assert "sexe_2" in conditions.candidates(project, CATALOGUE["t_test"], "group")
    result = run(project, prepared, "t_test", {"outcome": ["stress"], "group": ["sexe_2"]})
    assert "pour Femme et M = " in result.sentence
    assert "pour Homme ou autre identité" in result.sentence
    assert Codebook(project, prepared).label("sexe_2") == "Sexe, en deux groupes"


def test_in_an_exclusion_rule_and_a_subgroup(base):
    project, table = base
    _added(project, AGE_CLASSE)
    rules.add(project, {"kind": "condition", "recode": "R1", "operator": "eq", "value": 1,
                        "reason": "Moins de 30 ans"})  # fmt: skip
    prepared = apply_decisions(project, table)
    assert len(prepared.table) == 240 - 49
    assert prepared.flow.excluded == (("E1", 49),)
    project.exclusions.rules.clear()
    prepared = apply_decisions(project, table)
    result = run(project, prepared, "paired", {"first": ["anxiete_t1"],
                                               "second": ["anxiete_t2"]},
                 None, Subset("age_classe", "eq", 3.0))  # fmt: skip
    assert result.sentence.startswith("Sous-groupe (age_classe = 45 et plus).")


def test_in_the_notebook_by_its_rule(base):
    project, _table = base
    _added(project, SEXE_2)
    ref = nb.reference(project, "sexe_2")
    assert ref == VariableRef(recode="R1")
    project.recodes.rules[0].name = "sexe_deux"  # a renaming keeps the reference
    assert nb.resolve(project, ref) == "sexe_deux"
    project.recodes.rules.clear()
    with pytest.raises(nb.NotebookError) as error:
        nb.resolve(project, ref)
    assert nb.reason_text(error.value) == "R1 : le recodage n'est plus dans le projet."


def test_removal_refused_while_used(base):
    project, _table = base
    _added(project, AGE_CLASSE)
    rules.add(project, {"kind": "condition", "recode": "R1", "operator": "eq", "value": 1,
                        "reason": "Moins de 30 ans"})  # fmt: skip
    with pytest.raises(recodes.RecodeRefused) as refused:
        recodes.remove(project, "R1")
    assert refused.value.reasons == ["R1 : utilisé par E1 ; retirez-les ou changez-les d'abord."]
    project.exclusions.rules.clear()
    recodes.remove(project, "R1")
    assert project.recodes.rules == [] and project.recodes.next_id == 2


def test_in_a_hypothesis_and_its_frozen_notebook(base, tmp_path):
    from ddict.analysis import plan_notebook as pn
    from ddict.analysis import planning as pl

    project, table = base
    _added(project, SEXE_2, AGE_CLASSE)
    prepared = apply_decisions(project, table)
    data = pl.data_from_names(project, "Le stress diffère selon le sexe.", "t_test",
                              {"outcome": ["stress"], "group": ["sexe_2"]}, {},
                              subgroup=Subset("age_classe", "ge", 2.0))  # fmt: skip
    hypothesis = pl.add(project, data, prepared, None, "2026-10-18")
    assert hypothesis.roles["group"] == [VariableRef(recode="R1")]
    assert hypothesis.subgroup.variable == VariableRef(recode="R2")
    assert pl.review(project, hypothesis, prepared).ok
    with pytest.raises(recodes.RecodeRefused) as refused:
        recodes.remove(project, "R1")
    assert refused.value.reasons == ["R1 : utilisé par H1 ; retirez-les ou changez-les d'abord."]
    notebook = nb.new_notebook(project, tmp_path / "etude-suivi.dd.json")
    project.plan.alpha, project.plan.data_origin = 0.05, "primary"
    project.plan.data_status = "before_analysis"
    pn.freeze(notebook, project, prepared, "2026-10-18")
    statistics = notebook.entries[0].computed.statistics
    assert statistics["n1"] + statistics["n2"] == 240 - 49


# --- In the preparation scripts, R and PSPP (step 3). ---


def _r_master(project, path, tmp_path, cases):
    """The tool's statistics and R's, by notebook entry, from the R master script."""
    import io
    import subprocess
    import zipfile

    from ddict.analysis import notebook_export as export
    from ddict.analysis.data import prepare
    from test_syntax import _rscript

    prepared = prepare(project, path)
    notebook = nb.new_notebook(project, path)
    tool = {}
    for key, chosen, subset in cases:
        result = run(project, prepared, key, chosen, None, subset)
        tool[nb.add(notebook, project, result, today="2026-10-18").id] = result.statistics
    analyses = [(e.id, e.analysis, nb.resolved(project, e), dict(e.options), e.subgroup)
                for e in notebook.entries]  # fmt: skip
    analyses = [
        (i, k, c, o, nb.resolved_subgroup(project, s)) if s else (i, k, c, o)
        for i, k, c, o, s in analyses
    ]
    bundle = export.r_bundle(project, path, analyses, today="2026-10-18")
    with zipfile.ZipFile(io.BytesIO(bundle)) as z:
        z.extractall(path.parent)
        assert f"{path.name.removesuffix('.dd.json')}-preparation.R" in z.namelist()
    (path.parent / "run.R").write_text(
        f'source("{path.name.removesuffix(".dd.json")}-article.R", encoding = "UTF-8")\n'
        "valeurs <- unlist(resultats)\n"
        'write.csv(data.frame(nom = names(valeurs), valeur = valeurs), "sortie.csv", '
        "row.names = FALSE)\n",
        encoding="utf-8",
    )
    done = subprocess.run([_rscript(), "run.R"], cwd=path.parent, capture_output=True,
                          text=True, encoding="utf-8", errors="replace")  # fmt: skip
    if "there is no package called" in done.stderr:
        pytest.skip("haven is not installed in R")
    assert done.returncode == 0, done.stderr[-2000:]
    out = pd.read_csv(path.parent / "sortie.csv")
    from_r: dict[str, dict[str, float]] = {}
    for name, value in zip(out["nom"], out["valeur"], strict=True):
        entry_id, key = name.split(".", 1)
        from_r.setdefault(entry_id, {})[key] = value
    return tool, from_r


def _same(tool, from_r, keys):
    for entry_id, wanted in keys.items():
        for key in wanted:
            assert tool[entry_id][key] == pytest.approx(from_r[entry_id][key], rel=1e-6), (
                entry_id, key)  # fmt: skip


def test_r_gives_the_tool_values(tmp_path):
    from test_syntax import _rscript

    if _rscript() is None:
        pytest.skip("R is not installed")
    import shutil

    shutil.copy(EXAMPLES / KEY["sav"], tmp_path / KEY["sav"])
    project = build_project(read_database(tmp_path / KEY["sav"]))
    _added(project, SEXE_2, AGE_CLASSE)
    rules.add(project, {"kind": "condition", "recode": "R2", "operator": "eq", "value": 1,
                        "reason": "Moins de 30 ans"})  # fmt: skip
    path = tmp_path / "etude-suivi.dd.json"
    save_project(project, path)
    cases = [
        ("t_test", {"outcome": ["stress"], "group": ["sexe_2"]}, None),
        ("chi_square", {"rows": ["sexe_2"], "columns": ["age_classe"]}, None),
        ("paired", {"first": ["anxiete_t1"], "second": ["anxiete_t2"]},
         Subset("age_classe", "eq", 3.0)),
    ]  # fmt: skip
    tool, from_r = _r_master(project, path, tmp_path, cases)
    assert tool["A1"]["n1"] + tool["A1"]["n2"] == 240 - 49
    _same(tool, from_r, {"A1": ["t", "df", "p", "n1", "n2"], "A2": ["chi2", "p"],
                         "A3": ["t", "n", "p"]})  # fmt: skip


def test_r_regroups_text_categories(tmp_path):
    from test_syntax import _rscript

    if _rscript() is None:
        pytest.skip("R is not installed")
    project, path = _base(tmp_path)
    recodes.add(project, recodes.validated(project, {
        "kind": "groups", "source": {"variable": "region"}, "name": "region_2",
        "label": "Région, en deux", "reason": "Comparer le centre et le reste.",
        "groups": [{"code": 1, "label": "Centre-ville", "values": ["centre-ville"]},
                   {"code": 2, "label": "Ailleurs", "values": ["nord-est", "sud", "ouest"]}],
    }))  # fmt: skip
    save_project(project, path)
    cases = [("t_test", {"outcome": ["satisfaction"], "group": ["region_2"]}, None)]
    tool, from_r = _r_master(project, path, tmp_path, cases)
    _same(tool, from_r, {"A1": ["t", "df", "p", "n1", "n2"]})


def test_pspp_gives_the_same_columns(tmp_path):
    import subprocess

    import pyreadstat

    from ddict.export.preparation import render_preparation
    from ddict.export.syntax import write_syntax
    from test_syntax import _pspp

    if _pspp() is None:
        pytest.skip("PSPP is not installed (D-027)")
    sav = EXAMPLES / KEY["sav"]
    data = read_database(sav)
    project = build_project(data)
    _added(project, SEXE_2, AGE_CLASSE)
    rules.add(project, {"kind": "condition", "recode": "R2", "operator": "eq", "value": 1,
                        "reason": "Moins de 30 ans"})  # fmt: skip
    raw, _meta = pyreadstat.read_sav(str(sav), user_missing=True)
    bare = tmp_path / "bare.sav"
    pyreadstat.write_sav(raw, str(bare))
    write_syntax(project, {"spss": tmp_path / "m.sps"}, "2026-10-18")
    (tmp_path / "p.sps").write_text(render_preparation(project, "spss", "2026-10-18"),
                                    encoding="utf-8")  # fmt: skip
    job = tmp_path / "job.sps"
    job.write_text(
        f"GET FILE='{bare}'.\nINSERT FILE='{tmp_path / 'm.sps'}' ENCODING='UTF-8'.\n"
        f"INSERT FILE='{tmp_path / 'p.sps'}' ENCODING='UTF-8'.\n"
        f"SAVE OUTFILE='{tmp_path / 'out.sav'}'.\n",
        encoding="utf-8",
    )
    done = subprocess.run([_pspp(), str(job)], capture_output=True, text=True,
                          errors="replace")  # fmt: skip
    said = (done.stdout + done.stderr).lower()
    assert done.returncode == 0 and "error" not in said and "warning" not in said, said
    out, meta = pyreadstat.read_sav(str(tmp_path / "out.sav"))
    tool = apply_decisions(project, data.table).table
    assert len(out) == len(tool) == 240 - 49
    for name in ("sexe_2", "age_classe"):
        assert out[name].value_counts().to_dict() == tool[name].value_counts().to_dict()
    assert meta.variable_value_labels["sexe_2"] == {1.0: "Femme", 2.0: "Homme ou autre identité"}
    assert meta.variable_measure["age_classe"] == "ordinal"


# --- Scripts, codebook, Method, command and page (step 3). ---


def test_the_preparation_scripts(base):
    from ddict.export.preparation import needed, render_preparation

    project, _table = base
    assert not needed(project)
    _added(project, SEXE_2, AGE_CLASSE)
    assert needed(project)
    spss = render_preparation(project, "spss", "2026-10-18")
    assert (
        "* R1 : Trop peu de personnes d'une autre identité pour un groupe à part.\n"
        "RECODE sexe (1=1) (2, 3=2) (ELSE=SYSMIS) INTO sexe_2.\n"
        "VARIABLE LABELS sexe_2 'Sexe, en deux groupes'.\n"
        "VALUE LABELS sexe_2 1 'Femme' 2 'Homme ou autre identité'.\n"
        "VARIABLE LEVEL sexe_2 (NOMINAL).\n"
    ) in spss
    assert "IF (NOT MISSING(age) AND age >= 30 AND age < 45) age_classe = 2.\n" in spss
    r = render_preparation(project, "r", "2026-10-18")
    assert (
        'donnees[["age_classe"]] <- haven::labelled(findInterval(ddict_valeurs(donnees[["age"]]), '
        'c(30, 45)) + 1, labels = c("moins de 30" = 1, "30 à moins de 45" = 2, "45 et plus" = 3), '
        'label = "Âge, en trois classes")'
    ) in r
    assert r.index("sexe_2") < r.index("rm(list")


def test_the_codebook_and_the_context(base, tmp_path):
    from openpyxl import load_workbook

    from ddict.export.context import context_package, render_context
    from ddict.export.markdown import render_markdown
    from ddict.export.xlsx import write_xlsx

    project, _table = base
    write_xlsx(project, tmp_path / "sans.xlsx")
    assert "Recodages" not in load_workbook(tmp_path / "sans.xlsx").sheetnames
    assert "recodes" not in context_package(project)["_ddict"]
    _added(project, SEXE_2, AGE_CLASSE)
    write_xlsx(project, tmp_path / "avec.xlsx")
    rows = list(load_workbook(tmp_path / "avec.xlsx")["Recodages"].values)
    assert rows[0] == ("Recodage", "Variable", "Libellé", "Niveau", "Source", "Règle", "Raison")
    assert rows[2] == (
        "R2", "age_classe", "Âge, en trois classes", "ordinal", "age",
        "moins de 30 (< 30) ; 30 à moins de 45 (≥ 30 et < 45) ; 45 et plus (≥ 45)",
        "Classes d'âge du rapport de l'organisme.",
    )  # fmt: skip
    markdown = render_markdown(project)
    assert "## Variables recodées" in markdown
    assert "- Règle : Femme (1) ; Homme ou autre identité (2, 3)" in markdown
    context = render_context(project)
    assert "- `sexe_2` (nominal) — Sexe, en deux groupes : tirée de `sexe`" in context
    package = context_package(project)["_ddict"]["recodes"]
    assert package[0]["source"] == "sexe" and list(package[0])[:3] == ["id", "kind", "source"]


def test_the_method_sentence(base, tmp_path):
    from ddict.analysis import method, method_text
    from ddict.analysis.apa import plain

    project, table = base
    _added(project, SEXE_2, AGE_CLASSE)
    prepared = apply_decisions(project, table)
    notebook = nb.new_notebook(project, tmp_path / "etude-suivi.dd.json")
    result = run(project, prepared, "t_test", {"outcome": ["stress"], "group": ["sexe_2"]})
    nb.add(notebook, project, result, today="2026-10-18")
    found = method.facts(project, prepared, notebook)
    assert found.recodes == (
        "La variable sexe_2 a été tirée de sexe : Femme (1) ; Homme ou autre identité (2, 3).",
    )  # only the recodings the Method speaks of
    assert found.recodes[0] in plain(method_text.analyses(found, "x.R"))


def _project_file(tmp_path):
    import shutil

    shutil.copy(EXAMPLES / KEY["sav"], tmp_path / KEY["sav"])
    path = tmp_path / "etude-suivi.dd.json"
    save_project(build_project(read_database(tmp_path / KEY["sav"])), path)
    return path


def test_the_command(tmp_path, capsys):
    from ddict.cli import main

    path = _project_file(tmp_path)
    assert main(["recoder", "regrouper", str(path), "sexe", "--nom", "sexe_2", "--libelle",
                 "Sexe, en deux groupes", "--groupe", "1", "Femme", "--groupe", "2,3",
                 "Homme ou autre identité", "--raison", SEXE_2["reason"]]) == 0  # fmt: skip
    out = capsys.readouterr().out
    assert "Recodage enregistré : R1 (sexe_2)." in out and "Femme : n = 133" in out
    assert main(["recode", "cut", str(path), "age", "--name", "age_classe", "--label",
                 "Âge, en trois classes", "--cuts", "30", "45", "--reason",
                 AGE_CLASSE["reason"]]) == 0  # fmt: skip
    project = load_project(path)
    expected = _added(build_project(read_database(EXAMPLES / KEY["sav"])), SEXE_2, AGE_CLASSE)
    assert project.recodes == expected.recodes  # the default labels are the sheet's
    capsys.readouterr()
    assert main(["recoder", "regrouper", str(path), "sexe", "--nom", "sexe_x", "--groupe", "1",
                 "F", "--groupe", "2", "H", "--raison", "x"]) == 2  # fmt: skip
    assert "Valeurs sans groupe : 3 (Autre identité, 14 personnes)." in capsys.readouterr().err
    assert main(["recoder", "lister", str(path)]) == 0
    assert "R2 — age_classe (ordinal) : Âge, en trois classes." in capsys.readouterr().out
    assert main(["recoder", "retirer", str(path), "R2"]) == 0
    assert [r.id for r in load_project(path).recodes.rules] == ["R1"]


def test_a_frozen_hypothesis_holds_its_recoding(base, tmp_path):
    from ddict.analysis import plan_notebook as pn
    from ddict.analysis import planning as pl

    project, table = base
    _added(project, SEXE_2)
    prepared = apply_decisions(project, table)
    data = pl.data_from_names(project, "Le stress diffère selon le sexe.", "t_test",
                              {"outcome": ["stress"], "group": ["sexe_2"]}, {})  # fmt: skip
    pl.add(project, data, prepared, None, "2026-10-18")
    project.plan.alpha, project.plan.data_origin = 0.05, "primary"
    project.plan.data_status = "before_analysis"
    pn.freeze(nb.new_notebook(project, tmp_path / "x.dd.json"), project, prepared, "2026-10-18")
    with pytest.raises(recodes.RecodeRefused) as refused:
        recode.prepared_rule(project, table, {**SEXE_2, "label": "Autre"}, "R1")
    assert refused.value.reasons == [
        "R1 : le plan est figé et H1 en dépend ; ce recodage ne change plus."
    ]


@pytest.mark.allow_hosts(["127.0.0.1"])
def test_the_page_keeps_the_rule_of_the_command(tmp_path):
    from streamlit.testing.v1 import AppTest

    from test_analysis_page import APP

    path = _project_file(tmp_path)
    app = AppTest.from_file(APP, default_timeout=60)
    app.session_state["project_path"] = str(path)
    app.session_state["mode"] = "variables"
    app.session_state["filter"] = "all"
    app.session_state["current"] = "sexe"
    app.run()
    assert not app.exception, app.exception
    for i, label in enumerate(["Femme", "Homme ou autre identité", "Homme ou autre identité"]):
        app.text_input(key=f"recode-sexe-group-{i}").set_value(label)
    app.text_input(key="recode-sexe-name").set_value("sexe_2")
    app.text_input(key="recode-sexe-label").set_value("Sexe, en deux groupes")
    app.text_input(key="recode-sexe-reason").set_value(SEXE_2["reason"]).run()
    assert not app.exception, app.exception
    assert not app.warning, [w.value for w in app.warning]
    app.button(key="recode-sexe-save").click().run()
    assert not app.exception, app.exception
    expected = _added(build_project(read_database(EXAMPLES / KEY["sav"])), SEXE_2)
    assert load_project(path).recodes == expected.recodes
    assert app.selectbox(key="recode-sexe-which").value == "R1"
    app.session_state["mode"] = "groups"
    app.run()
    assert not app.exception, app.exception
    assert any("R1 — `sexe_2`" in m.value for m in app.markdown)


def test_a_subgroup_by_the_labels_of_a_recoding(tmp_path, capsys):
    """v0.15, step 7, the AMATUS trial: "--sous-groupe enseignante eq oui" was refused, and
    the Method wrote the code ("enseignante = 2") instead of the label."""
    from ddict.analysis import prereg
    from ddict.cli import main

    path = _project_file(tmp_path)
    assert main(["recoder", "regrouper", str(path), "sexe", "--nom", "sexe_2", "--groupe", "1",
                 "Femme", "--groupe", "2,3", "Homme ou autre",
                 "--raison", "Deux groupes."]) == 0  # fmt: skip
    assert main(["plan", "ajouter", str(path), "apparie", "--enonce", "L'anxiété baisse.",
                 "--premiere", "anxiete_t1", "--seconde", "anxiete_t2",
                 "--sous-groupe", "sexe_2", "eq", "Femme"]) == 0  # fmt: skip
    project = load_project(path)
    subgroup = project.plan.hypotheses[0].subgroup
    assert subgroup.variable == VariableRef(recode="R1") and subgroup.value == 1.0
    assert prereg.subgroup_plain(project, subgroup) == "sexe_2 = Femme"
