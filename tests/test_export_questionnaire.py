"""v0.2, step 9: the codebook holds the questions, the glossary and their provenance, and
nothing else of the questionnaire."""

import zipfile
from pathlib import Path

import pytest
from openpyxl import load_workbook
from streamlit.testing.v1 import AppTest

from ddict import glossary
from ddict.cli import main
from ddict.export import render_markdown, write_xlsx
from ddict.export import wording as w
from ddict.io import read_database
from ddict.project import build_project, load_project, save_project
from ddict.project.model import TrackedValue
from ddict.questionnaire import load_questionnaire
from ddict.review import logic

NOW = "2026-10-14T18:00:00"
APP = str(Path(__file__).resolve().parent.parent / "src" / "ddict" / "review" / "app.py")
SLEEP_Q3 = (
    "Au cours du dernier mois, dans quelle mesure les énoncés suivants s'appliquent-ils à "
    "vous ? — Je me réveille plusieurs fois pendant la nuit sans pouvoir me rendormir "
    "facilement."
)


@pytest.fixture
def project(fictional_files, generator):
    """The CSV project, every link of the answer key accepted, as after a full review."""
    project = build_project(read_database(fictional_files["etude-sommeil.csv"]))
    project.questionnaire = load_questionnaire(fictional_files["questionnaire-sommeil.txt"])
    for name, (_page, rank, row) in generator.questionnaire_key().items():
        question = project.questionnaire.questions[rank - 1]
        target = question.question_id
        if row is not None:
            target += f".{'r' if question.kind == 'grid' else 'o'}{row}"
        logic.link_question(project, name, target, now=NOW)
    logic.link_question(project, "id_participant", "none", now=NOW)
    # The consent question, linked on purpose: its paragraphs must still stay out.
    logic.link_question(project, "vague", "Q1", now=NOW)
    glossary.propose_from_questionnaire(project)
    glossary.decide(project, "TCC", "accept", now=NOW)
    glossary.decide(project, "EFA", "reject", now=NOW)  # QFS stays proposed
    glossary.write(project, "ANX", "anxiété", now=NOW)
    logic.decide(project, "sommeil_q1", "value_labels", "accept", now=NOW)
    return project


def _sheet(path, title):
    rows = list(load_workbook(path)[title].iter_rows(values_only=True))
    return [dict(zip(rows[0], row, strict=True)) for row in rows[1:]]


def _variables(path):
    return {row["Nom"]: row for row in _sheet(path, "Variables")}


def _all_text(path) -> str:
    with zipfile.ZipFile(path) as archive:
        return "".join(archive.read(n).decode("utf-8") for n in archive.namelist())


def test_exact_question_text(project, tmp_path):
    path = tmp_path / "codebook.xlsx"
    write_xlsx(project, path)
    rows = _variables(path)
    assert rows["sommeil_q3"]["Question"] == "Q7.r3"  # a text file has no pages
    assert rows["sommeil_q3"]["Texte de la question"] == SLEEP_Q3
    assert rows["aide_tcc"]["Texte de la question"].endswith(
        "— Une thérapie cognitivo-comportementale (TCC)"
    )
    assert rows["id_participant"]["Question"] == "aucune question du questionnaire"
    assert rows["id_participant"]["Texte de la question"] is None
    report = render_markdown(project)
    assert f"> Question (Q7.r3) : {SLEEP_Q3}" in report


def test_pages_of_a_pdf_questionnaire(project, fictional_files, tmp_path):
    project.questionnaire = load_questionnaire(fictional_files["questionnaire-sommeil.pdf"])
    path = tmp_path / "codebook.xlsx"
    write_xlsx(project, path)
    assert _variables(path)["sommeil_q3"]["Question"] == "Q7.r3, p. 3"
    about = {r["Élément"]: r["Valeur"] for r in _sheet(path, "Informations")}
    assert about["Questionnaire"] == "questionnaire-sommeil.pdf (4 pages)"


def test_nothing_else_of_the_questionnaire(project, generator, tmp_path):
    """The consent page and the other paragraphs never reach an export (D-012)."""
    path = tmp_path / "codebook.xlsx"
    write_xlsx(project, path, kind="work")
    for text in (_all_text(path), render_markdown(project, kind="work")):
        assert generator.CANARY_CONSENT not in text
        assert generator.CANARY_CONSENT.lower() not in text
        assert "Laboratoire fictif" not in text and "555-0199" not in text
        assert "Merci de votre participation" not in text
        assert "Acceptez-vous de participer" in text  # the linked question itself


def test_proposed_links_only_in_a_work_export(project, tmp_path):
    project.variable("age").question = TrackedValue[str](
        value="Q3", origin="ai", status="proposed", ai_run="run-0001"
    )
    final, work = tmp_path / "final.xlsx", tmp_path / "work.xlsx"
    write_xlsx(project, final)
    write_xlsx(project, work, kind="work")
    assert _variables(final)["age"]["Question"] is None
    assert _variables(work)["age"]["Question"] == "Q3 [proposé]"
    assert "[rattachement proposé]" in _variables(work)["age"]["Texte de la question"]


def test_glossary_sheet(project, tmp_path):
    final, work = tmp_path / "final.xlsx", tmp_path / "work.xlsx"
    write_xlsx(project, final)
    write_xlsx(project, work, kind="work")
    assert [(r["Sigle"], r["Origine"]) for r in _sheet(final, "Glossaire")] == [
        ("TCC", "Questionnaire"),
        ("ANX", "Équipe"),
    ]
    assert "QFS" in [r["Sigle"] for r in _sheet(work, "Glossaire")]  # waiting: work only
    assert "EFA" not in [r["Sigle"] for r in _sheet(work, "Glossaire")]  # rejected
    assert "| TCC | thérapie cognitivo-comportementale |" in render_markdown(project)


def test_value_labels_say_where_they_come_from(project, tmp_path):
    final, work = tmp_path / "final.xlsx", tmp_path / "work.xlsx"
    write_xlsx(project, final)
    write_xlsx(project, work, kind="work")
    q1 = [r for r in _sheet(final, "Valeurs") if r["Variable"] == "sommeil_q1"]
    assert q1[0]["Étiquette"] == "Pas du tout"
    assert q1[0]["Source de l'étiquette"] == w.LABELS_FROM_QUESTIONNAIRE
    assert not any(
        r["Étiquette"] for r in _sheet(final, "Valeurs") if r["Variable"] == "sommeil_q2"
    )
    proposed = [
        r for r in _sheet(work, "Valeurs") if r["Variable"] == "sommeil_q2" and r["Étiquette"]
    ]
    assert proposed and proposed[0]["Source de l'étiquette"] == w.LABELS_PROPOSED


def test_questionnaire_provenance(project, tmp_path):
    path = tmp_path / "codebook.xlsx"
    write_xlsx(project, path)
    about = {r["Élément"]: r["Valeur"] for r in _sheet(path, "Informations")}
    assert about["Questionnaire"] == "questionnaire-sommeil.txt"
    assert about["Empreinte SHA-256 du questionnaire"] == project.questionnaire.sha256
    assert about["Texte des questions"] == "exporté pour toutes les variables rattachées"


@pytest.mark.parametrize("how", ["option", "project"])
def test_protected_instrument_text_is_left_out(project, tmp_path, how):
    """D-014: the reference stays, the words of the items do not."""
    path = tmp_path / "codebook.xlsx"
    if how == "project":
        logic.set_omit_question_text(project, "G1", True)
        write_xlsx(project, path)
        report = render_markdown(project)
    else:
        write_xlsx(project, path, omit_groups=("G1",))
        report = render_markdown(project, omit_groups=("G1",))
    rows = _variables(path)
    assert rows["sommeil_q3"]["Question"] == "Q7.r3"
    assert rows["sommeil_q3"]["Texte de la question"] == w.OMITTED_TEXT
    assert rows["anx1"]["Texte de la question"].endswith("Je me suis senti.e inquiet.e.")
    for text in (_all_text(path), report):
        assert "Je me réveille plusieurs fois" not in text
    about = {r["Élément"]: r["Valeur"] for r in _sheet(path, "Informations")}
    assert "G1" in about["Texte des questions"]


def test_export_command_option(project, tmp_path, capsys):
    path = tmp_path / "etude.dd.json"
    save_project(project, path)
    assert main(["exporter", str(path), "--sans-texte", "G1", "G2"]) == 0
    rows = _variables(tmp_path / "etude.xlsx")
    assert rows["anx1"]["Texte de la question"] == w.OMITTED_TEXT
    assert main(["exporter", str(path), "--sans-texte", "G9"]) == 2
    assert "Groupe inconnu : G9" in capsys.readouterr().err


@pytest.mark.allow_hosts(["127.0.0.1"])
def test_omit_choice_in_the_interface(project, tmp_path):
    path = tmp_path / "etude.dd.json"
    save_project(project, path)
    test = AppTest.from_file(APP, default_timeout=30)
    test.session_state["project_path"] = str(path)
    test.run()
    test.session_state["mode"] = "groups"  # v0.16: the pages are a navigation, no longer a radio
    test.run()
    test.checkbox(key="omit-text-G1").check()
    test.button(key="save-omit-text-G1").click().run()
    assert not test.exception, test.exception
    assert load_project(path).groups[0].omit_question_text is True
