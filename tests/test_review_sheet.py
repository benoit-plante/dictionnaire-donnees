"""v0.5, step 8: the review spreadsheet of the English translation, and back (D-034)."""

import importlib.util
import shutil
import sys
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent


def _scripts():
    sys.path.insert(0, str(ROOT / "scripts"))
    spec = importlib.util.spec_from_file_location(
        "tableur_relecture", ROOT / "scripts" / "tableur_relecture.py"
    )
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module, sys.modules["traductions"]


def test_round_trip(tmp_path, monkeypatch):
    sheet_module, traductions = _scripts()
    locale = tmp_path / "locale"
    shutil.copytree(traductions.LOCALE, locale)
    monkeypatch.setattr(traductions, "LOCALE", locale)  # the repository's catalog is not touched
    path = tmp_path / "relecture.xlsx"
    sheet_module.write(path)
    book = load_workbook(path)
    assert book.sheetnames == ["Instructions", "Strings", "Glossary"]
    rows = list(book["Strings"].iter_rows(values_only=True))
    assert rows[0] == (
        "#",
        "Where",
        "French (source)",
        "English (proposed)",
        "Correction",
        "Comment",
    )
    assert rows[1][1].startswith("Home page")  # in the order of the screens
    assert all(r[1] for r in rows[1:])  # every string says where it appears
    found = {r[2]: i for i, r in enumerate(rows) if i}
    sheet = book["Strings"]
    sheet.cell(row=found["Nouveau projet"] + 1, column=5, value="Start a project")
    sheet.cell(
        row=found["Projet créé : {path}"] + 1, column=5, value="Project created"
    )  # {path} lost
    book.save(path)
    sheet_module.read(path)
    catalog = traductions.read("en")
    assert catalog["Nouveau projet"].string == "Start a project"
    assert catalog["Projet créé : {path}"].string == "Project created: {path}"  # refused
    assert traductions.read("en") is not None
    assert (ROOT / "src" / "ddict" / "locale" / "en" / "LC_MESSAGES" / "ddict.po").read_text(
        encoding="utf-8"
    ).count('msgstr "New project"') == 1
