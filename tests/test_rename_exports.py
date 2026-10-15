"""v0.7, step 8: the accepted new names in the exports (D-048).

The codebook shows the new name and the original one, the command files rename first, the
documented copy is written with the new names. Only accepted names apply; a work export
shows the proposals, marked. Checked as in v0.4: the R script and the SPSS syntax (in PSPP)
give back the documented copy of the tool.
"""

import json
import shutil
import subprocess
from pathlib import Path

import pyreadstat
import pytest
from openpyxl import load_workbook

from ddict import naming
from ddict.export import render_markdown, write_xlsx
from ddict.export.documented import CopyError, differences, write_documented_copy
from ddict.export.syntax import render_syntax, write_syntax
from ddict.io import read_database
from ddict.project import build_project
from test_syntax import _pspp, _rscript

EXAMPLES = Path(__file__).resolve().parent.parent / "exemples" / "bases-fictives"
KEY = json.loads((EXAMPLES / "corrige-surveymonkey.json").read_text(encoding="utf-8"))
NOW = "2026-10-15T12:00:00"
TODAY = "2026-10-15"
BASE = "travail-t2_surveymonkey.sav"


@pytest.fixture
def folder(tmp_path):
    shutil.copy(EXAMPLES / BASE, tmp_path / BASE)
    return tmp_path


@pytest.fixture
def project(folder):
    """Wave 2, its names proposed with the prefixes and stems of the answer key; all
    accepted but `q0001` (left proposed) and `q0006` (corrected)."""
    project = build_project(read_database(folder / BASE))
    naming.set_wave(project, "t2")
    expected = KEY["waves"]["t2"]
    for group in project.groups:
        prefix = expected["prefixes"].get(group.members.value[0][:5])
        if prefix:
            naming.set_prefix(project, group.group_id, prefix, now=NOW)
    for key, stem in expected["stems"].items():
        naming.decide_stem(project, key, "correct", value=stem, now=NOW)
    naming.propose(project)
    for entry in project.variables:
        if entry.name == "q0006":
            naming.decide(project, "q0006", "correct", value="formation_t2", now=NOW)
        elif entry.name != "q0001":
            naming.decide(project, entry.name, "accept", now=NOW)
    return project


def _sheet(path: Path, title: str) -> list[tuple]:
    book = load_workbook(path, read_only=True)
    rows = list(book[title].iter_rows(values_only=True))
    book.close()
    return rows


# --- The codebook. ---


def test_the_codebook_shows_the_new_and_the_original_names(project, tmp_path):
    path = tmp_path / "codebook.xlsx"
    write_xlsx(project, path)
    rows = _sheet(path, "Variables")
    assert rows[0][:3] == ("Position", "Nom", "Nom d'origine")
    names = {row[2]: row[1] for row in rows[1:]}
    assert names["q0003_0001"] == "eft_1_t2"
    assert names["q0006"] == "formation_t2"
    assert names["q0001"] == "q0001"  # only proposed: not applied in a final export
    correspondence = _sheet(path, "Correspondance")
    assert correspondence[0] == ("Nom d'origine", "Nouveau nom", "Statut", "Origine")
    assert ("q0006", "formation_t2", "corrigé", "règle de nommage") in correspondence
    assert not any(row[0] == "q0001" for row in correspondence)
    # The other sheets name the variables as the data now does.
    assert any(row[0] == "eft_1_t2" for row in _sheet(path, "Valeurs")[1:])
    groups = _sheet(path, "Groupes")
    assert groups[1][2].startswith("eft_1_t2, eft_2_t2")


def test_the_work_codebook_shows_proposals_marked(project, tmp_path):
    path = tmp_path / "travail.xlsx"
    write_xlsx(project, path, kind="work")
    names = {row[2]: row[1] for row in _sheet(path, "Variables")[1:]}
    assert names["q0001"] == "secteur_t2 [proposé]"
    statuses = {row[0]: row[2] for row in _sheet(path, "Correspondance")[1:]}
    assert statuses["q0001"] == "proposé"  # by the rules, not "proposé par IA"


def test_without_new_names_the_codebook_does_not_change(folder, tmp_path):
    project = build_project(read_database(folder / BASE))
    naming.propose(project)  # proposals only
    path = tmp_path / "codebook.xlsx"
    write_xlsx(project, path)
    assert "Nom d'origine" not in _sheet(path, "Variables")[0]
    assert "Correspondance" not in load_workbook(path, read_only=True).sheetnames


def test_the_markdown_report(project):
    text = render_markdown(project)
    assert "### `eft_1_t2`" in text
    assert "Nom d'origine : `q0003_0001`" in text
    assert "## Correspondance des noms" in text
    assert "| `q0006` | `formation_t2` | corrigé | règle de nommage |" in text
    assert "### `q0001`" in text  # not renamed in a final export
    assert "### `secteur_t2 [proposé]`" in render_markdown(project, kind="work")


# --- The command files. ---


def test_the_spss_syntax_renames_first(project):
    text = render_syntax(project, "spss", TODAY)
    start = text.index("RENAME VARIABLES")
    assert start < text.index("VARIABLE LABELS")
    assert "  (q0003_0001 = eft_1_t2)\n" in text
    assert "(q0001 = " not in text  # a proposal is not applied
    assert "VARIABLE LABELS eft_1_t2 " in text
    assert "VARIABLE LABELS q0003_0001 " not in text
    assert "VARIABLE LABELS q0001 " in text


def test_the_r_script_renames_first(project):
    text = render_syntax(project, "r", TODAY)
    assert text.index("ddict_noms <- c(") < text.index("labelled_spss(")
    assert '  "q0003_0001" = "eft_1_t2",\n' in text
    assert "names(donnees)[match(names(ddict_noms), names(donnees))] <- ddict_noms" in text
    assert 'donnees[["eft_1_t2"]] <- labelled_spss(' in text


# --- The documented copy, and the command files that give it back. ---


def test_the_documented_copy_has_the_new_names(project, folder):
    result = write_documented_copy(project, folder / BASE)
    copy, meta = pyreadstat.read_sav(str(result.path), user_missing=True)
    original, _meta = pyreadstat.read_sav(str(folder / BASE), user_missing=True)
    assert list(copy.columns) == [naming.final_name(e) for e in project.variables]
    assert "eft_1_t2" in copy.columns and "q0001" in copy.columns
    assert (
        differences(
            original.rename(columns=dict(zip(original.columns, copy.columns, strict=True))), copy
        )
        == 0
    )
    assert (
        meta.column_names_to_labels["eft_1_t2"]
        == "J'ai eu trop de tâches à accomplir en peu de temps."
    )


def test_two_equal_names_stop_the_copy(project, folder):
    project.variable("q0007").new_name = project.variable("q0006").new_name
    with pytest.raises(CopyError) as error:
        write_documented_copy(project, folder / BASE)
    assert error.value.reason == "name_conflict"


def _metadata(path: Path):
    _table, meta = pyreadstat.read_sav(str(path), user_missing=True, metadataonly=True)
    return (
        meta.column_names,
        meta.column_names_to_labels,
        meta.variable_value_labels,
        meta.missing_ranges,
    )


@pytest.mark.skipif(_rscript() is None, reason="R is not installed")
def test_the_r_script_gives_back_the_documented_copy(project, folder, tmp_path):
    copy = write_documented_copy(project, folder / BASE).path
    write_syntax(project, {"r": tmp_path / "m.R"}, TODAY)
    runner = tmp_path / "run.R"
    runner.write_text(
        "args <- commandArgs(trailingOnly = TRUE)\n"
        "donnees <- haven::read_sav(args[1], user_na = TRUE)\n"
        'source(args[2], encoding = "UTF-8")\n'
        "haven::write_sav(donnees, args[3])\n",
        encoding="utf-8",
    )
    out = tmp_path / "out.sav"
    done = subprocess.run(
        [_rscript(), str(runner), str(folder / BASE), str(tmp_path / "m.R"), str(out)],
        capture_output=True,
        text=True,
        encoding="utf-8",
        errors="replace",
    )
    if "there is no package called" in done.stderr:
        pytest.skip("haven is not installed in R")
    assert done.returncode == 0, done.stderr
    assert _metadata(out) == _metadata(copy)
    raw = {"user_missing": True}
    a, _ = pyreadstat.read_sav(str(copy), **raw)
    b, _ = pyreadstat.read_sav(str(out), **raw)
    assert differences(a, b) == 0


@pytest.mark.skipif(_pspp() is None, reason="PSPP is not installed (D-027)")
def test_the_spss_syntax_gives_back_the_documented_copy_in_pspp(project, folder, tmp_path):
    copy = write_documented_copy(project, folder / BASE).path
    data, _meta = pyreadstat.read_sav(str(folder / BASE), user_missing=True)
    bare = tmp_path / "bare.sav"
    pyreadstat.write_sav(data, str(bare))  # the same data, without the metadata
    write_syntax(project, {"spss": tmp_path / "m.sps"}, TODAY)
    out = tmp_path / "out.sav"
    job = tmp_path / "job.sps"
    job.write_text(
        f"GET FILE='{bare}'.\nINSERT FILE='{tmp_path / 'm.sps'}' ENCODING='UTF-8'.\n"
        f"SAVE OUTFILE='{out}'.\n",
        encoding="utf-8",
    )
    done = subprocess.run([_pspp(), str(job)], capture_output=True, text=True, errors="replace")
    said = (done.stdout + done.stderr).lower()
    assert done.returncode == 0 and "error" not in said and "warning" not in said, said
    assert _metadata(out) == _metadata(copy)


def test_the_export_command_writes_the_table_of_names(project, folder, capsys):
    from ddict.cli import main
    from ddict.project import save_project
    from ddict.waves import read_name_table

    path = folder / "travail-t2.dd.json"
    save_project(project, path)
    assert main(["exporter", str(path)]) == 0
    table = folder / "travail-t2-correspondance.csv"
    assert "Table de correspondance des noms" in capsys.readouterr().out
    rows = dict(read_name_table(table))
    assert rows["q0003_0001"] == "eft_1_t2" and rows["q0001"] == "q0001"
