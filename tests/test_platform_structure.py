"""v0.7, step 4: groups and links from the structure of a platform export (D-044), and the
header of its Excel export as the questionnaire (D-045)."""

import json
from pathlib import Path

import pyreadstat
import pytest
from openpyxl import Workbook

from ddict import i18n
from ddict.export.wording import group_description
from ddict.io import read_database
from ddict.project import build_project
from ddict.questionnaire import link, load_questionnaire
from ddict.questionnaire.link import texts_agree
from ddict.workflow import add_questionnaire

EXAMPLES = Path(__file__).resolve().parent.parent / "exemples" / "bases-fictives"
KEY = json.loads((EXAMPLES / "corrige-surveymonkey.json").read_text(encoding="utf-8"))
HEADER = EXAMPLES / "travail-t2_surveymonkey.xlsx"


def _project(file_name: str, questionnaire: Path | None = None):
    project = build_project(read_database(EXAMPLES / file_name))
    if questionnaire is not None:
        add_questionnaire(project, questionnaire)
    return project


def _groups(project) -> list[dict]:
    return [
        {"kind": g.kind, "members": g.members.value}
        | ({"attached": g.attached} if g.kind == "multiple_response" else {})
        for g in project.groups
    ]


# --- Groups. ---


@pytest.mark.parametrize("wave", ["t1", "t2"])
def test_the_groups_are_those_of_the_survey(wave):
    project = _project(KEY["waves"][wave]["sav"])
    assert _groups(project) == KEY["waves"][wave]["groups"]


@pytest.mark.parametrize(
    "file_name", ["travail-t2_surveymonkey.xlsx", "travail-t2_surveymonkey.csv"]
)
def test_the_excel_and_csv_exports_give_the_same_groups(file_name):
    assert _groups(_project(file_name)) == KEY["waves"]["t2"]["groups"]


def test_a_set_of_checkboxes_has_no_item_statistics():
    project = _project("travail-t2_surveymonkey.sav")
    group = next(g for g in project.groups if g.kind == "multiple_response")
    assert group.cronbach_alpha is None and group.basis == ["platform_structure"]
    assert group.descriptive_label.value == "5 options à cocher, préfixe « q0005_ »"
    with i18n.using("en"):
        assert group_description(group, "final") == '5 checkbox options, prefix "q0005_"'
    member = project.variable("q0005_0001")
    assert member.group_id.value == group.group_id
    assert project.variable("q0005_other").group_id.value is None  # attached, not a member
    grid = project.groups[0]
    assert grid.kind == "scale" and grid.cronbach_alpha is not None
    assert grid.descriptive_label.value == "12 items, codes de 1 à 5, préfixe « q0003_ »"


def test_a_grid_with_reversed_items_shows_them():
    project = _project("travail-t2_surveymonkey.sav")
    suspects = {e.name for e in project.variables if e.profile.reverse_suspect}
    # Items 2, 7 and 11 of the fictional grid describe a satisfying situation.
    assert suspects == {f"q0003_{k:04d}" for k in (2, 7, 11)}


# --- The header of the Excel export as the questionnaire. ---


def test_the_header_is_read_as_the_questionnaire():
    questionnaire = load_questionnaire(HEADER)
    assert questionnaire.format == "xlsx" and questionnaire.n_pages is None
    expected = [(q["kind"], q["n_rows"]) for q in KEY["waves"]["t2"]["questions"]]
    found = [(q.kind, len(q.rows)) for q in questionnaire.questions]
    # The questions of the paper questionnaire, then the column the platform adds ("id").
    assert found == [*expected, ("open", 0)]
    grid = questionnaire.questions[2]
    assert grid.question_id == "Q3" and grid.rows[3].row_id == "Q3.r4"
    assert grid.rows[0].text == "J'ai eu trop de tâches à accomplir en peu de temps."
    tasks = questionnaire.questions[4]
    assert [o.option_id for o in tasks.options][-1] == "Q5.o6"
    assert tasks.options[-1].text == "Other (please specify)"
    assert questionnaire.texts == []  # nothing but the header: no consent page


def test_the_header_reads_no_answer(tmp_path):
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Respondent ID", "Votre couleur ?", "Votre ville ?"])
    sheet.append([None, "Response", "Open-Ended Response"])
    sheet.append([1, "Bleu", "ZQXREPONSE"])
    path = tmp_path / "export.xlsx"
    workbook.save(path)
    questionnaire = load_questionnaire(path)
    assert "ZQXREPONSE" not in questionnaire.model_dump_json()
    assert [(q.kind, q.text) for q in questionnaire.questions] == [
        ("single", "Votre couleur ?"),
        ("open", "Votre ville ?"),
    ]


def test_a_choice_question_keeps_its_other_option(tmp_path):
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Respondent ID", "Votre secteur ?", None])
    sheet.append([None, "Response", "Autre (veuillez préciser)"])
    path = tmp_path / "export.xlsx"
    workbook.save(path)
    (question,) = load_questionnaire(path).questions
    assert question.kind == "single"
    assert [(o.option_id, o.text) for o in question.options] == [
        ("Q1.o1", "Autre (veuillez préciser)")
    ]


def test_the_kind_of_a_question_comes_from_the_database_when_known(tmp_path):
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Respondent ID", "Une question", None, None])
    sheet.append([None, "A", "B", "C"])
    path = tmp_path / "export.xlsx"
    workbook.save(path)
    assert load_questionnaire(path).questions[0].kind == "grid"  # no other clue
    assert load_questionnaire(path, {1: "checkbox"}).questions[0].kind == "checkbox"


def test_an_ordinary_workbook_is_not_a_questionnaire(tmp_path):
    from ddict.io import ReadError

    with pytest.raises(ReadError):
        load_questionnaire(EXAMPLES / "etude-sommeil.xlsx")


# --- Links. ---


def test_every_variable_is_linked_by_the_structure():
    project = _project("travail-t2_surveymonkey.sav", HEADER)
    assert link.platform_choices(project) == KEY["waves"]["t2"]["links"]
    assert link.sure_choices(project) == KEY["waves"]["t2"]["links"]
    assert link.best_guess(project) == KEY["waves"]["t2"]["links"]


def test_the_excel_export_gets_its_labels_from_its_header():
    excel = read_database(EXAMPLES / "travail-t2_surveymonkey.xlsx")
    _df, meta = pyreadstat.read_sav(EXAMPLES / "travail-t2_surveymonkey.sav", metadataonly=True)
    labels = meta.column_names_to_labels
    for variable in excel.variables:
        name = variable.name
        if name in ("q0003_0001", "q0004_0002", "q0001", "q0009", "q0005_other"):
            assert variable.label == labels[name], name
    assert excel.variable("q0005_0001").label == "Accueil de la clientèle"  # the option
    assert excel.variable("respondent_id").label is None


@pytest.mark.parametrize(
    "file_name", ["travail-t2_surveymonkey.xlsx", "travail-t2_surveymonkey.csv"]
)
def test_an_excel_or_csv_base_is_linked_by_the_structure(file_name):
    project = _project(file_name, HEADER)
    expected = {k: v for k, v in KEY["waves"]["t2"]["links"].items() if k != "CollectorNm"}
    assert link.platform_choices(project) == expected


def test_with_a_pdf_every_structural_link_has_agreeing_texts():
    project = _project("travail-t2_surveymonkey.sav", EXAMPLES / "questionnaire-travail.pdf")
    choices = link.platform_choices(project)
    for name, target_id in choices.items():
        if target_id == link.NO_QUESTION:
            continue
        found = link.target(project.questionnaire, target_id)
        entry = project.variable(name)
        texts = [entry.file.label, *(v.label for v in entry.file.value_labels)]
        if name.endswith("_other"):  # an "other" option, in any language
            texts = ["Other (please specify)", "Autre (veuillez préciser)"]
        assert any(texts_agree(t, found.part or found.question.text) for t in texts), name


def test_wave_1_is_linked_to_its_own_numbering():
    # The questionnaire of wave 2 has a question wave 1 did not: q0002 of wave 1 is Q3.
    project = _project("travail-t1_surveymonkey.sav", HEADER)
    choices = link.platform_choices(project)
    assert choices["q0001"] == "Q1"
    assert choices["q0002_0001"] == "Q3.r1"
    assert "q0002_0010" not in choices  # the item wave 2 removed
    assert choices["q0002_0011"] == "Q3.r10"
    assert choices["q0006"] == "Q7"  # the changed labels do not matter: the question agrees


@pytest.mark.parametrize(
    ("first", "second", "agree"),
    [
        ("L'éclairage", "l'eclairage", True),
        ("…je me sens fatigué", "Je me sens fatigué.", True),
        ("Un texte assez long pour être coupé par SPSS à la fin", "Un texte assez long pour", True),
        ("Une option coupée au milieu d'un mo", "Une option coupée au milieu d'un mot", True),
        ("Oui", "Non", False),
        ("Le", "Le bruit", False),  # too short to be a cut text
        ("Jamais", None, False),
    ],
)  # fmt: skip
def test_texts_agree(first, second, agree):
    assert texts_agree(first, second) is agree


# --- In the interface. ---


@pytest.mark.allow_hosts(["127.0.0.1"])
def test_the_header_questionnaire_and_the_groups_in_the_interface(tmp_path):
    from streamlit.testing.v1 import AppTest

    from ddict.workflow import profile_database

    for name in ("travail-t2_surveymonkey.sav", "travail-t2_surveymonkey.xlsx"):
        (tmp_path / name).write_bytes((EXAMPLES / name).read_bytes())
    _created, path = profile_database(tmp_path / "travail-t2_surveymonkey.sav")
    app_path = Path(__file__).resolve().parent.parent / "src" / "ddict" / "review" / "app.py"
    app = AppTest.from_file(str(app_path), default_timeout=30)
    app.session_state["project_path"] = str(path)
    app.session_state["mode"] = "questionnaire"
    app.run()
    app.text_input(key="questionnaire-path").input(str(tmp_path / "travail-t2_surveymonkey.xlsx"))
    app.button(key="questionnaire-read").click().run()
    assert not app.exception, app.exception
    assert any("12 questions" in m.value for m in app.markdown)

    app.session_state["mode"] = "groups"
    app.run()
    app.sidebar.selectbox(key="group").select("G3").run()
    assert not app.exception, app.exception
    texts = [m.value for m in app.markdown]
    assert any("une variable par option cochée" in t for t in texts)
    assert any("q0005_other" in t for t in texts)
    assert not app.multiselect  # no reversed items for checkboxes
