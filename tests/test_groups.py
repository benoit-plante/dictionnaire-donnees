"""Step 8: probable scales and reversed items, found without any model."""

import numpy as np
import pandas as pd
import pytest
from openpyxl import load_workbook

from ddict.export import render_markdown, write_xlsx
from ddict.groups import find_groups, item_prefix
from ddict.io import LoadedData, SourceFile, VariableMetadata, read_csv, read_sav
from ddict.privacy import classify_database
from ddict.profile import profile_database
from ddict.project import build_project


def _groups(data):
    profiles = profile_database(data)
    return find_groups(data, profiles, classify_database(data, profiles))


@pytest.fixture(scope="module", params=["etude-sommeil.sav", "etude-sommeil_excel-fr.csv"])
def groups(request, fictional_files):
    path = fictional_files[request.param]
    data = read_sav(path) if path.suffix == ".sav" else read_csv(path)
    return {g.group_id: g for g in _groups(data)}


def test_both_fictional_scales_are_found(groups, generator):
    assert {g.members for g in groups.values()} == {
        tuple(generator.SLEEP_ITEMS),
        tuple(generator.ANXIETY_ITEMS),
    }
    assert groups["G1"].members[0] == "sommeil_q1"


def test_non_contiguous_scale_is_found_by_its_prefix(groups):
    anxiety = groups["G2"]
    assert anxiety.prefix == "anx"
    assert "common_prefix" in anxiety.basis
    assert "contiguous" not in anxiety.basis
    assert "contiguous" in groups["G1"].basis


def test_reversed_items(groups, generator):
    suspects = {i.name for g in groups.values() for i in g.items if i.reverse_suspect}
    assert suspects == set(generator.TRAPS["reversed_items"])
    for group in groups.values():
        for item in group.items:
            assert item.item_rest_r is not None
            assert (item.item_rest_r < 0) == item.reverse_suspect


def test_alpha_rises_when_suspect_items_are_reversed(groups):
    sleep, anxiety = groups["G1"], groups["G2"]
    assert sleep.cronbach_alpha_reversed > sleep.cronbach_alpha
    assert sleep.cronbach_alpha_reversed > 0.7
    assert anxiety.cronbach_alpha > 0.7
    assert anxiety.cronbach_alpha_reversed is None


def test_missing_codes_are_left_out_of_the_statistics(groups):
    """Declared -9 and the unlabeled 9 would distort the codes and the correlations."""
    assert (groups["G1"].code_min, groups["G1"].code_max) == (1.0, 5.0)
    assert (groups["G2"].code_min, groups["G2"].code_max) == (0.0, 3.0)
    assert groups["G1"].n_complete_rows < 200


def test_alpha_matches_a_direct_computation(fictional_files, generator):
    data = read_sav(fictional_files["etude-sommeil.sav"])
    items = data.table[generator.ANXIETY_ITEMS].replace(-9, np.nan).dropna()
    k = items.shape[1]
    expected = k / (k - 1) * (1 - items.var(ddof=1).sum() / items.sum(axis=1).var(ddof=1))
    (anxiety,) = [g for g in _groups(data) if g.prefix == "anx"]
    assert anxiety.cronbach_alpha == pytest.approx(expected, abs=1e-3)
    assert anxiety.n_complete_rows == len(items)


@pytest.mark.parametrize(
    ("name", "prefix"),
    [
        ("sommeil_q1", "sommeil_q"),
        ("anx12", "anx"),
        ("bdi.3", "bdi"),
        ("item-4", "item"),
        ("heures_sommeil", None),
        ("q_retiree", None),
        ("42", None),
    ],
)
def test_item_prefix(name, prefix):
    assert item_prefix(name) == prefix


def _data(columns, labels=None):
    table = pd.DataFrame(columns)
    variables = tuple(
        VariableMetadata(
            name=name,
            position=i,
            storage_type="numeric",
            value_labels=(labels or {}).get(name, {}),
        )
        for i, name in enumerate(table.columns, start=1)
    )
    source = SourceFile(name="t.sav", format="sav", sha256="0", n_rows=len(table), n_columns=0)
    return LoadedData(source=source, variables=variables, table=table)


def _scale(rng, n=50):
    latent = rng.normal(size=n)
    return lambda: np.clip(np.round(3 + latent + rng.normal(scale=0.7, size=n)), 1, 5)


def test_neighbours_with_the_same_labels_form_a_group_without_prefix():
    rng = np.random.default_rng(1)
    item = _scale(rng)
    likert = {1: "Jamais", 2: "Rarement", 3: "Parfois", 4: "Souvent", 5: "Toujours"}
    columns = {"fatigue": item(), "humeur": item(), "appetit": item(), "age": rng.normal(size=50)}
    data = _data(columns, labels={name: likert for name in ["fatigue", "humeur", "appetit"]})
    (group,) = _groups(data)
    assert group.members == ("fatigue", "humeur", "appetit")
    assert group.basis == ("shared_value_labels", "contiguous")
    assert group.prefix is None


def test_same_labels_far_apart_do_not_make_a_group():
    rng = np.random.default_rng(2)
    item = _scale(rng)
    yes_no = {1: "Oui", 2: "Non"}
    columns = {
        "a": item(),
        "x": rng.normal(size=50),
        "b": item(),
        "y": rng.normal(size=50),
        "c": item(),
    }
    data = _data(columns, labels={name: yes_no for name in ["a", "b", "c"]})
    assert _groups(data) == ()


def test_two_items_are_not_a_group():
    rng = np.random.default_rng(3)
    item = _scale(rng)
    assert _groups(_data({"q1": item(), "q2": item()})) == ()


def test_project_and_exports_show_the_groups(fictional_files, tmp_path):
    project = build_project(read_sav(fictional_files["etude-sommeil.sav"]))
    assert [g.group_id for g in project.groups] == ["G1", "G2"]
    sleep = project.groups[0]
    assert sleep.descriptive_label.value == "8 items, codes de 1 à 5, préfixe « sommeil_q »"
    assert sleep.members.origin == "computed"
    assert sleep.instrument_name is None
    q3 = project.variable("sommeil_q3")
    assert q3.group_id.value == "G1" and q3.profile.reverse_suspect
    assert project.variable("age").group_id.value is None

    report = render_markdown(project)
    assert "## Groupes de variables" in report
    assert "Items probablement inversés : `sommeil_q3`, `sommeil_q6`" in report
    assert "Item probablement inversé (corrélation avec le reste de l'échelle" in report

    path = tmp_path / "codebook.xlsx"
    write_xlsx(project, path)
    rows = list(load_workbook(path)["Groupes"].iter_rows(values_only=True))
    header, sleep_row = rows[0], dict(zip(rows[0], rows[1], strict=True))
    assert len(rows) == 3
    assert sleep_row["Items inversés probables"] == "sommeil_q3, sommeil_q6"
    assert sleep_row["Instrument (source)"] is None
    assert "Alpha de Cronbach observé" in header
