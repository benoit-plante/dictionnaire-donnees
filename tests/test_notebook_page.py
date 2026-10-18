"""v0.11, step 6: the Notebook page, the "Ajouter au carnet" button of the Analyses page,
`ddict carnet` and `ddict resultats`; the same files from the page and from the command
line."""

import io
import json
import shutil
import zipfile
from pathlib import Path

import pytest
from docx import Document
from streamlit.testing.v1 import AppTest

from ddict.analysis import notebook as nb
from ddict.cli import main
from ddict.io import read_database
from ddict.project import build_project, load_project, save_project
from ddict.review import logic
from test_analysis_data import EXAMPLES, KEY, analyses_project

APP = str(Path(__file__).resolve().parent.parent / "src" / "ddict" / "review" / "app.py")
NOW = "2026-10-16T12:00:00"


def _folder(tmp_path, name="projet"):
    folder = tmp_path / name
    folder.mkdir()
    shutil.copy(EXAMPLES / KEY["sav"], folder / KEY["sav"])
    path = folder / "etude-analyses.dd.json"
    save_project(analyses_project(), path)
    return path


def _sleep(tmp_path):
    """The sleep study: its code 999 of age is suspect, not declared in the file."""
    folder = tmp_path / "sommeil"
    folder.mkdir()
    shutil.copy(EXAMPLES / "etude-sommeil.sav", folder / "etude-sommeil.sav")
    path = folder / "etude-sommeil.dd.json"
    save_project(build_project(read_database(folder / "etude-sommeil.sav")), path)
    return path


def _confirm_999(path):
    project = load_project(path)
    logic.decide_sentinel(project, "age", 999, True, now=NOW)
    save_project(project, path)


@pytest.fixture
def project_path(tmp_path):
    return _folder(tmp_path)


def _build_by_command(path):
    """The notebook of the tests, by the command line."""
    for arguments in (
        ["test_t", "--dependante", "bienetre_total", "--groupe", "groupe"],
        ["correlation", "--variables", "heures_travail", "stress", "soutien"],
        ["tableau1", "--variables", "age", "sexe", "--groupe", "groupe"],
    ):
        assert main(["carnet", "ajouter", str(path), *arguments]) == 0
    assert main(["carnet", "deplacer", str(path), "A3", "1"]) == 0
    assert main(["carnet", "ecarter", str(path), "A2"]) == 0


# --- The command line. ---


def test_the_commands(project_path, capsys):
    assert main(["carnet", "lister", str(project_path)]) == 0
    assert "Le carnet est vide" in capsys.readouterr().out
    assert not nb.notebook_path(project_path).exists()  # nothing written for an empty one
    _build_by_command(project_path)
    out = capsys.readouterr().out
    assert "Ajoutée au carnet : A1" in out and "t(297,29) = 5,00" in out
    assert "A3 est maintenant au rang 1." in out
    notebook = nb.load_notebook(nb.notebook_path(project_path))
    assert [(e.id, e.retained) for e in notebook.entries] == [
        ("A3", True),
        ("A1", True),
        ("A2", False),
    ]
    assert main(["notebook", "numbers", str(project_path), "--table", "3"]) == 0
    assert main(["carnet", "lister", str(project_path)]) == 0
    out = capsys.readouterr().out
    assert "1. A3 — Tableau 1 : décrire l'échantillon — retenue, à jour" in out
    assert "3. A2 — Corrélation" in out and "écartée" in out
    assert "Numérotation : premier tableau 3, première figure 1." in out
    assert main(["carnet", "retenir", str(project_path), "A2"]) == 0
    assert main(["carnet", "retirer", str(project_path), "A2"]) == 0
    assert main(["carnet", "retirer", str(project_path), "A2"]) == 2
    assert "A2 : aucune analyse du carnet ne porte ce numéro." in capsys.readouterr().err
    assert main(["carnet", "ajouter", str(project_path), "test_t", "--dependante", "stress",
                 "--groupe", "groupe"]) == 0  # fmt: skip
    notebook = nb.load_notebook(nb.notebook_path(project_path))
    assert [e.id for e in notebook.entries] == ["A3", "A1", "A4"]  # A2 never comes back
    assert notebook.first_table == 3
    assert main(["resultats", str(project_path)]) == 0
    out = capsys.readouterr().out
    for name in ("article.docx", "article-graphiques.xlsx", "article-R.zip"):
        assert project_path.with_name(f"etude-analyses-{name}").is_file()
        assert f"etude-analyses-{name}" in out
    word = Document(project_path.with_name("etude-analyses-article.docx"))
    assert next(p.text for p in word.paragraphs if p.text.startswith("Tableau ")) == "Tableau 3"


def test_the_command_refuses(project_path, capsys):
    assert main(["resultats", str(project_path)]) == 2
    assert "Aucune analyse retenue" in capsys.readouterr().err
    code = main(["carnet", "ajouter", str(project_path), "correlation", "--variables", "stress",
                 "id_participant"])  # fmt: skip
    assert code == 2 and "identifiant" in capsys.readouterr().err
    assert not nb.notebook_path(project_path).exists()
    nb.notebook_path(project_path).write_text('{"carnet_version": "2.0.0"}', encoding="utf-8")
    assert main(["carnet", "lister", str(project_path)]) == 2
    assert "version plus récente" in capsys.readouterr().err


def test_a_change_to_confirm_by_command(tmp_path, capsys):
    path = _sleep(tmp_path)
    assert main(["carnet", "ajouter", str(path), "correlation", "--variables", "age",
                 "heures_sommeil"]) == 0  # fmt: skip
    _confirm_999(path)
    assert main(["resultats", str(path)]) == 2
    assert "A1" in capsys.readouterr().err
    assert not path.with_name("etude-sommeil-article.docx").exists()
    assert main(["carnet", "lister", str(path)]) == 0
    out = capsys.readouterr().out
    assert "changée, à confirmer" in out
    assert "n_used : 200 → 197" in out and "r : ,04 → ,07" in out
    assert "Nouveau : " in out
    assert main(["carnet", "confirmer", str(path), "--tout"]) == 0
    assert "Confirmée : A1" in capsys.readouterr().out
    assert main(["carnet", "confirmer", str(path)]) == 0
    assert "Aucun changement à confirmer." in capsys.readouterr().out
    assert main(["resultats", str(path)]) == 0
    saved = json.loads(nb.notebook_path(path).read_text(encoding="utf-8"))
    assert saved["entries"][0]["computed"]["n_used"] == 197


def test_without_the_database_the_list_still_shows(project_path, capsys):
    _build_by_command(project_path)
    project_path.with_name(KEY["sav"]).unlink()
    capsys.readouterr()
    assert main(["carnet", "lister", str(project_path)]) == 0
    out = capsys.readouterr().out
    assert "non recalculée" in out and "introuvable" in out
    assert "t(297,29) = 5,00" in out  # the sentence kept
    assert main(["resultats", str(project_path)]) == 2
    assert "introuvable" in capsys.readouterr().err


# --- The page. ---


def _open(path, mode="notebook"):
    app = AppTest.from_file(APP, default_timeout=60)
    app.session_state["project_path"] = str(path)
    app.session_state["mode"] = mode
    app.run()
    assert not app.exception, app.exception
    return app


def _add_from_the_analyses_page(app, analysis, roles):
    app.selectbox(key="analysis-key").set_value(analysis).run()
    for role, value in roles.items():
        key = f"analysis-{analysis}-{role}"
        if isinstance(value, list):
            app.multiselect(key=key).set_value(value).run()
        else:
            app.selectbox(key=key).set_value(value).run()
    app.button(key="analysis-run").click().run()
    assert not app.exception, app.exception
    app.button(key="analysis-add").click().run()
    assert not app.exception, app.exception


@pytest.mark.allow_hosts(["127.0.0.1"])
def test_the_page_gives_the_command_files(project_path, tmp_path):
    app = _open(project_path)
    assert app.header[0].value == "Carnet"
    assert "Le carnet est vide" in app.info[0].value
    app.session_state["mode"] = "analyses"
    app.run()
    _add_from_the_analyses_page(app, "t_test", {"outcome": "bienetre_total", "group": "groupe"})
    assert "Ajoutée au carnet : A1" in app.success[0].value
    assert app.button(key="analysis-add").disabled  # the same run, once
    _add_from_the_analyses_page(
        app, "correlation", {"variables": ["heures_travail", "stress", "soutien"]}
    )
    _add_from_the_analyses_page(app, "table1", {"variables": ["age", "sexe"], "group": "groupe"})
    app.session_state["mode"] = "notebook"
    app.run()
    assert not app.exception, app.exception
    assert app.button(key="notebook-up-A1").disabled and app.button(key="notebook-down-A3").disabled
    app.button(key="notebook-up-A3").click().run()
    app.button(key="notebook-up-A3").click().run()
    app.toggle(key="notebook-retain-A2").set_value(False).run()
    assert not app.exception, app.exception
    notebook = nb.load_notebook(nb.notebook_path(project_path))
    assert [(e.id, e.retained) for e in notebook.entries] == [
        ("A3", True),
        ("A1", True),
        ("A2", False),
    ]
    app.button(key="notebook-write").click().run()
    assert "Fichiers écrits à côté du projet" in app.success[0].value

    other = _folder(tmp_path, "commande")
    _build_by_command(other)
    assert main(["resultats", str(other)]) == 0
    name = "etude-analyses-article-R.zip"
    assert project_path.with_name(name).read_bytes() == other.with_name(name).read_bytes()
    name = "etude-analyses-article.docx"
    page = Document(project_path.with_name(name))
    command = Document(other.with_name(name))
    assert [p.text for p in page.paragraphs] == [p.text for p in command.paragraphs]
    assert [[c.text for c in t._cells] for t in page.tables] == [
        [c.text for c in t._cells] for t in command.tables
    ]
    with zipfile.ZipFile(io.BytesIO(project_path.with_name(
        "etude-analyses-article-graphiques.xlsx").read_bytes())) as book:  # fmt: skip
        assert "xl/charts/chart1.xml" in book.namelist()


@pytest.mark.allow_hosts(["127.0.0.1"])
def test_a_title_changed_on_the_analyses_page_is_kept(project_path):
    app = _open(project_path, "analyses")
    app.selectbox(key="analysis-key").set_value("t_test").run()
    app.selectbox(key="analysis-t_test-outcome").set_value("bienetre_total").run()
    app.selectbox(key="analysis-t_test-group").set_value("groupe").run()
    app.button(key="analysis-run").click().run()
    title = next(t for t in app.text_input if t.key.endswith("-table-1-title"))
    title.set_value("Bien-être selon le groupe").run()
    number = next(n for n in app.number_input if n.key.endswith("-table-1-number"))
    number.set_value(7).run()
    app.button(key="analysis-add").click().run()
    [entry] = nb.load_notebook(nb.notebook_path(project_path)).entries
    assert entry.edits["table-1"].title == "Bien-être selon le groupe"
    assert "number" not in entry.edits["table-1"].model_dump()  # the order numbers
    app.session_state["mode"] = "notebook"
    app.run()
    app.button(key="notebook-write").click().run()
    word = Document(project_path.with_name("etude-analyses-article.docx"))
    texts = [p.text for p in word.paragraphs]
    assert texts[texts.index("Tableau 1") + 1] == "Bien-être selon le groupe"


@pytest.mark.allow_hosts(["127.0.0.1"])
def test_the_page_shows_a_change_to_confirm(tmp_path):
    path = _sleep(tmp_path)
    assert main(["carnet", "ajouter", str(path), "correlation", "--variables", "age",
                 "heures_sommeil"]) == 0  # fmt: skip
    _confirm_999(path)
    app = _open(path)
    assert "changée, à confirmer" in " ".join(m.value for m in app.markdown)
    assert any("n_used : 200 → 197" in m.value for m in app.markdown)
    assert app.button(key="notebook-write").disabled
    assert "A1" in app.warning[0].value
    app.button(key="notebook-confirm-A1").click().run()
    assert not app.exception, app.exception
    assert not app.button(key="notebook-write").disabled
    assert nb.load_notebook(nb.notebook_path(path)).entries[0].computed.n_used == 197


@pytest.mark.allow_hosts(["127.0.0.1"])
def test_the_page_without_the_database(project_path):
    _build_by_command(project_path)
    project_path.with_name(KEY["sav"]).unlink()
    app = _open(project_path)
    assert "introuvable" in app.info[0].value
    assert app.button(key="notebook-choose")
    assert "non recalculée" in " ".join(m.value for m in app.markdown)
    assert app.button(key="notebook-write").disabled
    app.button(key="notebook-remove-A2").click().run()  # still possible without the base
    assert [e.id for e in nb.load_notebook(nb.notebook_path(project_path)).entries] == [
        "A3",
        "A1",
    ]
