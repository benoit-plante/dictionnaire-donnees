"""v0.7, step 1: the fictional SurveyMonkey export and what v0.6 does with it.

The first tests check that the fictional files follow SurveyMonkey's conventions, as seen in
two real exports (docs/plan-v0.7.md, § 1). The others state, one finding of that trial at a
time, what the tool should do; they fail with v0.6 (strict xfail) and are turned into plain
tests by the step that fixes them.
"""

import hashlib
import json
from pathlib import Path

import pandas as pd
import pyreadstat
import pytest
from openpyxl import load_workbook

from conftest import load_counter
from ddict.questionnaire import load_questionnaire

ROOT = Path(__file__).resolve().parent.parent
EXAMPLES = ROOT / "exemples" / "bases-fictives"
FILES = [
    "travail-t1_surveymonkey.sav",
    "travail-t2_surveymonkey.sav",
    "travail-t1-documente.sav",
    "travail-t1-correspondance.csv",
    "travail-t1_surveymonkey.xlsx",
    "travail-t2_surveymonkey.xlsx",
    "travail-t2_surveymonkey.csv",
    "questionnaire-travail.txt",
    "questionnaire-travail.pdf",
    "corrige-surveymonkey.json",
]


def _sha(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


@pytest.fixture(scope="module")
def key():
    return json.loads((EXAMPLES / "corrige-surveymonkey.json").read_text(encoding="utf-8"))


def test_committed_files_match_the_generator(surveymonkey_files):
    assert sorted(surveymonkey_files) == sorted(FILES)
    for name in FILES:
        assert _sha(surveymonkey_files[name]) == _sha(EXAMPLES / name), (
            f"{name} differs: run uv run python scripts/generer_base_surveymonkey.py"
        )


def test_the_key_covers_every_variable(key):
    for wave in ("t1", "t2"):
        expected = key["waves"][wave]
        _df, meta = pyreadstat.read_sav(EXAMPLES / expected["sav"], metadataonly=True)
        names = meta.column_names
        for part in ("structure", "links", "names"):
            covered = set(expected[part]) | (
                set(expected["platform_variables"]) if part == "structure" else set()
            )
            assert covered == set(names), (wave, part)
        assert len(set(expected["names"].values())) == len(names)
    between = key["between_waves"]
    assert between["added"] == ["q0002"]
    assert between["removed"] == ["q0002_0010"]
    assert between["value_labels_changed"] == ["q0007"]
    assert len(between["pairs_t2_to_t1"]) == 52 - 1  # all but the added question


def test_the_sav_files_follow_surveymonkey_conventions(key):
    _t1, meta1 = pyreadstat.read_sav(EXAMPLES / "travail-t1_surveymonkey.sav", metadataonly=True)
    _t2, meta2 = pyreadstat.read_sav(
        EXAMPLES / "travail-t2_surveymonkey.sav", metadataonly=True, user_missing=True
    )
    # Platform variables first, in their old and recent names.
    assert meta1.column_names[:4] == ["CollectorNm", "RespondentID", "CollectorID", "StartDate"]
    assert meta2.column_names[:4] == [
        "CollectorNm",
        "respondent_id",
        "collector_id",
        "date_created",
    ]
    types = meta2.original_variable_types
    assert types["date_created"] == "A100"
    assert meta1.original_variable_types["StartDate"] == "A22"
    # Codes in F8.2, one value label set per labelled variable.
    assert types["q0001"] == "F8.2" and types["q0003_0001"] == "F8.2"
    assert len(meta2.value_labels) == len(meta2.variable_value_labels)
    # Likert items at the scale level; a checkbox has one value label, its question as label.
    assert meta2.variable_measure["q0003_0001"] == "scale"
    assert meta2.variable_value_labels["q0005_0001"] == {1.0: "Accueil de la clientèle"}
    labels = meta2.column_names_to_labels
    assert labels["q0005_0001"] == labels["q0005_0002"]
    assert labels["q0005_other"] == "Other (please specify)"
    assert "identifier" in labels["q0003_0005"]
    # "Ne sais pas" and "Préfère ne pas répondre" after the scale, not declared missing.
    assert meta2.variable_value_labels["q0004_0001"][6.0] == "Ne sais pas"
    assert not meta2.missing_ranges
    assert meta2.column_names[-1] == "q0012" and labels["q0012"] == "id"


def test_the_survey_dates_are_text(key):
    table, _meta = pyreadstat.read_sav(EXAMPLES / "travail-t2_surveymonkey.sav")
    first = table["date_created"].iloc[0]
    assert isinstance(first, str)
    assert pd.to_datetime(table["date_created"], format="%m/%d/%Y %I:%M:%S %p").notna().all()


def test_the_excel_export_has_two_header_rows():
    workbook = load_workbook(EXAMPLES / "travail-t2_surveymonkey.xlsx", read_only=True)
    rows = list(workbook.worksheets[0].iter_rows(values_only=True))
    workbook.close()
    first, second = rows[0], rows[1]
    assert first[0] == "respondent_id" and second[0] is None
    grid = first.index(
        "Au cours du dernier mois, à quelle fréquence avez-vous vécu les "
        "situations suivantes au travail ?"
    )
    assert all(cell is None for cell in first[grid + 1 : grid + 12])
    assert second[grid] == "J'ai eu trop de tâches à accomplir en peu de temps."
    assert "Response" in second and "Open-Ended Response" in second
    assert "Other (please specify)" in second
    assert rows[2][grid] in ("Jamais", "Rarement", "Parfois", "Souvent", "Toujours")
    assert len(rows) == 2 + 110


def test_the_csv_export_has_the_same_rows():
    workbook = load_workbook(EXAMPLES / "travail-t2_surveymonkey.xlsx", read_only=True)
    expected = [
        ["" if c is None else str(c) for c in row]
        for row in workbook.worksheets[0].iter_rows(values_only=True)
    ]
    workbook.close()
    text = (EXAMPLES / "travail-t2_surveymonkey.csv").read_text(encoding="utf-8-sig")
    import csv
    import io

    assert list(csv.reader(io.StringIO(text))) == expected


def test_wave_1_is_given_renamed(key):
    _df, meta = pyreadstat.read_sav(
        EXAMPLES / "travail-t1-documente.sav", metadataonly=True, user_missing=True
    )
    assert meta.column_names == list(key["waves"]["t1"]["names"].values())
    assert meta.missing_ranges["milieu_1_t1"]  # the team declared 6 and 7 missing
    lines = (EXAMPLES / "travail-t1-correspondance.csv").read_text(encoding="utf-8").splitlines()
    assert lines[0] == "nom_origine,nouveau_nom"
    assert lines[1] == "CollectorNm,nom_collecteur_t1"
    assert len(lines) == 1 + 52


# --- What v0.6 does with these files: the findings of the trial, to be fixed. -------------


@pytest.fixture(scope="module")
def counts(tmp_path_factory):
    """The counts of scripts/compter_structure.py on wave 2, as on a real export."""
    folder = tmp_path_factory.mktemp("vague-2")
    for name in (
        "travail-t2_surveymonkey.sav",
        "travail-t2_surveymonkey.xlsx",
        "questionnaire-travail.pdf",
        "questionnaire-travail.txt",
    ):
        (folder / name).write_bytes((EXAMPLES / name).read_bytes())
    found = load_counter().count_base(folder)
    assert not [k for k in found if k.startswith("erreur")]
    return found


def test_the_counts_see_the_structure(counts):
    assert counts["variables_plateforme"] == 10
    assert counts["grilles"] == 2
    assert counts["ensembles_a_cocher"] == 3
    assert counts["dates_de_passation"] == 2
    assert counts["taux_min_valeurs_de_date_lisibles_pct"] == 100.0
    assert counts["codes_a_etiquette_de_manquant"] == 10


FINDING = pytest.mark.xfail(strict=True, reason="v0.6: finding of the trial, fixed in v0.7")


def test_the_platform_is_recognized(counts):
    assert counts["plateforme_reconnue"] == 1


def test_the_second_header_row_is_not_a_participant(counts):
    assert counts["xlsx_lignes_en_trop"] == 0
    assert counts["xlsx_colonnes_sans_nom"] == 0
    assert counts["xlsx_noms_absents_du_sav"] == 0


@FINDING
def test_single_questions_do_not_make_a_group(counts):
    assert counts["groupes_melant_des_questions"] == 0


@FINDING
def test_a_grid_is_one_group(counts):
    assert counts["grilles_coupees"] == 0
    assert counts["grilles_entieres"] == counts["grilles"]


@FINDING
def test_a_checkbox_set_is_a_multiple_response_group(counts):
    assert counts["ensembles_a_cocher_groupes"] == counts["ensembles_a_cocher"]


@FINDING
def test_the_verb_identifier_does_not_make_an_identifier(counts):
    assert counts["items_de_grille_identifiant_probable"] == 0


@FINDING
def test_the_survey_dates_are_read_as_dates(counts):
    assert counts["dates_lues_comme_dates"] == counts["dates_de_passation"]
    assert counts["dates_classees_texte_libre"] == 0


@FINDING
def test_dont_know_codes_are_suspect_missing_codes(counts):
    assert (
        counts["codes_a_etiquette_de_manquant_signales"] == counts["codes_a_etiquette_de_manquant"]
    )


@FINDING
def test_the_scale_level_of_labelled_items_is_flagged(counts):
    assert counts["items_niveau_scale_signales"] == counts["items_de_grille_niveau_scale"]


def _expected_questions(key) -> list[tuple[str, int]]:
    return [(q["kind"], q["n_rows"]) for q in key["waves"]["t2"]["questions"]]


@FINDING
def test_the_numbered_text_questionnaire_is_read(key):
    found = load_questionnaire(EXAMPLES / "questionnaire-travail.txt").questions
    assert [(q.kind, len(q.rows)) for q in found] == _expected_questions(key)


@FINDING
def test_the_surveymonkey_pdf_is_read(key):
    found = load_questionnaire(EXAMPLES / "questionnaire-travail.pdf").questions
    assert [(q.kind, len(q.rows)) for q in found] == _expected_questions(key)
