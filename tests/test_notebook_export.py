"""v0.11, step 4: the results document of the notebook (D-089, D-090, D-092): the retained
analyses recomputed, numbered one after the other, a person's titles and notes kept, the
common test of the formats (D-083) passed on a document of several analyses."""

import io
from pathlib import Path

import pytest
from docx import Document
from openpyxl import load_workbook

from ddict.analysis import notebook as nb
from ddict.analysis import notebook_export as export
from ddict.analysis.apa import Edit, plain
from ddict.analysis.compute import run
from ddict.analysis.data import apply_decisions
from ddict.analysis.formats import FORMATS, available
from ddict.io import read_database
from ddict.project import build_project
from ddict.project.model import TrackedValue
from ddict.review import logic
from test_analysis_data import KEY, _prepared, analyses_project

EXAMPLES = Path(__file__).resolve().parent.parent / "exemples" / "bases-fictives"
TODAY = "2026-10-16"
CASES = [
    ("table1", {"variables": ["age", "sexe", "niveau_etudes", "bienetre_total"],
                "group": ["groupe"]}, None),
    ("t_test", {"outcome": ["bienetre_total"], "group": ["groupe"]}, None),
    ("anova", {"outcome": ["stress"], "group": ["niveau_etudes"]}, {"post_hoc": "tukey"}),
    ("correlation", {"variables": ["heures_travail", "stress", "soutien"]}, None),
    ("chi_square", {"rows": ["teletravail"], "columns": ["secteur"]}, None),
]  # fmt: skip


@pytest.fixture
def ready():
    project = analyses_project()
    prepared = _prepared(project, KEY["sav"])
    notebook = nb.new_notebook(project, Path("etude-analyses.dd.json"))
    for key, chosen, options in CASES:
        edits = (
            {"table-1": Edit(number=9, title="Bien-être selon le groupe")}
            if key == "t_test"
            else None
        )
        nb.add(notebook, project, run(project, prepared, key, chosen, options), edits=edits,
               today=TODAY)  # fmt: skip
    nb.remove(notebook, "A5")
    nb.move(notebook, "A2", 1)
    nb.move(notebook, "A3", 2)
    nb.retain(notebook, "A4", False)
    return project, prepared, notebook


def _analyses(document):
    """The sections of the analyses, after those of the Method (v0.12)."""
    return [s for s in document.sections if s.key[:1] == "A" and s.key[1:].isdigit()]


def _numbers(document):
    return [(s.key, [t.number for t in s.tables], [f.number for f in s.figures])
            for s in _analyses(document)]  # fmt: skip


# --- Numbering (D-090), as the sheet's example (§ 4). ---


def test_the_sheets_numbering(ready):
    document = export.results_document(*ready, today=TODAY)
    assert _numbers(document) == [("A2", [1], [1]), ("A3", [2, 3], [2]), ("A1", [4], [])]
    assert _analyses(document)[0].heading == "A2 — Test t : comparer deux groupes"
    t_test = _analyses(document)[0]
    assert t_test.tables[0].title == "Bien-être selon le groupe"  # a person's title kept
    assert t_test.tables[0].number == 1  # but never their number: the order numbers
    assert t_test.sentence.startswith("bienetre_total (score de G1) : M = 17,01")


def test_moving_an_analysis_renumbers(ready):
    project, prepared, notebook = ready
    nb.move(notebook, "A1", 1)
    document = export.results_document(project, prepared, notebook, today=TODAY)
    assert _numbers(document) == [("A1", [1], []), ("A2", [2], [1]), ("A3", [3, 4], [2])]
    notebook.first_table, notebook.first_figure = 3, 2
    document = export.results_document(project, prepared, notebook, today=TODAY)
    assert _numbers(document) == [("A1", [3], []), ("A2", [4], [2]), ("A3", [5, 6], [3])]


def test_the_closing_lines(ready):
    document = export.results_document(*ready, today=TODAY)
    closing = list(document.closing)
    assert closing[0].startswith("Empreinte du codebook : ")
    assert closing[1].startswith("Produit le 2026-10-16 par dictionnaire-donnees ")
    assert closing[2] == "Analyses du document :"
    assert closing[3].startswith("A2 — Test t : comparer deux groupes — Variable dépendante : ")
    assert not any("A4" in line for line in closing)  # not retained


# --- Recomputed, never copied (D-089). ---


def test_an_impossible_analysis_is_left_out_and_said(ready):
    project, prepared, notebook = ready
    project.variable("stress").measure_level = TrackedValue[str](
        value="ordinal", origin="human", status="none"
    )
    document = export.results_document(project, prepared, notebook, today=TODAY)
    assert [s.key for s in _analyses(document)] == ["A2", "A1"]
    assert _numbers(document) == [("A2", [1], [1]), ("A1", [2], [])]  # no hole
    assert "Analyses retenues mais écartées, impossibles à calculer :" in document.closing
    assert any(line.startswith("A3 — ANOVA") and "ordinal" in line for line in document.closing)


def test_a_change_waiting_for_a_confirmation_blocks_the_export(tmp_path):
    database = read_database(EXAMPLES / "etude-sommeil.sav")
    project = build_project(database)
    notebook = nb.new_notebook(project, tmp_path / "etude-sommeil.dd.json")
    chosen = {"variables": ["age", "heures_sommeil"]}
    nb.add(notebook, project, run(project, apply_decisions(project, database.table),
                                  "correlation", chosen), today=TODAY)  # fmt: skip
    logic.decide_sentinel(project, "age", 999, True, now="2026-10-16T12:00:00")
    prepared = apply_decisions(project, database.table)
    with pytest.raises(export.ExportBlocked) as error:
        export.results_document(project, prepared, notebook, today=TODAY)
    assert error.value.ids == ["A1"]
    from ddict.analysis import notebook_check

    notebook_check.confirm(notebook, "A1")
    document = export.results_document(project, prepared, notebook, today=TODAY)
    assert plain(document.figures[0].note).startswith("N = 197.")  # recomputed, confirmed
    assert "r(195) = ,07" in _analyses(document)[0].sentence


# --- The files. ---


def test_the_files_next_to_the_project(ready, tmp_path):
    project, prepared, notebook = ready
    project_path = tmp_path / "etude-analyses.dd.json"
    paths = export.results_paths(project_path)
    assert [p.name for p in paths.values()] == [
        "etude-analyses-article.docx",
        "etude-analyses-article-graphiques.xlsx",
        "etude-analyses-article-R.zip",
    ]
    written = export.write_results(project, prepared, notebook, project_path, today=TODAY)
    assert written == [paths["docx"], paths["xlsx"], paths["r"]]
    word = Document(io.BytesIO(paths["docx"].read_bytes()))
    headings = [p.text for p in word.paragraphs if p.style.name.startswith("Heading")]
    assert headings == ["Résumé", "[à compléter : titre]", "Méthode", "Participants",
                        "Instruments", "Analyses", "Résultats",
                        "A2 — Test t : comparer deux groupes",
                        "A3 — ANOVA : comparer plusieurs groupes",
                        "A1 — Tableau 1 : décrire l'échantillon", "Discussion",
                        "Références"]  # fmt: skip
    texts = [p.text for p in word.paragraphs]
    assert [t for t in texts if t.startswith("Tableau ")] == [f"Tableau {n}" for n in range(1, 5)]
    assert [t for t in texts if t.startswith("Figure ")] == ["Figure 1", "Figure 2"]
    assert len(word.tables) == 4 and len(word.inline_shapes) == 2
    book = load_workbook(io.BytesIO(paths["xlsx"].read_bytes()))
    assert book.sheetnames == ["Figure 1", "Figure 2"]
    export.write_results(project, prepared, notebook, project_path, today=TODAY)  # replaced
    assert paths["docx"].is_file()


# --- The common test of the formats (D-083), on several analyses. ---


@pytest.mark.parametrize("language", ["fr", "en"])
@pytest.mark.parametrize("key", list(FORMATS))
def test_every_format_renders_the_results(ready, key, language):
    document = export.results_document(*ready, language=language, today=TODAY)
    found = FORMATS[key]
    assert found in available(document)
    text = " ".join(found.read_text(found.render(document)).replace(chr(0xA0), " ").split())
    expected = []
    if found.tables:
        for table in document.tables:
            expected += [table.title, plain(table.note)]
            expected += [plain(c.header) for c in table.columns]
            expected += [plain(cell) for row in table.rows for cell in row.cells]
        expected += [s.heading for s in document.sections] + [s.sentence for s in document.sections]
        expected += [plain(p) for s in document.sections for p in s.paragraphs]  # v0.12
        expected += list(document.closing)
    if found.figures:
        for figure in document.figures:
            expected += [figure.title, plain(figure.note)]
    flat = [" ".join(t.replace(chr(0xA0), " ").split()) for t in expected if t.strip()]
    missing = [t for t in flat if t not in text]
    assert not missing, missing
