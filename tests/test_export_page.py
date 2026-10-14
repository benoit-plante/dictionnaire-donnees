"""v0.3, step 5: the Export page gives the codebook `ddict exporter` gives."""

import shutil
from pathlib import Path

import pytest
from openpyxl import load_workbook
from streamlit.testing.v1 import AppTest

from ddict.cli import main
from ddict.project import load_project, save_project
from ddict.project.model import TrackedValue
from ddict.review import logic
from ddict.review.ui import export_view
from ddict.workflow import UnknownGroups, add_questionnaire, export_codebook, profile_database

APP = str(Path(__file__).resolve().parent.parent / "src" / "ddict" / "review" / "app.py")


@pytest.fixture
def project_path(fictional_files, tmp_path):
    folder = tmp_path / "projet"
    folder.mkdir()
    for name in ("etude-sommeil.sav", "questionnaire-sommeil.txt"):
        shutil.copy(fictional_files[name], folder / name)
    project, path = profile_database(folder / "etude-sommeil.sav")
    add_questionnaire(project, folder / "questionnaire-sommeil.txt")
    logic.link_question(project, "age", "Q3", now="2026-10-14T12:00:00")
    logic.set_omit_question_text(project, "G1", True)
    project.variable("genre").label = TrackedValue[str](
        value="Genre déclaré", origin="ai", status="proposed", ai_run="run-0001"
    )
    save_project(project, path)
    return path


def _open(path):
    app = AppTest.from_file(APP, default_timeout=30)
    app.session_state["project_path"] = str(path)
    app.session_state["mode"] = "export"
    app.run()
    assert not app.exception, app.exception
    return app


def _sheets(path):
    book = load_workbook(path)
    return {s.title: list(s.iter_rows(values_only=True)) for s in book.worksheets}


@pytest.mark.allow_hosts(["127.0.0.1"])
def test_same_codebook_as_the_command(project_path, tmp_path):
    app = _open(project_path)
    assert app.header[0].value == "Export"
    assert app.multiselect(key="export-omit").value == ["G1"]  # the team's choice, by default
    assert "propositions sans décision" in app.info[0].value
    app.button(key="export-write").click().run()
    assert "Fichiers écrits à côté du projet" in app.success[0].value

    command = tmp_path / "commande"
    command.mkdir()
    main(
        [
            "exporter",
            str(project_path),
            "--xlsx",
            str(command / "c.xlsx"),
            "--md",
            str(command / "c.md"),
            "--sans-texte",
            "G1",
        ]
    )
    page_xlsx = project_path.with_name("etude-sommeil.xlsx")
    assert _sheets(page_xlsx) == _sheets(command / "c.xlsx")
    assert project_path.with_name("etude-sommeil.md").read_text(encoding="utf-8") == (
        command / "c.md"
    ).read_text(encoding="utf-8")


@pytest.mark.allow_hosts(["127.0.0.1"])
def test_work_version(project_path):
    app = _open(project_path)
    app.radio(key="export-kind").set_value("work").run()
    app.button(key="export-write").click().run()
    report = project_path.with_name("etude-sommeil-travail.md").read_text(encoding="utf-8")
    assert "Genre déclaré [proposé par IA]" in report
    assert not project_path.with_name("etude-sommeil.md").exists()  # the final one untouched


@pytest.mark.allow_hosts(["127.0.0.1"])
def test_a_file_open_in_excel_is_explained(project_path):
    project_path.with_name("etude-sommeil.xlsx").mkdir()  # cannot be written over
    app = _open(project_path)
    app.button(key="export-write").click().run()
    assert "Impossible d'écrire les fichiers" in app.error[0].value
    assert "ouvert dans Excel" in app.error[0].value


@pytest.mark.allow_hosts(["127.0.0.1"])
def test_open_folder(project_path, monkeypatch):
    opened = []
    monkeypatch.setattr(export_view.os, "startfile", opened.append, raising=False)
    monkeypatch.setattr(export_view.sys, "platform", "win32")
    app = _open(project_path)
    app.button(key="export-folder").click().run()
    assert opened == [str(project_path.resolve().parent)]


def test_unknown_groups(project_path):
    with pytest.raises(UnknownGroups) as error:
        export_codebook(load_project(project_path), project_path, omit_groups=("G9",))
    assert error.value.known == ["G1", "G2"]
