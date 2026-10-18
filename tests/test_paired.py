"""v0.14, step 2: the paired t test and Wilcoxon's signed-rank test (docs/catalogue-elargi.md,
§ 2), checked against the answer key of the follow-up base, against R, in every format,
and in the plan and the article."""

import io
import json
import shutil
import subprocess
from pathlib import Path

import numpy as np
import pandas as pd
import pytest
from docx import Document
from streamlit.testing.v1 import AppTest

from ddict.analysis import article_text as article
from ddict.analysis import conditions, prereg
from ddict.analysis import planning as pl
from ddict.analysis.apa import describe, plain
from ddict.analysis.book import Codebook
from ddict.analysis.compute import run
from ddict.analysis.data import apply_decisions
from ddict.analysis.formats import FORMATS, available
from ddict.analysis.method_text import _method_sentences
from ddict.analysis.models import wilcoxon
from ddict.analysis.rscript import r_script
from ddict.cli import main
from ddict.export.syntax import write_syntax
from ddict.io import read_database
from ddict.project import build_project, save_project
from ddict.project.model import TrackedValue
from test_analysis_page import APP
from test_syntax import _rscript

ROOT = Path(__file__).resolve().parent.parent
EXAMPLES = ROOT / "exemples" / "bases-fictives"
KEY = json.loads((EXAMPLES / "corrige-suivi.json").read_text(encoding="utf-8"))
EXACT = KEY["exact"]
TODAY = "2026-10-18"
PAIR = {"first": ["anxiete_t1"], "second": ["anxiete_t2"]}


@pytest.fixture(scope="module")
def ready():
    data = read_database(EXAMPLES / KEY["sav"])
    project = build_project(data)
    return project, apply_decisions(project, data.table)


def _run(ready, options=None, chosen=PAIR):
    project, prepared = ready
    return run(project, prepared, "paired", chosen, options)


# --- Against the answer key and the sheet. ---


def test_against_the_key(ready):
    s = _run(ready).statistics
    exact = EXACT["t_apparie_t1_t2"]
    assert s["n"] == exact["n"]
    assert s["diff"] == pytest.approx(exact["mean_diff"], abs=1e-9)
    assert s["t"] == pytest.approx(exact["t"], abs=1e-8)
    assert s["df"] == exact["df"]
    assert s["d"] == pytest.approx(exact["d_z"], abs=1e-9)


def test_the_sentence_of_the_sheet(ready):
    result = _run(ready)
    assert result.sentence == (
        "Anxiété au début du programme (0 à 40) : M = 21,56 (ÉT = 6,38) ; anxiété à la fin du "
        "programme (0 à 40) : M = 18,56 (ÉT = 6,65) ; différence 3,00, IC 95 % [2,28 ; 3,72], "
        "t(227) = 8,18, p < ,001, d_z = 0,54, IC 95 % [0,40 ; 0,68] (n = 228)."
    )
    assert (result.n_used, result.n_excluded) == (228, 12)
    wilcox = _run(ready, {"method": "wilcoxon"})
    assert "V = 20169,0, p < ,001, r = ,56 (n = 228)." in wilcox.sentence
    assert wilcox.warnings == ["1 différences nulles écartées par le test de Wilcoxon."]


# --- The conditions. ---


def test_the_conditions(ready):
    project, prepared = ready
    same = conditions.check(project, "paired", {"first": ["anxiete_t1"],
                                                "second": ["anxiete_t1"]})  # fmt: skip
    assert same.refusals == ["anxiete_t1 : la même variable dans deux rôles."]
    nominal = conditions.check(project, "paired", {"first": ["sexe"], "second": ["anxiete_t1"]})
    assert any("ne convient pas" in r for r in nominal.refusals)
    assert conditions.check(project, "paired", PAIR, data=prepared).ok
    long = project.model_copy(deep=True)
    long.dataset.structure = TrackedValue[str](value="long", origin="human", status="validated")
    refused = conditions.check(long, "paired", PAIR)
    assert any("format long" in r for r in refused.refusals)
    # Not a warning about independence: the measures are the same people's on purpose.
    assert not conditions.check(project, "paired", PAIR).warnings


# --- Wilcoxon, as R 4.6 computes it. ---


def test_wilcoxon_as_r_computes_it():
    """Values from R 4.6.1, wilcox.test(d, exact = FALSE, correct = TRUE): with ties, without
    ties, a small sample, zeros; the three alternatives."""
    rng = np.random.default_rng(1)
    cases = {
        "ties": np.round(rng.normal(0.3, 1, 40), 1),
        "noties": rng.normal(0.4, 1, 25),
        "small": rng.normal(-0.5, 1, 8),
        "zeros": np.r_[np.zeros(3), rng.normal(0.2, 1, 30)],
    }
    from_r = {
        "ties": (502.5, 0.0230939976673659, 0.0115469988336829, 0.988901258853407),
        "noties": (212, 0.187356840856879, 0.0936784204284394, 0.910743170997295),
        "small": (6, 0.107327557761612, 0.959971040007783, 0.053663778880806),
        "zeros": (292, 0.224926479146903, 0.112463239573451, 0.891417245534212),
    }
    for name, d in cases.items():
        v, two, greater, less = from_r[name]
        assert wilcoxon(d, "two-sided")["V"] == v
        for alternative, p in (("two-sided", two), ("greater", greater), ("less", less)):
            assert wilcoxon(d, alternative)["p"] == pytest.approx(p, rel=1e-12), (name, alternative)


# --- Against the equivalent R script. ---


def _r_results(ready, options, tmp_path) -> dict[str, float]:
    project = ready[0]
    shutil.copy(EXAMPLES / KEY["sav"], tmp_path / KEY["sav"])
    write_syntax(project, {"r": tmp_path / "etude-suivi-metadonnees.R"}, TODAY)
    script = r_script(project, "etude-suivi", "paired", PAIR, options, TODAY)
    (tmp_path / "analyse.R").write_text(script, encoding="utf-8")
    (tmp_path / "run.R").write_text(
        'source("analyse.R", encoding = "UTF-8")\n'
        "write.csv(data.frame(nom = names(unlist(resultats)), valeur = unlist(resultats)), "
        '"sortie.csv", row.names = FALSE)\n',
        encoding="utf-8",
    )
    done = subprocess.run([_rscript(), "run.R"], cwd=tmp_path, capture_output=True, text=True,
                          encoding="utf-8", errors="replace")  # fmt: skip
    if "there is no package called" in done.stderr:
        pytest.skip("haven is not installed in R")
    assert done.returncode == 0, done.stderr
    out = pd.read_csv(tmp_path / "sortie.csv")
    return dict(zip(out["nom"], out["valeur"], strict=True))


@pytest.mark.skipif(_rscript() is None, reason="R is not installed")
@pytest.mark.parametrize(
    "options",
    [None, {"alternative": "greater"}, {"method": "wilcoxon"},
     {"method": "wilcoxon", "alternative": "less"}],
)  # fmt: skip
def test_the_same_results_as_r(ready, options, tmp_path):
    here = _run(ready, options).statistics
    from_r = _r_results(ready, options, tmp_path)
    common = [k for k in here if k in from_r]
    assert len(common) >= 9
    for name in common:
        assert here[name] == pytest.approx(from_r[name], rel=1e-6, abs=1e-10), name


# --- Tables and figures. ---


@pytest.mark.parametrize("language", ["fr", "en"])
@pytest.mark.parametrize("options", [None, {"method": "wilcoxon"}])
@pytest.mark.parametrize("key", list(FORMATS))
def test_every_format_renders_everything(ready, key, options, language):
    project, prepared = ready
    document = describe(project, prepared, _run(ready, options), language=language)
    found = FORMATS[key]
    assert found in available(document)
    text = " ".join(found.read_text(found.render(document)).replace(chr(0xA0), " ").split())
    expected = [document.figures[0].title, plain(document.figures[0].note)]
    if found.tables:
        table = document.tables[0]
        expected += [table.title, plain(table.note), *[plain(c.header) for c in table.columns]]
        expected += [plain(cell) for cell in table.rows[0].cells]
    missing = [t for t in expected if " ".join(t.replace(chr(0xA0), " ").split()) not in text]
    assert not missing, missing


def test_the_word_table_and_figure(ready):
    project, prepared = ready
    document = describe(project, prepared, _run(ready))
    assert document.figures[0].kind == "line"
    word = Document(io.BytesIO(FORMATS["docx"].render(document)))
    header = word.tables[0].rows[1].cells
    dz = [(r.text, r.italic, r.font.subscript) for r in header[8].paragraphs[0].runs]
    assert dz == [("d", True, None), ("z", False, True)]  # d_z, the z in subscript
    assert [c.text for c in word.tables[0].rows[2].cells][:3] == ["n = 228", "21,56", "6,38"]


# --- In the Method, the plan and the article. ---


def test_the_method_sentence():
    assert plain(_method_sentences("paired", "paired_t", "none")[0]).startswith(
        "Deux mesures d'une même personne ont été comparées par le test t apparié"
    )
    assert "Wilcoxon" in plain(_method_sentences("paired", "wilcoxon", "none")[0])


def test_in_the_plan_and_the_article(ready):
    project, prepared = ready
    data = pl.data_from_names(project, "L'anxiété baisse.", "paired", PAIR, {},
                              pair=None, predictor=None, model=False,
                              direction="first_higher", kind="confirmatory",
                              family="")  # fmt: skip
    hypothesis = pl.build(project, data, "H1")
    assert hypothesis.effect.kind == "difference"
    assert hypothesis.options == {"method": "paired_t", "alternative": "greater"}
    assert prereg._effect(project, hypothesis) == "la différence entre les deux mesures"
    assert prereg._direction(project, hypothesis) == (
        "Anxiété au début du programme (0 à 40) plus élevé que Anxiété à la fin du programme "
        "(0 à 40)"
    )
    result = run(project, prepared, "paired", PAIR, hypothesis.options)
    assert result.statistics["p"] < 1e-13  # one-sided, the expected direction
    tested = article.Tested(hypothesis, result, "p")
    book = Codebook(project, prepared)
    assert article.observed(book, prepared, tested) == (
        "La moyenne observée était plus élevée pour anxiété au début du programme (0 à 40)."
    )


# --- The page and the command line. ---


def _suivi(tmp_path, name):
    folder = tmp_path / name
    folder.mkdir()
    shutil.copy(EXAMPLES / KEY["sav"], folder / KEY["sav"])
    path = folder / "etude-suivi.dd.json"
    save_project(build_project(read_database(folder / KEY["sav"])), path)
    return path


@pytest.mark.allow_hosts(["127.0.0.1"])
def test_the_page_gives_the_command_result(tmp_path):
    path = _suivi(tmp_path, "page")
    app = AppTest.from_file(APP, default_timeout=60)
    app.session_state["project_path"] = str(path)
    app.session_state["mode"] = "analyses"
    app.run()
    app.selectbox(key="analysis-key").set_value("paired").run()
    app.selectbox(key="analysis-paired-first").set_value("anxiete_t1").run()
    app.selectbox(key="analysis-paired-second").set_value("anxiete_t2").run()
    app.button(key="analysis-run").click().run()
    assert not app.exception, app.exception
    sentence = next(m.value for m in app.markdown if m.value.startswith("> "))
    assert "t(227) = 8,18" in sentence
    app.button(key="analysis-write").click().run()
    other = _suivi(tmp_path, "commande")
    assert main(["analyser", str(other), "apparie", "--premiere", "anxiete_t1", "--seconde",
                 "anxiete_t2"]) == 0  # fmt: skip
    for suffix in (".md", ".R"):
        name = f"etude-suivi-analyse-1{suffix}"
        assert path.with_name(name).read_text(encoding="utf-8") == other.with_name(name).read_text(
            encoding="utf-8"
        )
