"""v0.12, step 7: the Method written by templates (D-097), the document "Méthode et
Résultats" (D-099) and the Method's figures confirmed in the notebook (D-100), as the
validated sheet docs/methode-gabarits.md gives them, on the fictional database of
administration."""

import io
import json
import re
from pathlib import Path

import pytest
from docx import Document
from docx.enum.text import WD_COLOR_INDEX

from ddict.analysis import method, method_text
from ddict.analysis import notebook as nb
from ddict.analysis import notebook_check as check
from ddict.analysis import notebook_export as export
from ddict.analysis.apa import plain
from ddict.analysis.compute import run
from ddict.analysis.data import apply_decisions
from ddict.project import exclusions as ex
from ddict.project.model import SourcedText, TrackedValue
from test_exclusions import KEY, _project

TODAY = "2026-10-16"
EXACT = KEY["exact"]
INSTRUMENT = SourcedText(value="Échelle d'anxiété fictive", source="Laboratoire fictif, 2024")


def _ready(rules=True, named=True):
    project, table = _project(rules=rules)
    group = project.groups[0]
    if named:
        group.instrument_name = INSTRUMENT
        group.descriptive_label = TrackedValue[str](
            value="L'anxiété au cours du dernier mois", origin="human", status="none"
        )
    prepared = apply_decisions(project, table)
    notebook = nb.new_notebook(project, Path("etude-passation.dd.json"))
    notebook.method.variables = [nb.reference(project, "age"), nb.reference(project, "genre")]
    chosen = {"outcome": ["anx_total"], "group": ["genre"]}
    nb.add(notebook, project, run(project, prepared, "t_test", chosen, {"levels": "1,2"}),
           today=TODAY)  # fmt: skip
    nb.add(notebook, project, run(project, prepared, "correlation",
                                  {"variables": ["stress", "anx_total"]}), today=TODAY)  # fmt: skip
    return project, table, prepared, notebook


@pytest.fixture(scope="module")
def ready():
    return _ready()


# --- The facts, equal to the answer key. ---


def test_the_facts(ready):
    project, _table, prepared, notebook = ready
    found = method.facts(project, prepared, notebook)
    assert dict(found.flow.excluded) == EXACT["flow"]["excluded"]
    age, genre = found.described
    assert age.n == EXACT["sample"]["age"]["n"] == 341
    assert age.mean == pytest.approx(EXACT["sample"]["age"]["mean"], abs=1e-9)
    assert age.sd == pytest.approx(EXACT["sample"]["age"]["sd"], abs=1e-9)
    shares = {str(int(code)): pct for code, _label, pct in genre.categories}
    assert shares == pytest.approx(EXACT["sample"]["genre"]["percent"], abs=1e-9)
    [alpha] = found.reliabilities
    assert alpha.alpha == pytest.approx(EXACT["anxiety"]["alpha"], abs=1e-9)
    assert found.missing == (("age", (999,)),)
    assert found.methods == (("t_test", "welch", "none"), ("correlation", "pearson", "none"))


def test_the_statistics_kept(ready):
    project, _table, prepared, notebook = ready
    stats = method.statistics(method.facts(project, prepared, notebook))
    assert stats["n_start"] == 400 and stats["n_analysed"] == 347 and stats["excluded_E3"] == 14
    assert stats["v1_n"] == 341 and stats["v2_pct_1"] == pytest.approx(55.6195965418)
    assert stats["alpha_n_G1"] == 343
    assert not [k for k in stats if "min" in k or "max" in k]  # no observed range (D-088)


# --- The paragraphs, as the sheet gives them. ---


def _paragraphs(ready, language="fr"):
    from ddict import i18n

    project, _table, prepared, notebook = ready
    found = method.facts(project, prepared, notebook)
    with i18n.using(language):
        return (
            plain(method_text.participants(project, found, 1)),
            [plain(p) for p in method_text.instruments(project, prepared, found)],
            plain(method_text.analyses(found, "etude-passation-resultats.R")),
        )


def test_participants(ready):
    text, _instruments, _analyses = _paragraphs(ready)
    assert text == (
        "Au total, 400 personnes ont répondu au questionnaire. 53 ont été exclues, pour les "
        "raisons suivantes, dans cet ordre : refus de consentement (n = 12) ; moins de 18 ans "
        "(n = 7) ; moins de 5 items d'anxiété répondus sur 6 (n = 14) ; passation de moins de 3 "
        "minutes (n = 10) ; grille de satisfaction en ligne droite (n = 5) ; courriel en double "
        "(n = 5) (Figure 1). L'échantillon analysé compte 347 personnes. Elles sont décrites par "
        "les variables suivantes : âge (années) (M = 45,07, ÉT = 15,04, n = 341) ; genre "
        "(femme, 55,6 % ; homme, 42,9 % ; autre, 1,4 %)."
    )


def test_instruments(ready):
    _participants, [text], _analyses = _paragraphs(ready)
    assert text == (
        "L'instrument Échelle d'anxiété fictive (Laboratoire fictif, 2024) mesure l'anxiété au "
        "cours du dernier mois au moyen de 6 items, notés de 1 (jamais) à 4 (très souvent). "
        "L'item anx3 est inversé avant le calcul du score. Le score (anx_total) est la somme des "
        "items, de 6 à 24. Il est calculé quand au moins 5 items sur 6 sont répondus ; un item "
        "manquant est remplacé par la moyenne des items répondus de la personne. La cotation "
        "suit Laboratoire fictif (2024). Sur l'échantillon analysé, l'alpha de Cronbach est de "
        ",89 (IC 95 % [,87 ; ,91], n = 343)."
    )


def test_analyses(ready):
    _participants, _instruments, text = _paragraphs(ready)
    assert text.startswith(
        "Le code 999 (age) a été traité comme une valeur manquante. Les valeurs manquantes ont "
        "été retirées analyse par analyse"
    )
    assert "par le test t de Welch, qui ne suppose pas des variances égales" in text
    assert "corrélation de Pearson, avec son intervalle de confiance à 95 %" in text
    assert text.endswith("un script R joint (etude-passation-resultats.R) les refait et donne "
                         "les mêmes valeurs.")  # fmt: skip


def test_in_english(ready):
    participants, [instrument], analyses = _paragraphs(ready, "en")
    assert participants.startswith("In total, 400 people answered the questionnaire. 53 were")
    assert "(women" not in participants and "1,4" not in participants and "1.4%" in participants
    assert instrument.startswith(
        "The Échelle d'anxiété fictive (Laboratoire fictif, 2024) measures"
    )
    assert "Cronbach's alpha is .89 (95% CI [.87, .91], n = 343)" in instrument
    assert analyses.startswith("Code 999 (age) was treated as a missing value.")


# --- The holes (D-097): never a guess. ---


def test_the_holes():
    project, _table, prepared, notebook = _ready(named=False)
    found = method.facts(project, prepared, notebook)
    [text] = method_text.instruments(project, prepared, found)
    holes = [run.text for run in text if run.hole]
    assert holes == [
        "[à compléter : nom et source de l'instrument du groupe G1]",
        "[à compléter : ce que mesure le groupe G1]",
    ]
    group = project.groups[0]
    group.scoring = group.scoring.model_copy(update={"items": [], "note": "Somme des items."})
    [text] = method_text.instruments(project, prepared, method.facts(project, prepared, notebook))
    assert (
        plain(text)
        == "[à compléter : description de l'instrument du groupe G1] « Somme des items. »"
    )


def test_no_alpha_gives_a_hole(ready):
    project, _table, prepared, notebook = ready
    small = prepared.__class__(prepared.table.head(1), prepared.scores, prepared.reversed,
                               prepared.flow)  # fmt: skip
    [text] = method_text.instruments(project, small, method.facts(project, small, notebook))
    assert text[-1].hole and "trop peu de réponses complètes" in text[-1].text


def test_no_rule_no_exclusion_sentence():
    project, _table, prepared, notebook = _ready(rules=False)
    text = plain(method_text.participants(project, method.facts(project, prepared, notebook), None))
    assert text.startswith("Au total, 400 personnes ont répondu au questionnaire. L'échantillon "
                           "analysé compte 400 personnes.")  # fmt: skip


def test_every_number_has_an_origin(ready):
    """Each number of the text is a count, a statistic, a code, a bound or a version that
    the codebook or the facts hold: nothing else (D-097)."""
    project, _table, prepared, notebook = ready
    found = method.facts(project, prepared, notebook)
    participants, instruments, analyses = _paragraphs(ready)
    allowed = {"95", "001", "3", "18", "5", "6", "1", "4", "24", "999", "0", "2024", "2"}
    from ddict.analysis.report import versions
    from ddict.analysis.results import bounded, number

    for value in versions().values():
        allowed |= set(re.findall(r"\d+", value))
    stats = method.statistics(found)
    for value in stats.values():
        allowed |= {number(value, 0), number(value), number(value, 1), bounded(value)}
    allowed.add(number(stats["n_start"] - stats["n_analysed"], 0))  # the people excluded
    allowed |= {s.lstrip(",") for s in list(allowed)}
    text = " ".join([participants, *instruments, analyses])
    numbers = set(re.findall(r"\d+(?:,\d+)?", text))
    assert numbers - allowed == set()


def test_no_judging_words(ready):
    text = " ".join([p if isinstance(p, str) else " ".join(p) for p in _paragraphs(ready)]).lower()
    for word in ("excellent", "bonne fidélité", "satisfaisant", "significati", "élevé", "faible"):
        assert word not in text


# --- The document "Méthode et Résultats" (D-099). ---


def test_the_document(tmp_path):
    project, _table, prepared, notebook = _ready()
    document = export.results_document(project, prepared, notebook, today=TODAY)
    keys = [s.key for s in document.sections]
    assert keys == ["methode", "participants", "instruments", "analyses", "resultats", "A1", "A2"]
    assert [s.level for s in document.sections][:5] == [1, 2, 2, 2, 1]
    assert [f.number for f in document.figures] == [1, 2, 3]  # the flow, then the results
    assert document.figures[0].kind == "flow"
    assert [t.number for t in document.tables] == [1, 2]
    assert "Règles d'exclusion :" in document.closing
    assert (
        "E4 — Passation de moins de 3 minutes (Seuil fixé par l'équipe (fictif))"
        in document.closing
    )
    notebook.first_figure = 4
    document = export.results_document(project, prepared, notebook, today=TODAY)
    assert [f.number for f in document.figures] == [4, 5, 6]
    path = tmp_path / "etude-passation.dd.json"
    export.write_results(project, prepared, notebook, path, today=TODAY)
    word = Document(io.BytesIO(path.with_name("etude-passation-resultats.docx").read_bytes()))
    texts = [p.text for p in word.paragraphs]
    assert texts[:2] == ["Méthode", "Participants"]
    assert texts[2].startswith("Au total, 400 personnes")


def test_holes_are_highlighted_in_word(tmp_path):
    project, _table, prepared, notebook = _ready(named=False)
    path = tmp_path / "etude-passation.dd.json"
    export.write_results(project, prepared, notebook, path, today=TODAY)
    word = Document(io.BytesIO(path.with_name("etude-passation-resultats.docx").read_bytes()))
    marked = [r.text for p in word.paragraphs for r in p.runs
              if r.font.highlight_color == WD_COLOR_INDEX.YELLOW]  # fmt: skip
    assert marked == [
        "[à compléter : nom et source de l'instrument du groupe G1]",
        "[à compléter : ce que mesure le groupe G1]",
    ]


# --- Kept and confirmed in the notebook (D-100). ---


def test_the_method_is_confirmed_like_an_analysis():
    project, table, prepared, notebook = _ready()
    first = check.check_method(notebook, project, prepared, today=TODAY)
    assert first.state == "current" and notebook.method.computed.n_used == 347
    # A variable added to describe the sample: a new figure, kept at once.
    notebook.method.variables.append(nb.reference(project, "stress"))
    assert check.check_method(notebook, project, prepared, today=TODAY).state == "current"
    assert "v3_m" in notebook.method.computed.statistics
    # A rule removed: the figures change, and wait for a person.
    ex.remove(project, "E6")
    prepared = apply_decisions(project, table)
    status = check.check_method(notebook, project, prepared, today=TODAY)
    assert status.state == "changed"
    shown = {d.key: (d.shown_before, d.shown_after) for d in status.differences}
    assert shown["n_analysed"] == ("347", "352")
    with pytest.raises(export.ExportBlocked) as error:
        export.results_document(project, prepared, notebook, today=TODAY)
    assert nb.METHOD in error.value.ids
    check.confirm(notebook, nb.METHOD)
    for entry in notebook.entries:
        if entry.pending is not None:
            check.confirm(notebook, entry.id)
    document = export.results_document(project, prepared, notebook, today=TODAY)
    assert "L'échantillon analysé compte 352 personnes." in plain(
        document.sections[1].paragraphs[0]
    )


def test_no_individual_data_in_the_notebook():
    """The canaries of the database (identifier, e-mail address) and the dates never enter
    the notebook's Method."""
    project, _table, prepared, notebook = _ready()
    check.check_method(notebook, project, prepared, today=TODAY)
    text = nb.dumps(notebook)
    assert "ZQX" not in text and "zqxmail" not in text and "2026-03" not in text
    assert json.loads(text)["method"]["variables"] == [{"variable": "age"}, {"variable": "genre"}]
