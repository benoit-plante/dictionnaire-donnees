"""v0.13, step 4: the preregistration (D-109), as the validated sheet
docs/plan-analyse-article.md (§ 8) gives it: every field filled is found in the plan or the
codebook, the holes are flagged, and no count nor result enters it."""

import io
import re
from pathlib import Path

import pytest
from docx import Document

from ddict import __version__
from ddict.analysis import notebook as nb
from ddict.analysis import prereg
from ddict.project import exclusions
from test_analysis_data import KEY, _prepared, analyses_project
from test_plan import EXAMPLE, RULES
from test_plan_notebook import _frozen

TODAY = "2026-10-21"
PATH = Path("etude-analyses.dd.json")


@pytest.fixture
def frozen():
    project = analyses_project()
    prepared = _prepared(project, KEY["sav"])
    notebook = nb.new_notebook(project, PATH)
    exclusions.add(
        project,
        {"kind": "condition", "variable": "age", "operator": "lt", "value": 18,
         "reason": "Moins de 18 ans"},
    )  # fmt: skip
    prepared = _prepared(project, KEY["sav"])
    _frozen((project, prepared, notebook))
    return project, prepared, notebook


def _text(project, language=None):
    return prereg.markdown(prereg.preregistration(project, PATH, language=language, today=TODAY))


def _headings(text):
    return [line for line in text.splitlines() if line.startswith("#")]


# --- Every field filled comes from the plan or the codebook. ---


def test_primary_data_follow_osf_prereg(frozen):
    project = frozen[0]
    text = _text(project)
    headings = _headings(text)
    assert headings[0] == "# Préenregistrement selon le gabarit OSF Prereg"
    assert [h for h in headings if h.startswith("## ")] == [
        "## Informations sur l'étude", "## Plan de l'étude", "## Échantillonnage",
        "## Variables", "## Plan d'analyse", "## Autre",
    ]  # fmt: skip
    assert len([h for h in headings if h.startswith("### ")]) == 24
    assert "Questions adaptées de Bowman et al. (2020), OSF Prereg Template" in text
    assert "sous licence CC0 1.0" in text
    for data in EXAMPLE:
        assert data["statement"] in text
    assert (
        "H1 (confirmatoire, famille F1). Le bien-être au travail est plus élevé dans le groupe "
        "d'intervention que dans le groupe contrôle. Effet testé : la différence entre les deux "
        "groupes. Sens attendu : Intervention plus élevé que Contrôle."
    ) in text
    assert "Effet testé : le coefficient de soutien. Sens attendu : un coefficient positif." in text
    assert (
        "H1. Test t : comparer deux groupes ; méthode : Welch (variances inégales permises) ; "
        "variable dépendante : bienetre_total ; variable de groupe : groupe ; effet testé : la "
        "différence entre les deux groupes ; test unilatéral, le premier groupe plus élevé."
    ) in text
    assert "stress (Score de stress (0 à 40)) : variable de la corrélation (H2) ; variable " in text
    assert "Le seuil α est fixé à ,05." in text
    assert (
        "Les valeurs de *p* seront corrigées par la méthode de Holm dans chaque famille : "
        "F1 (H1, H2, H3)."
    ) in text
    assert (
        "Avant l'analyse : l'équipe avait accédé aux données, sans faire d'analyse liée au "
        "plan. (Déclaré par l'équipe.)"
    ) in text
    assert RULES["sample_size"] in text
    assert "E1. Moins de 18 ans (age < 18)." in text
    assert "Les codes -9 (stress) ; 999 (age) seront traités comme des valeurs manquantes." in text
    assert "au moyen de 5 items, notés de 1 (pas du tout d'accord) à 5" in text
    assert "alpha" not in text.lower()  # the observed alpha is a result
    assert (
        "les analyses prévues pour H3 (A2, ajoutée le 18 octobre 2026)." in text
    )  # the planned analysis the notebook already held
    assert "Plan arrêté dans l'outil le 20 octobre 2026 ; empreinte du plan : " in text


def test_secondary_data_follow_van_den_akker(frozen):
    project = frozen[0]
    project.plan.data_origin = "secondary"
    text = _text(project)
    headings = _headings(text)
    assert headings[0] == "# Préenregistrement d'une analyse de données secondaires"
    numbered = [h for h in headings if re.match(r"### Q\d+\. ", h)]
    assert [int(re.match(r"### Q(\d+)", h).group(1)) for h in numbered] == list(range(1, 26))
    assert "### Q18. Connaissance préalable des données" in headings
    assert headings[-1] == "## Déclaration d'intégrité"
    assert "van den Akker et al. (2021), Meta-Psychology, 5" in text and "CC BY 4.0" in text
    assert "Le dictionnaire de données joint (etude-analyses.xlsx) décrit chaque variable." in text
    assert "La base compte 19 variables." in text
    q18 = text.split("### Q18.")[1].split("## ")[0]
    assert "(Déclaré par l'équipe.)" in q18
    assert "[à compléter : connaissance préalable des données]" in q18
    assert "Aucune pondération n'est utilisée." in text


# --- The holes. ---

HOLES = [
    "titre",
    "auteurs",
    "variables manipulées",
    "ce que l'équipe a vu des données par ailleurs",
    "ce que l'équipe sait des taux de valeurs manquantes",
    "justification de la taille d'échantillon",
]


@pytest.mark.parametrize("origin", ["primary", "secondary"])
def test_what_the_tool_cannot_know_stays_a_hole(frozen, origin):
    project = frozen[0]
    project.plan.data_origin = origin
    text = _text(project)
    for what in HOLES if origin == "primary" else HOLES[:5]:
        assert f"[à compléter : {what}]" in text


def test_a_plan_not_frozen_nor_declared_says_so(frozen):
    project = analyses_project()
    for data in EXAMPLE:
        from ddict.analysis import planning

        planning.add(project, data)
    text = _text(project)
    assert text.startswith("# Préenregistrement selon le gabarit OSF Prereg")
    assert "[à compléter : origine des données]" in text
    assert "[à compléter : ce plan n'est pas encore arrêté dans l'outil]" in text
    assert "[à compléter : seuil α]" in text
    assert "[à compléter : état des données quand le plan a été écrit]" in text
    assert "Plan non arrêté ; empreinte du plan : " in text
    assert "contenait déjà" not in text


# --- No count, no result (§ 8.3). ---


def _allowed(project) -> set[str]:
    """Every number the plan, the codebook and the template may write."""
    found = {__version__, "95", "2020", "2021", "5", "4.0", "1.0", "05"}  # CI, sources, α
    found |= {"18", "20", "21", "2026"}  # dates; the exclusion rule's value
    found |= {str(len(project.variables))}
    found |= {str(n) for n in range(1, 26)}  # the questions
    for entry in project.variables:
        found |= set(re.findall(r"\d+", entry.file.label or ""))
        for missing in entry.file.declared_missing or []:
            found |= set(re.findall(r"\d+", str(missing)))
    for group in project.groups:
        if group.scoring and group.scoring.items:
            s = group.scoring
            k = len(s.items)
            found |= {str(int(v)) for v in (k, s.scale_min, s.scale_max, k * s.scale_min,
                                             k * s.scale_max, s.min_items_answered)}  # fmt: skip
    found |= {"999", "9"}  # the confirmed missing codes
    return found


@pytest.mark.parametrize("origin", ["primary", "secondary"])
def test_no_count_and_no_result(frozen, origin):
    project, prepared, notebook = frozen
    project.plan.data_origin = origin
    text = _text(project)
    text = re.sub(r"https?://\S+", "", text)
    text = re.sub(r"empreinte d[a-z ]+ : [0-9a-f]{12}", "", text)
    numbers = set(re.findall(r"(?<![\w.])\d+(?:[.,]\d+)*(?![\w])", text))
    assert numbers <= _allowed(project), numbers - _allowed(project)
    # the counts the tool knows are nowhere: the sample, each analysis's
    for entry in notebook.entries:
        assert str(entry.computed.n_used) not in numbers
    assert str(len(prepared.table)) not in numbers


def test_the_tool_never_says_the_study_was_preregistered(frozen):
    project = frozen[0]
    for origin in ("primary", "secondary"):
        project.plan.data_origin = origin
        for language in ("fr", "en"):
            text = _text(project, language).lower()
            assert "a été préenregistrée" not in text and "was preregistered" not in text


# --- English, Word and the files. ---


def test_in_english(frozen):
    project = frozen[0]
    project.plan.data_origin = "secondary"
    text = _text(project, "en")
    assert text.startswith("# Preregistration of a secondary data analysis")
    assert "### Q18. Prior knowledge of the data" in text
    assert "[to complete: prior knowledge of the data]" in text
    assert "The α level is set at .05." in text
    assert "Expected direction: Intervention higher than Contrôle." in text
    assert "à compléter" not in text


def test_the_files(frozen, tmp_path):
    project = frozen[0]
    paths = prereg.write_preregistration(project, tmp_path / PATH.name, today=TODAY)
    assert [p.name for p in paths] == [
        "etude-analyses-preenregistrement.docx",
        "etude-analyses-preenregistrement.md",
    ]
    markdown = paths[1].read_bytes()
    assert b"\r\n" not in markdown and markdown.endswith(b"\n")
    document = Document(io.BytesIO(paths[0].read_bytes()))
    texts = [p.text for p in document.paragraphs]
    assert texts[0] == "Préenregistrement selon le gabarit OSF Prereg"
    assert EXAMPLE[0]["statement"] in " ".join(texts)
    holes = [r for p in document.paragraphs for r in p.runs if r.text.startswith("[à compléter")]
    assert holes and all(r.font.highlight_color is not None for r in holes)
    styles = {p.style.name for p in document.paragraphs}
    assert {"Heading 1", "Heading 2", "Heading 3"} <= styles
    assert document.styles["Heading 3"].font.italic
