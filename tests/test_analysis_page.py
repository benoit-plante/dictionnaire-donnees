"""v0.9, step 6: the Analyses page and `ddict analyser` give the same result, the same
report, and an R script that runs next to the database (D-078: only the last analysis is
kept, in memory)."""

import io
import re
import shutil
import subprocess
import zipfile
from pathlib import Path

import pytest
from openpyxl import load_workbook
from streamlit.testing.v1 import AppTest

from ddict import i18n
from ddict.analysis.compute import run
from ddict.analysis.report import r_bundle, render_report, report_xlsx, versions
from ddict.cli import main
from ddict.project import fingerprint, load_project, save_project
from ddict.workflow import analysis_files
from test_analysis_data import EXAMPLES, KEY, _prepared, analyses_project
from test_syntax import _rscript

APP = str(Path(__file__).resolve().parent.parent / "src" / "ddict" / "review" / "app.py")
TODAY = "2026-10-16"
NBSP = chr(0xA0)
T_TEST = {"outcome": ["bienetre_total"], "group": ["groupe"]}


def _folder(tmp_path, name="projet"):
    folder = tmp_path / name
    folder.mkdir()
    shutil.copy(EXAMPLES / KEY["sav"], folder / KEY["sav"])
    path = folder / "etude-analyses.dd.json"
    save_project(analyses_project(), path)
    return path


@pytest.fixture
def project_path(tmp_path):
    return _folder(tmp_path)


@pytest.fixture(scope="module")
def ready():
    project = analyses_project()
    prepared = _prepared(project, KEY["sav"])
    return project, prepared, run(project, prepared, "t_test", T_TEST)


# --- The report. ---


def test_the_report(ready, tmp_path):
    project, prepared, result = ready
    files = analysis_files(tmp_path / "etude-analyses.dd.json")
    assert files.markdown.name == "etude-analyses-analyse-1.md"
    text = render_report(project, prepared, result, files, TODAY)
    assert text.startswith("# Test t : comparer deux groupes\n")
    assert f"- **Empreinte du codebook** : {fingerprint.short(project)}" in text
    assert f"statsmodels {versions()['statsmodels']}" in text
    assert "- **Méthode** : Welch (variances inégales permises)" in text
    assert "- **Variable de groupe** : groupe (Groupe de l'étude)" in text
    assert "- **Variable dépendante** : bienetre_total (score de G1)\n" in text
    assert "- **Observations écartées** : 0 (valeurs manquantes)" in text
    assert result.sentence in text
    assert "| Groupe de l'étude | n | M (ÉT) |" in text
    assert "etude-analyses-analyse-1.R" in text and "etude-analyses-preparation.R" in text
    with i18n.using("en"):
        english = render_report(project, prepared, result, files, TODAY)
    assert "- **Method**: Welch" in english and "Codebook fingerprint" in english


def test_categories_and_exclusions_in_the_report(ready, tmp_path):
    project, prepared, _result = ready
    files = analysis_files(tmp_path / "etude-analyses.dd.json")
    table1 = run(project, prepared, "table1", {"variables": ["region"]})
    text = render_report(project, prepared, table1, files, TODAY)
    assert "| " + NBSP * 4 + "Ouest |" in text  # still under its variable
    assert "## Résultat\n\n###" in text  # Table 1 has no sentence
    chosen = {"outcome": ["stress"], "group": ["niveau_etudes"]}
    two = run(project, prepared, "t_test", chosen, {"levels": "1,3"})
    text = render_report(project, prepared, two, files, TODAY)
    assert "- **Modalités comparées** : Secondaire, Universitaire" in text
    assert "(valeurs manquantes, ou hors des deux groupes comparés)" in text


def test_the_excel_tables(ready, tmp_path):
    project, prepared, result = ready
    files = analysis_files(tmp_path / "etude-analyses.dd.json")
    book = load_workbook(io.BytesIO(report_xlsx(project, prepared, result, files, TODAY)))
    assert book.sheetnames[0] == "Informations"
    about = dict(list(book.worksheets[0].iter_rows(values_only=True))[1:])
    assert about["Résultat"] == result.sentence
    assert about["Empreinte du codebook"] == fingerprint.short(project)
    rows = list(book.worksheets[1].iter_rows(values_only=True))
    assert rows[0] == tuple(result.tables[0].columns)
    assert [list(r) for r in rows[1:]] == result.tables[0].rows


def test_the_next_number(tmp_path):
    path = tmp_path / "etude.dd.json"
    (tmp_path / "etude-analyse-1.xlsx").write_bytes(b"")
    (tmp_path / "etude-analyse-2.R").write_text("", encoding="utf-8")
    assert analysis_files(path).r.name == "etude-analyse-3.R"


@pytest.mark.skipif(_rscript() is None, reason="R is not installed")
def test_the_r_zip_runs_next_to_the_database(ready, tmp_path):
    project, _prepared_data, result = ready
    files = analysis_files(tmp_path / "etude-analyses.dd.json")
    shutil.copy(EXAMPLES / KEY["sav"], tmp_path / KEY["sav"])
    archive = zipfile.ZipFile(io.BytesIO(r_bundle(project, files, "etude-analyses", result, TODAY)))
    assert sorted(archive.namelist()) == [
        "etude-analyses-analyse-1.R",
        "etude-analyses-metadonnees.R",
        "etude-analyses-preparation.R",
    ]
    archive.extractall(tmp_path)
    command = [
        _rscript(),
        "-e",
        'source("etude-analyses-analyse-1.R"); cat(sprintf("T=%.10f", resultats$t))',
    ]
    done = subprocess.run(command, cwd=tmp_path, capture_output=True, text=True,
                          encoding="utf-8", errors="replace")  # fmt: skip
    if "there is no package called" in done.stderr:
        pytest.skip("haven is not installed in R")
    assert done.returncode == 0, done.stderr
    from_r = float(re.search(r"T=([-0-9.]+)", done.stdout).group(1))
    assert from_r == pytest.approx(result.statistics["t"], rel=1e-8)


# --- The command. ---


def test_the_command(project_path, capsys):
    code = main(["analyser", str(project_path), "test_t", "--dependante", "bienetre_total",
                 "--groupe", "groupe"])  # fmt: skip
    assert code == 0
    out = capsys.readouterr().out
    assert "t(297,29) = 5,00" in out
    for name in ("analyse-1.md", "analyse-1.xlsx", "analyse-1.R", "metadonnees.R"):
        assert project_path.with_name(f"etude-analyses-{name}").is_file()
    main(["analyser", str(project_path), "khi2", "--lignes", "teletravail", "--colonnes",
          "secteur", "--methode", "fisher"])  # fmt: skip
    assert "Test exact de Fisher" in project_path.with_name(
        "etude-analyses-analyse-2.md"
    ).read_text(encoding="utf-8")
    assert load_project(project_path).variables == analyses_project().variables  # D-078


def test_the_command_refuses(project_path, capsys):
    code = main(["analyser", str(project_path), "correlation", "--variables", "stress",
                 "id_participant"])  # fmt: skip
    assert code == 2
    assert "identifiant" in capsys.readouterr().err
    assert not project_path.with_name("etude-analyses-analyse-1.md").exists()
    project_path.with_name(KEY["sav"]).unlink()
    code = main(["analyser", str(project_path), "anova", "--dependante", "stress",
                 "--groupe", "niveau_etudes"])  # fmt: skip
    assert code == 2 and "introuvable" in capsys.readouterr().err


# --- The page. ---


def _open(path):
    app = AppTest.from_file(APP, default_timeout=60)
    app.session_state["project_path"] = str(path)
    app.session_state["mode"] = "analyses"
    app.run()
    assert not app.exception, app.exception
    return app


@pytest.mark.allow_hosts(["127.0.0.1"])
def test_the_page_gives_the_command_result(project_path, tmp_path):
    app = _open(project_path)
    assert app.header[0].value == "Analyses"
    assert app.button(key="analysis-run").disabled  # nothing chosen yet, and no error shown
    assert not app.error
    app.selectbox(key="analysis-key").set_value("t_test").run()
    app.selectbox(key="analysis-t_test-outcome").set_value("bienetre_total").run()
    app.selectbox(key="analysis-t_test-group").set_value("groupe").run()
    app.button(key="analysis-run").click().run()
    assert not app.exception, app.exception
    sentence = next(m.value for m in app.markdown if m.value.startswith("> "))
    assert "t(297,29) = 5,00" in sentence
    app.button(key="analysis-write").click().run()
    assert "Fichiers écrits à côté du projet" in app.success[0].value

    other = _folder(tmp_path, "commande")
    main(["analyser", str(other), "test_t", "--dependante", "bienetre_total", "--groupe",
          "groupe"])  # fmt: skip
    for suffix in (".md", ".R"):
        name = f"etude-analyses-analyse-1{suffix}"
        page = project_path.with_name(name).read_text(encoding="utf-8")
        assert page == other.with_name(name).read_text(encoding="utf-8")


@pytest.mark.allow_hosts(["127.0.0.1"])
def test_the_page_refuses_and_warns(project_path):
    app = _open(project_path)
    app.selectbox(key="analysis-key").set_value("anova").run()
    app.selectbox(key="analysis-anova-outcome").set_value("stress").run()
    app.selectbox(key="analysis-anova-group").set_value("region").run()
    assert "moins de 5 personnes dans 4 Ouest" in app.warning[0].value
    app.selectbox(key="analysis-key").set_value("t_test").run()
    app.selectbox(key="analysis-t_test-outcome").set_value("stress").run()
    app.selectbox(key="analysis-t_test-group").set_value("niveau_etudes").run()
    assert "choisissez-en deux" in app.error[0].value
    assert app.button(key="analysis-run").disabled
    app.multiselect(key="analysis-t_test-option-levels-niveau_etudes").set_value([1.0, 3.0]).run()
    assert not app.error and not app.button(key="analysis-run").disabled


@pytest.mark.allow_hosts(["127.0.0.1"])
def test_the_page_without_the_database(project_path):
    project_path.with_name(KEY["sav"]).unlink()
    app = _open(project_path)
    assert "introuvable" in app.info[0].value
    assert app.button(key="analysis-choose")
