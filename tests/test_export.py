"""Step 7: from the fictional database to the codebook, without any AI ("useful without AI")."""

import os
import subprocess
import sys
import zipfile

import pytest
from openpyxl import load_workbook

from ddict.cli import main
from ddict.export import render_markdown, write_xlsx
from ddict.io import read_csv, read_sav
from ddict.project import build_project, correct, load_project, save_project
from ddict.project.model import TrackedValue

AI_MODULES = ("httpx", "openai", "streamlit", "requests")


@pytest.fixture(scope="module")
def run(fictional_files, tmp_path_factory):
    """Run the two commands as a user would."""
    folder = tmp_path_factory.mktemp("cli")
    project = folder / "etude.dd.json"
    assert main(["profiler", str(fictional_files["etude-sommeil.sav"]), "-o", str(project)]) == 0
    assert main(["exporter", str(project)]) == 0
    return {"project": project, "xlsx": folder / "etude.xlsx", "md": folder / "etude.md"}


def _xlsx_text(path):
    with zipfile.ZipFile(path) as archive:
        return "".join(
            archive.read(name).decode("utf-8")
            for name in archive.namelist()
            if name.endswith(".xml")
        )


def _rows(path, sheet):
    return list(load_workbook(path)[sheet].iter_rows(values_only=True))


def test_commands_produce_the_three_files(run):
    for path in run.values():
        assert path.exists() and path.stat().st_size > 0


def test_whole_chain_works_with_ai_libraries_missing(fictional_files, tmp_path):
    """In a separate process where importing an AI library fails, as if not installed."""
    script = (
        "import sys\n"
        f"for name in {AI_MODULES!r}:\n"
        "    sys.modules[name] = None\n"
        "from ddict.cli import main\n"
        "base, project = sys.argv[1], sys.argv[2]\n"
        "assert main(['profiler', base, '-o', project]) == 0\n"
        "assert main(['exporter', project]) == 0\n"
        "assert main(['enrichir', project, '--apercu']) == 0\n"
    )
    project = tmp_path / "sans-ia.dd.json"
    result = subprocess.run(
        [sys.executable, "-c", script, str(fictional_files["etude-sommeil.sav"]), str(project)],
        capture_output=True,
        text=True,
        encoding="utf-8",
        env={**os.environ, "PYTHONIOENCODING": "utf-8"},
        check=False,
    )
    assert result.returncode == 0, result.stderr
    assert (tmp_path / "sans-ia.xlsx").exists()


@pytest.mark.parametrize("output", ["project", "xlsx", "md"])
def test_no_canary_in_any_output(run, generator, output):
    path = run[output]
    text = _xlsx_text(path) if output == "xlsx" else path.read_text(encoding="utf-8")
    for canary in generator.CANARIES:
        assert canary not in text


def test_xlsx_sheets_and_variables(run):
    book = load_workbook(run["xlsx"])
    assert book.sheetnames == ["Informations", "Variables", "Valeurs", "Groupes", "Provenance"]
    rows = _rows(run["xlsx"], "Variables")
    header, body = rows[0], rows[1:]
    assert len(body) == 34
    by_name = {row[1]: dict(zip(header, row, strict=True)) for row in body}
    q1 = by_name["sommeil_q1"]
    assert q1["Libellé"] == "Q1 endormissement"
    assert q1["Type"] == "entier"
    assert q1["Niveau de mesure"] == "ordinal"
    assert "Code 9 observé sans étiquette" in by_name["sommeil_q3"]["Signalements"]
    assert "Valeur 999 probablement manquante" in by_name["age"]["Signalements"]
    assert "Identique à heures_sommeil" in by_name["heures_sommeil_2"]["Signalements"]
    courriel = by_name["courriel"]
    assert courriel["Confidentialité"].startswith("identifiant probable")
    assert courriel["Valeurs distinctes"] is None
    assert by_name["genre"]["Moyenne"] is None


def test_xlsx_values_sheet_masks_small_counts(run):
    rows = [r for r in _rows(run["xlsx"], "Valeurs") if r[0] == "genre"]
    assert rows == [
        ("genre", "1", "Femme", "114", None, None),
        ("genre", "2", "Homme", "masqué", None, None),
        ("genre", "3", "Autre identité", "masqué", None, None),
        ("genre", "9", "Préfère ne pas répondre", "4", "oui", None),
    ]
    unused = [r for r in _rows(run["xlsx"], "Valeurs") if r[0] == "satisfaction_globale"]
    assert ("satisfaction_globale", "99", "Ne sait pas", "0", None, None) in unused


def test_markdown_report(run):
    text = run["md"].read_text(encoding="utf-8")
    assert text.startswith("# Dictionnaire de données : etude-sommeil.sav\n")
    assert "### `sommeil_q1` — Q1 endormissement" in text
    assert "| 3 | Autre identité | masqué |  |" in text
    assert "Variable sensible : aucune statistique sur ses valeurs." in text
    assert "\r\n" not in run["md"].read_bytes().decode("utf-8")


def test_markdown_is_deterministic(run):
    project = load_project(run["project"])
    assert render_markdown(project) == render_markdown(project)
    assert render_markdown(project) == run["md"].read_text(encoding="utf-8")


@pytest.fixture
def with_ai(fictional_files):
    """The fictional project with an AI proposal pending and another one corrected."""
    project = build_project(read_sav(fictional_files["etude-sommeil.sav"]))
    q7 = project.variable("sommeil_q7")
    q7.label = TrackedValue[str](
        value="Routine du coucher", origin="ai", status="proposed", ai_run="run-0001"
    )
    q4 = project.variable("sommeil_q4")
    q4.description = correct(
        TrackedValue[str](value="Proposition", origin="ai", status="proposed", ai_run="run-0001"),
        "Satisfaction à l'égard du sommeil.",
        now="2026-10-13T14:00:00",
    )
    return project


def test_final_export_leaves_out_pending_proposals(with_ai):
    text = render_markdown(with_ai, kind="final")
    assert "Routine du coucher" not in text
    assert "### `sommeil_q7`\n" in text
    assert "Satisfaction à l'égard du sommeil." in text


def test_work_export_marks_pending_proposals(with_ai, tmp_path):
    text = render_markdown(with_ai, kind="work")
    assert "### `sommeil_q7` — Routine du coucher [proposé par IA]" in text
    path = tmp_path / "travail.xlsx"
    write_xlsx(with_ai, path, kind="work")
    provenance = _rows(path, "Provenance")
    assert (
        "sommeil_q7",
        "Libellé",
        "Routine du coucher",
        "IA",
        "proposé par IA",
        None,
        "run-0001",
    ) in provenance
    assert (
        "sommeil_q4",
        "Description",
        "Satisfaction à l'égard du sommeil.",
        "IA",
        "corrigé",
        "Proposition",
        "run-0001",
    ) in provenance


def test_formula_like_labels_stay_text(with_ai, tmp_path):
    with_ai.variable("age").label = TrackedValue[str](value="=1+1", origin="file")
    path = tmp_path / "formule.xlsx"
    write_xlsx(with_ai, path)
    sheet = load_workbook(path)["Variables"]
    cell = next(row[2] for row in sheet.iter_rows(min_row=2) if row[1].value == "age")
    assert cell.value == "=1+1"
    assert cell.data_type == "s"


def test_profiler_refuses_to_overwrite(run, fictional_files, capsys):
    code = main(["profiler", str(fictional_files["etude-sommeil.sav"]), "-o", str(run["project"])])
    assert code == 2
    assert "--remplacer" in capsys.readouterr().err


def test_profiler_default_output_next_to_the_database(fictional_files, tmp_path):
    database = tmp_path / "ma base.csv"
    database.write_bytes(fictional_files["etude-sommeil_excel-fr.csv"].read_bytes())
    assert main(["profiler", str(database)]) == 0
    assert (tmp_path / "ma base.dd.json").exists()
    assert main(["exporter", str(tmp_path / "ma base.dd.json"), "--travail"]) == 0
    assert (tmp_path / "ma base-travail.xlsx").exists()


def test_readable_errors(tmp_path, capsys):
    assert main(["profiler", str(tmp_path / "absente.sav")]) == 2
    assert "Impossible de lire la base" in capsys.readouterr().err
    broken = tmp_path / "abime.dd.json"
    broken.write_text("{", encoding="utf-8")
    assert main(["exporter", str(broken)]) == 2
    assert "Impossible de lire le projet" in capsys.readouterr().err


def test_saved_project_with_ai_round_trips(with_ai, tmp_path):
    path = tmp_path / "p.dd.json"
    save_project(with_ai, path)
    assert load_project(path) == with_ai


def test_missing_tokens_are_reported(tmp_path):
    """D-010: the codebook says which cells were read as empty."""
    path = tmp_path / "r.csv"
    path.write_text("score,groupe\n" + "2,a\nNA,b\n" * 6, encoding="utf-8")
    report = render_markdown(build_project(read_csv(path)), kind="final")
    assert "Codes manquants déclarés : « NA » lu comme vide (6)" in report
