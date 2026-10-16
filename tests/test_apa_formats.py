"""v0.10, step 3: the test every format of the APA tables and figures must pass (D-083),
and the APA 7 layout of the Word file (D-084, D-085).

The common test reads each format's file back and finds in it every number, header, title
and note of the description: a new format registered in FORMATS passes it, or fails."""

import io
import zipfile

import pytest
from docx import Document
from docx.oxml.ns import qn
from PIL import Image

from ddict.analysis.apa import Edit, describe, edited, plain
from ddict.analysis.compute import run
from ddict.analysis.formats import FORMATS, available
from ddict.analysis.formats.figures import image
from test_analysis_data import KEY, _prepared, analyses_project

CASES = {
    "table1": ({"variables": ["age", "sexe", "niveau_etudes", "bienetre_total"],
                "group": ["groupe"]}, None),
    "t_test": ({"outcome": ["bienetre_total"], "group": ["groupe"]}, None),
    "mann_whitney": ({"outcome": ["bienetre_total"], "group": ["groupe"]},
                     {"method": "mann_whitney"}),
    "anova": ({"outcome": ["stress"], "group": ["niveau_etudes"]}, {"post_hoc": "tukey"}),
    "kruskal": ({"outcome": ["stress"], "group": ["niveau_etudes"]},
                {"method": "kruskal_wallis"}),
    "chi_square": ({"rows": ["teletravail"], "columns": ["secteur"]}, None),
    "correlation": ({"variables": ["heures_travail", "stress", "soutien"]}, None),
    "scatter": ({"variables": ["heures_travail", "stress"]}, None),
    "regression": ({"outcome": ["satisfaction"], "predictors": ["soutien", "secteur"]}, None),
}  # fmt: skip
KEYS = {"mann_whitney": "t_test", "kruskal": "anova", "scatter": "correlation"}


@pytest.fixture(scope="module")
def ready():
    project = analyses_project()
    return project, _prepared(project, KEY["sav"])


def _document(ready, case, language="fr"):
    project, prepared = ready
    chosen, options = CASES[case]
    result = run(project, prepared, KEYS.get(case, case), chosen, options)
    return describe(project, prepared, result, language=language)


def _flat(text: str) -> str:
    return " ".join(text.replace(chr(0xA0), " ").split())


def _expected(document, found) -> list[str]:
    """Every text of the description a format renders."""
    texts = []
    if found.tables:
        for table in document.tables:
            texts += [table.title, plain(table.note)]
            texts += [plain(c.header) for c in table.columns]
            texts += [plain(s.label) for s in table.spanners]
            texts += [plain(cell) for row in table.rows for cell in row.cells]
    if found.figures:
        for figure in document.figures:
            texts += [figure.title, plain(figure.note)]
    return [_flat(t) for t in texts if t.strip()]


# --- The test common to every format. ---


@pytest.mark.parametrize("language", ["fr", "en"])
@pytest.mark.parametrize("case", list(CASES))
@pytest.mark.parametrize("key", list(FORMATS))
def test_every_format_renders_everything(ready, key, case, language):
    found = FORMATS[key]
    document = _document(ready, case, language)
    if found not in available(document):
        pytest.skip(f"{key}: nothing to render for {case}")
    text = _flat(found.read_text(found.render(document)))
    missing = [t for t in _expected(document, found) if t not in text]
    assert not missing, missing


@pytest.mark.parametrize("key", list(FORMATS))
def test_every_format_takes_a_persons_edits(ready, key):
    found = FORMATS[key]
    document = edited(_document(ready, "anova"),
                      {"table-1": Edit(number=7, title="Le stress", note="Ma note."),
                       "figure-1": Edit(number=3, title="Les moyennes")})  # fmt: skip
    text = _flat(found.read_text(found.render(document)))
    if found.tables:
        assert "Le stress" in text and "Ma note." in text
    if found.figures:
        assert "Les moyennes" in text


def test_the_formats_of_this_version():
    assert list(FORMATS) == ["docx"]  # D-085
    assert FORMATS["docx"].tables and FORMATS["docx"].figures


# --- Word, APA 7. ---


def _word(ready, case, language="fr"):
    return Document(io.BytesIO(FORMATS["docx"].render(_document(ready, case, language))))


def _borders(cell) -> set[str]:
    found = cell._tc.tcPr.find(qn("w:tcBorders"))
    return set() if found is None else {b.tag.split("}")[1] for b in found}


def test_word_tables_have_horizontal_rules_only(ready):
    word = _word(ready, "t_test")
    table = word.tables[0]
    spanners, header, data = table.rows
    for cell in spanners.cells:
        assert "top" in _borders(cell)
    assert "bottom" in _borders(spanners.cells[1])  # the short rule under "Intervention"
    for cell in header.cells:
        assert "bottom" in _borders(cell)
    for cell in data.cells:
        assert _borders(cell) == {"bottom"}  # no vertical rule anywhere
    grid = table._tbl.tblPr.find(qn("w:tblBorders"))
    assert all(b.get(qn("w:val")) == "nil" for b in grid)


def test_word_captions_symbols_and_notes(ready):
    word = _word(ready, "t_test")
    paragraphs = word.paragraphs
    number, title = paragraphs[0], paragraphs[1]
    assert number.text == "Tableau 1" and number.runs[0].bold
    assert title.runs[0].italic and title.text.startswith("bienetre_total")
    header = word.tables[0].rows[1].cells
    assert [r.text for r in header[1].paragraphs[0].runs if r.italic] == ["M"]
    assert header[0].paragraphs[0].alignment == 0  # left
    note = next(p for p in paragraphs if p.text.startswith("Note."))
    assert note.runs[0].italic and note.runs[0].text == "Note."
    assert "Figure 1" in [p.text for p in paragraphs]
    assert len(word.inline_shapes) == 1  # the figure, as an image
    assert "Empreinte du codebook" in paragraphs[-1].text


def test_word_in_english(ready):
    word = _word(ready, "regression", "en")
    texts = [p.text for p in word.paragraphs]
    assert texts[0] == "Table 1" and "Figure 1" in texts
    assert word.tables[0].rows[1].cells[0].text == "Intercept"


def test_numbers_never_wrap_where_they_fit(ready):
    """The columns of numbers get the width their numbers need (Word read in the browser
    trial: "14,2 / 4" before this rule)."""
    word = _word(ready, "regression")
    table = word.tables[0]
    widths = [cell.width.cm for cell in table.rows[0].cells]
    assert sum(widths) == pytest.approx(16.5, abs=0.01)
    assert widths[0] == max(widths)  # the labels
    assert min(widths[1:]) >= 1.2


def test_the_figure_images(ready):
    for case in ("t_test", "chi_square", "scatter", "regression"):
        figure = _document(ready, case).figures[0]
        png = Image.open(io.BytesIO(image(figure, "fr")))
        assert png.format == "PNG" and png.size[0] == pytest.approx(1890, abs=2)  # 16 cm, 300 dpi
    svg = image(_document(ready, "t_test").figures[0], "fr", kind="svg").decode("utf-8")
    assert "Intervention" in svg and "17,5" in svg  # a decimal comma on the axis


def test_word_is_a_valid_office_file(ready):
    content = FORMATS["docx"].render(_document(ready, "anova"))
    with zipfile.ZipFile(io.BytesIO(content)) as archive:
        assert "word/document.xml" in archive.namelist()
        assert any(n.startswith("word/media/") for n in archive.namelist())
