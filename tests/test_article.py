"""v0.13, step 5: the draft of an IMRaD article (D-106, D-107), as the validated sheet
docs/plan-analyse-article.md (§ 9 to 13) gives it: the title page and the abstract in holes,
the Introduction with the hypotheses, the paragraph "Plan d'analyse" of the Method, the
Results, the Discussion and the References, drawn from the sources only."""

import io
from pathlib import Path

import pytest
from docx import Document
from docx.enum.text import WD_COLOR_INDEX

from ddict.analysis import article_text as article
from ddict.analysis import notebook as nb
from ddict.analysis import notebook_check as check
from ddict.analysis import notebook_export as export
from ddict.analysis import plan_notebook as pn
from ddict.analysis import planning as pl
from ddict.analysis.apa import plain
from ddict.analysis.book import Codebook
from ddict.analysis.compute import run
from test_analysis_data import KEY, _prepared, analyses_project
from test_plan import EXAMPLE
from test_plan_notebook import LATER, _frozen

REASON = "Distribution des heures de travail très asymétrique."


def _confirmed(project, prepared, notebook):
    """What waits for a person (the adjusted p after an amendment) is confirmed."""
    check.check_all(notebook, project, prepared, today=LATER)
    for entry in notebook.entries:
        if entry.pending is not None:
            check.confirm(notebook, entry.id)


@pytest.fixture
def amended():
    """The sheet's example: the plan frozen, then H2 amended to Spearman."""
    project = analyses_project()
    prepared = _prepared(project, KEY["sav"])
    notebook = nb.new_notebook(project, Path("etude-analyses.dd.json"))
    _frozen((project, prepared, notebook))
    pl.modify(project, "H2", {"options": {"method": "spearman"}}, prepared, REASON, LATER)
    pn.sync(notebook, project, prepared, LATER)
    _confirmed(project, prepared, notebook)
    return project, prepared, notebook


def _document(ready, language=None):
    return export.results_document(*ready, language=language, today=LATER)


def _section(document, key):
    return next(s for s in document.sections if s.key == key)


def _texts(document, key):
    return [plain(p) for p in _section(document, key).paragraphs]


# --- The order of the article (§ 9). ---


def test_the_parts_in_order(amended):
    document = _document(amended)
    top = [(s.key, s.level) for s in document.sections if s.level == 1]
    assert top == [("titre", 1), ("resume", 1), ("introduction", 1), ("methode", 1),
                   ("resultats", 1), ("discussion", 1), ("references", 1)]  # fmt: skip
    method = [s.key for s in document.sections][4:8]
    assert method == ["participants", "instruments", "analyses", "plan"]
    assert _texts(document, "titre") == [
        "[à compléter : titre]", "[à compléter : auteurs et affiliations]",
    ]  # fmt: skip
    assert _texts(document, "resume") == [
        "[à compléter : résumé]", "Mots-clés : [à compléter : mots-clés]",
    ]  # fmt: skip
    assert _section(document, "introduction").heading == "[à compléter : titre]"


def test_the_introduction_holds_the_hypotheses(amended):
    assert _texts(_document(amended), "introduction") == [
        "[à compléter : contexte, revue des travaux et justification de l'étude]",
        "Les hypothèses suivantes ont été formulées :",
        *[f"H{i}. {EXAMPLE[i - 1]['statement']}" for i in range(1, 5)],
        "Les questions suivantes ont été examinées de façon exploratoire :",
        f"H5. {EXAMPLE[4]['statement']}",
    ]


# --- The paragraph "Plan d'analyse" (§ 10). ---


def test_the_plan_paragraph_as_the_sheet(amended):
    [text] = _texts(_document(amended), "plan")
    assert text == (
        "Les données sont primaires, recueillies par l'équipe pour cette étude. Selon "
        "l'équipe, le plan d'analyse a été écrit après avoir accédé aux données, avant de les "
        "analyser. Il a été arrêté dans l'outil le 20 octobre 2026. [à compléter : lieu et "
        "identifiant du préenregistrement, s'il y en a un] Il comptait quatre hypothèses "
        "confirmatoires et une exploratoire. Le seuil α était fixé à ,05. Les valeurs de p ont "
        "été corrigées par la méthode de Holm (1979) au sein de chaque famille d'hypothèses : "
        "F1 (H1, H2, H3). Le test de H1 était unilatéral, dans le sens attendu. Le plan a été "
        "modifié une fois après avoir été arrêté : le 2 novembre 2026, pour H2 (méthode : "
        "Spearman au lieu de Pearson) : distribution des heures de travail très asymétrique. "
        "Le jour où le plan a été arrêté, le carnet d'analyses contenait déjà l'analyse prévue "
        "pour H3. Les autres analyses sont présentées comme exploratoires."
    )


def test_the_plan_paragraph_in_english(amended):
    [text] = _texts(_document(amended, "en"), "plan")
    assert text.startswith(
        "The data are primary, collected by the team for this study. According to the team, "
        "the analysis plan was written after accessing the data, before analysing them. It "
        "was finalised in the tool on 20 October 2026."
    )
    assert "It included four confirmatory hypotheses and one exploratory." in text
    assert "The plan was modified once after being finalised: on 2 November 2026, for H2 " in text


def test_the_summary_of_each_change(amended):
    project, prepared, _notebook = amended
    pl.modify(project, "H4", {"roles": {"outcome": EXAMPLE[3]["roles"]["outcome"],
                                        "predictors": [{"variable": "soutien"}]}},
              prepared, "Charge retirée.", LATER)  # fmt: skip
    pl.modify(project, "H3", {"statement": "Le stress varie selon le niveau d'études."},
              prepared, "Reformulée.", LATER)  # fmt: skip
    pl.modify(project, "H5", {"kind": "confirmatory"}, prepared, "Revue.", LATER)
    pl.remove(project, "H5", "Retirée.", LATER)
    pl.set_rules(project, {"alpha": 0.01, "correction": "bonferroni"}, "Plus strict.", LATER)
    summaries = [article.amendment_summary(project, a) for a in project.plan.amendments]
    assert summaries == [
        "méthode : Spearman au lieu de Pearson",
        "prédicteurs : soutien au lieu de soutien, charge",
        "énoncé reformulé",
        "confirmatoire au lieu de exploratoire",
        "hypothèse retirée",
        "seuil α : ,01 au lieu de ,05, correction : Bonferroni au lieu de Holm",
    ]
    assert article.amendment_line(project, project.plan.amendments[-1]).startswith(
        "le 2 novembre 2026, pour les règles générales (seuil α"
    )


def test_no_correction_no_sentence_and_no_holm(amended):
    project, prepared, notebook = amended
    pl.set_rules(project, {"correction": "none"}, "Une seule famille.", LATER)
    pn.sync(notebook, project, prepared, LATER)
    _confirmed(project, prepared, notebook)
    document = _document(amended)
    [text] = _texts(document, "plan")
    assert "corrigées" not in text
    assert (
        "le 2 novembre 2026, pour les règles générales (correction : aucune au lieu de Holm)"
    ) in text
    assert "modifié deux fois" in text
    assert not any(t.startswith("Holm, S.") for t in _texts(document, "references"))


def test_the_closing_lines_hold_the_plan(amended):
    closing = list(_document(amended).closing)
    project = amended[0]
    assert (
        "Plan arrêté dans l'outil le 20 octobre 2026 ; empreinte du plan : "
        f"{project.plan.frozen.fingerprint[:12]}."
    ) in closing
    start = closing.index("Amendements du plan :")
    assert closing[start + 1].startswith("le 2 novembre 2026, pour H2 (méthode : Spearman")


# --- Without a frozen plan (§ 9, 11, 12). ---


def test_without_a_frozen_plan():
    project = analyses_project()
    prepared = _prepared(project, KEY["sav"])
    notebook = nb.new_notebook(project, Path("etude-analyses.dd.json"))
    from ddict.analysis.compute import run

    chosen = {"outcome": ["bienetre_total"], "group": ["groupe"]}
    nb.add(notebook, project, run(project, prepared, "t_test", chosen), today=LATER)
    for data in EXAMPLE:
        pl.add(project, data, prepared)  # a plan, not frozen
    document = export.results_document(project, prepared, notebook, today=LATER)
    assert "plan" not in [s.key for s in document.sections]
    assert _texts(document, "introduction") == [
        "[à compléter : contexte, revue des travaux et justification de l'étude]"
    ]
    assert _texts(document, "discussion") == [
        "[à compléter : rappel du but de l'étude]",
        "[à compléter : discussion des résultats]",
        "[à compléter : limites]",
        "[à compléter : conclusion]",
    ]
    assert not any("Plan arrêté" in line for line in document.closing)


# --- The References (§ 13): the sources, as written, and what the text cites. ---


def test_the_references(amended):
    document = _document(amended)
    found = _texts(document, "references")
    assert found[0].startswith("[à vérifier : les sources ci-dessous viennent du codebook")
    assert _section(document, "references").paragraphs[0][0].hole
    assert found[1:] == [
        "Feldt, L. S. (1965). The approximate sampling distribution of Kuder-Richardson "
        "reliability coefficient twenty. Psychometrika, 30(3), 357–370. "
        "https://doi.org/10.1007/BF02289499",
        "Holm, S. (1979). A simple sequentially rejective multiple test procedure. "
        "Scandinavian Journal of Statistics, 6(2), 65–70.",
        "Manuel fictif de l'échelle de sommeil, p. 4 (référence inventée pour l'exemple)",
    ]
    # Feldt is listed because I6 cites it.
    assert "selon Feldt, 1965" in " ".join(_texts(document, "instruments"))
    italics = [r.text for r in _section(document, "references").paragraphs[2] if r.italic]
    assert italics == ["Scandinavian Journal of Statistics", "6"]


def test_every_reference_comes_from_a_source(amended):
    """Each line is a source of the codebook, as the team wrote it, once, sorted without
    accents nor case; or one of the tool's two fixed references."""
    project, _prepared_data, notebook = amended
    group = project.groups[0]
    from ddict.project.model import SourcedText, Threshold

    group.instrument_name = SourcedText(value="Échelle fictive", source="Zeta, A. (2020)")
    group.thresholds = [Threshold(value=15, meaning="élevé", source="Ébert, B. (2019)")]
    project.variable("stress").thresholds = [
        Threshold(value=30, meaning="élevé", source="Zeta, A. (2020)")  # once only
    ]
    project.variable("secteur").thresholds = []  # a variable the article never names
    found = article.sources(project, notebook)
    assert found == [
        "Ébert, B. (2019)",
        "Manuel fictif de l'échelle de sommeil, p. 4 (référence inventée pour l'exemple)",
        "Zeta, A. (2020)",
    ]
    names = [plain(p) for p in article.references(project, notebook, feldt=False)[1:]]
    assert names == [found[0], "Holm, S. (1979). A simple sequentially rejective multiple "
                     "test procedure. Scandinavian Journal of Statistics, 6(2), 65–70.",
                     *found[1:]]  # fmt: skip


# --- Word. ---


def test_the_word_file(amended, tmp_path):
    project, prepared, notebook = amended
    path = tmp_path / "etude-analyses.dd.json"
    written = export.write_results(project, prepared, notebook, path, today=LATER)
    assert [p.name for p in written] == [
        "etude-analyses-article.docx",
        "etude-analyses-article-graphiques.xlsx",
        "etude-analyses-article-R.zip",
    ]
    word = Document(io.BytesIO(written[0].read_bytes()))
    paragraphs = word.paragraphs
    assert paragraphs[0].text == "[à compléter : titre]"
    title = next(p for p in paragraphs if p.style.name == "Heading 1"
                 and p.text == "[à compléter : titre]")  # fmt: skip
    assert all(r.font.highlight_color == WD_COLOR_INDEX.YELLOW for r in title.runs)
    breaks = sum(
        1 for p in paragraphs for r in p.runs for b in r._element.findall(
            "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}br")
        if b.get("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}type") == "page"
    )  # fmt: skip
    assert breaks == 3  # after the title page and the abstract, before the References
    headings = [p.text for p in paragraphs if p.style.name == "Heading 1"]
    assert headings == ["Résumé", "[à compléter : titre]", "Méthode", "Résultats",
                        "Discussion", "Références"]  # fmt: skip
    import zipfile

    with zipfile.ZipFile(written[2]) as archive:
        assert "etude-analyses-article.R" in archive.namelist()
    discussion = paragraphs[[p.text for p in paragraphs].index("Discussion") + 2]
    assert discussion.runs[0].text == "H1." and discussion.runs[0].bold


# --- The Results ordered by the plan (§ 11), the adjusted p (§ 7). ---


def _results(document):
    keys = [s.key for s in document.sections]
    return document.sections[keys.index("resultats") + 1 : keys.index("discussion")]


def test_the_results_by_hypothesis(amended):
    document = _document(amended)
    found = [(s.key, s.level) for s in _results(document)]
    assert found == [("A1", 2), ("H1", 2), ("H2", 2), ("H3", 2), ("H4", 2),
                     ("exploratoires", 2), ("H5", 3), ("A4", 3)]  # fmt: skip
    h1 = _section(document, "H1")
    assert h1.heading == f"H1 — {EXAMPLE[0]['statement']}"
    assert "t(297,29) = 5,00, p < ,001, p corrigé (Holm) < ,001, d = 0,58" in h1.sentence
    h2 = _section(document, "H2").sentence
    assert "ρ = ,42, p < ,001, p corrigé (Holm) < ,001 (n = 296)." in h2
    # A predictor: the model's sentence, then the effect's; H4 is in no family.
    h4 = _section(document, "H4").sentence
    assert h4.endswith(
        "Soutien des collègues (0 à 10) : b = 0,51, IC 95 % [0,44 ; 0,58], β = ,60, "
        "t(297) = 14,24, p < ,001."
    )
    assert "corrigé" not in h4
    assert _section(document, "exploratoires").heading == "Analyses exploratoires"
    # The identifiers stay at the end, to find each analysis in the R script.
    assert any(line.startswith("A5 — Test t") for line in document.closing)


def test_the_adjusted_p_in_the_tables(amended):
    document = _document(amended)
    [table] = _section(document, "H1").tables
    headers = [plain(c.header) for c in table.columns]
    at = headers.index("p")
    assert headers[at + 1] == "p corrigé (Holm)"
    assert [plain(c) for c in table.rows[0].cells][at : at + 2] == ["< ,001", "< ,001"]
    assert [s.first for s in table.spanners] == [1, 3]  # the spanners left of p stay
    assert plain(table.note).endswith(
        "p corrigé par la méthode de Holm dans la famille F1 (3 hypothèses)."
    )
    anova = _section(document, "H3").tables[0]
    assert plain(anova.note).endswith(
        "p corrigé par la méthode de Holm dans la famille F1 (3 hypothèses) : < ,001."
    )
    assert not any("corrigé" in plain(t.note) for t in _section(document, "H4").tables)


def test_the_adjusted_p_on_the_line_of_a_predictor(amended):
    project, prepared, _notebook = amended
    from ddict.analysis.apa import describe
    from ddict.analysis.compute import run

    chosen = {"outcome": ["satisfaction"], "predictors": ["soutien", "charge"]}
    result = run(project, prepared, "regression", chosen)
    h4 = project.plan.hypotheses[3].model_copy(update={"family": "F2"})
    tested = article.Tested(h4, result, "p_2", 0.0412, 2, "charge")
    document = article.adjusted_tables(describe(project, prepared, result), tested, "holm")
    table = document.tables[0]
    at = [plain(c.header) for c in table.columns].index("p corrigé (Holm)")
    assert [plain(r.cells[at]) for r in table.rows] == ["—", "—", ",041"]
    sentence = article.with_adjusted("t = 2,10, p = ,020, d = 0,40.", 0.02, 0.04, "bonferroni")
    assert sentence == "t = 2,10, p = ,020, p corrigé (Bonferroni) = ,040, d = 0,40."


def test_a_variant_and_a_deviation(amended):
    project, prepared, notebook = amended
    from ddict.analysis.compute import run

    chosen = {"outcome": ["bienetre_total"], "group": ["groupe"]}
    options = {"method": "student", "alternative": "greater"}
    nb.add(notebook, project, run(project, prepared, "t_test", chosen, options), today=LATER)
    nb.retain(notebook, next(e.id for e in notebook.entries if e.hypothesis == "H2"), False)
    _confirmed(project, prepared, notebook)
    document = _document(amended)
    assert [s.key for s in _results(document)][-1] == "A9"
    assert _section(document, "A9").sentence.startswith(
        "Variante de H1 : méthode : Student (variances égales) au lieu de Welch (variances "
        "inégales permises). bienetre_total (score de G1)"
    )
    [text] = _texts(document, "H2")
    assert text == (
        "H2 : l'analyse prévue n'est pas retenue dans le carnet. [à compléter : raison de "
        "l'écart au plan pour H2]"
    )
    [plan] = _texts(document, "plan")
    assert "La famille F1 comptait trois hypothèses ; deux ont pu être testées." in plan
    discussion = _texts(document, "discussion")
    assert discussion[2] == (
        f"H2. {EXAMPLE[1]['statement']} H2 : l'analyse prévue n'est pas retenue dans le "
        "carnet. [à compléter : raison de l'écart au plan pour H2] [à compléter : "
        "interprétation de H2]"
    )


# --- The Discussion (§ 12), and no word that judges (D-107). ---


def test_the_discussion_as_the_sheet(amended):
    document = _document(amended)
    found = _texts(document, "discussion")
    assert found[0] == "[à compléter : rappel du but de l'étude]"
    assert found[1] == (
        f"H1. {EXAMPLE[0]['statement']} Le sens attendu était : Intervention plus élevé que "
        "Contrôle. La moyenne observée était plus élevée dans le groupe Intervention. "
        "bienetre_total (score de G1) : M = 17,01 (ÉT = 4,18) pour Intervention et M = 14,54 "
        "(ÉT = 4,39) pour Contrôle, t(297,29) = 5,00, p < ,001, p corrigé (Holm) < ,001, d = "
        "0,58, IC 95 % [0,35 ; 0,81]. [à compléter : interprétation de H1]"
    )
    assert found[2].startswith(
        f"H2. {EXAMPLE[1]['statement']} Le sens attendu était : une corrélation positive. La "
        "corrélation observée était positive. La corrélation de Spearman"
    )
    assert found[3].startswith(f"H3. {EXAMPLE[2]['statement']} Score de stress")  # no direction
    assert "Le coefficient observé de Soutien des collègues (0 à 10) était positif." in found[4]
    assert found[5].startswith("H5. ")  # the exploratory hypotheses after
    assert found[-3:] == [
        "[à compléter : interprétation des analyses exploratoires]",
        "[à compléter : limites]",
        "[à compléter : conclusion]",
    ]
    bold = [r.text for r in _section(document, "discussion").paragraphs[1] if r.bold]
    assert bold == ["H1."]


def test_a_direction_shown_as_zero_goes_unsaid(amended):
    """Means equal to two decimals, as printed beside them, have no higher group (step 7,
    seen in the browser: 23,35 and 23,35)."""
    project, prepared, _notebook = amended
    result = run(project, prepared, "t_test", {"outcome": ["bienetre_total"], "group": ["groupe"]})
    result.statistics |= {"m1": 23.351, "m2": 23.349}
    book = Codebook(project, prepared)
    tested = article.Tested(project.plan.hypotheses[0], result, "p")
    assert article.observed(book, prepared, tested) == (
        "Les moyennes observées étaient égales au centième près."
    )
    result.statistics["m2"] = 23.34
    assert "plus élevée dans le groupe" in article.observed(book, prepared, tested)


FORBIDDEN = [
    "soutenu", "appuyé", "confirmé", "infirmé", "rejeté", "réfuté", "significati",
    "non significatif", "tendance", "marginal", "prouve", "démontre",
    "supported", "confirmed", "rejected", "refuted", "significant", "trend", "proves",
    "demonstrates",
]  # fmt: skip


def _written(document) -> str:
    """Everything the tool wrote, holes excluded."""
    parts = []
    for section in document.sections:
        parts += [section.heading, section.sentence]
        parts += ["".join(r.text for r in p if not r.hole) for p in section.paragraphs]
        for table in section.tables:
            parts += [table.title, plain(table.note), *[plain(c.header) for c in table.columns]]
            parts += [plain(cell) for row in table.rows for cell in row.cells]
        for figure in section.figures:
            parts += [figure.title, plain(figure.note)]
    parts += list(document.closing)
    return " ".join(parts).lower()


@pytest.mark.parametrize("language", ["fr", "en"])
def test_no_word_that_judges(amended, language):
    text = _written(_document(amended, language))
    found = [w for w in FORBIDDEN if w in text]
    assert not found, found


# --- The common test of the formats (D-083), on the article. ---


@pytest.mark.parametrize("language", ["fr", "en"])
def test_every_format_renders_the_article(amended, language):
    from ddict.analysis.formats import available

    document = _document(amended, language)
    for found in available(document):
        text = " ".join(found.read_text(found.render(document)).replace(chr(0xA0), " ").split())
        expected = []
        if found.tables:
            for table in document.tables:
                expected += [table.title, plain(table.note)]
                expected += [plain(c.header) for c in table.columns]
                expected += [plain(cell) for row in table.rows for cell in row.cells]
            expected += [s.heading for s in document.sections]
            expected += [s.sentence for s in document.sections]
            expected += [plain(p) for s in document.sections for p in s.paragraphs]
            expected += list(document.closing)
        if found.figures:
            for figure in document.figures:
                expected += [figure.title, plain(figure.note)]
        flat = [" ".join(t.replace(chr(0xA0), " ").split()) for t in expected if t.strip()]
        missing = [t for t in flat if t not in text]
        assert not missing, (found.key, missing)
