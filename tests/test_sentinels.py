"""v0.4, step 2: suspect missing codes confirmed by a person (D-028)."""

import json
from pathlib import Path

import httpx
import pytest
from openpyxl import load_workbook
from streamlit.testing.v1 import AppTest

from ddict.export import render_markdown, write_xlsx
from ddict.export import wording as w
from ddict.io import read_database
from ddict.llm.enrich import enrich
from ddict.llm.messages import build_blocks
from ddict.project import build_project, load_project, save_project
from ddict.review import logic
from ddict.workflow import tasks
from fake_ai import institution

NOW = "2026-10-14T15:00:00"
APP = str(Path(__file__).resolve().parent.parent / "src" / "ddict" / "review" / "app.py")


@pytest.fixture
def project(fictional_files):
    return build_project(read_database(fictional_files["etude-sommeil.sav"]))


def test_suspect_codes_of_the_fictional_base(project):
    suspects = {e.name: e.profile.suspect_missing_codes for e in project.variables}
    assert {n: c for n, c in suspects.items() if c} == {
        "age": [999],
        "sommeil_q3": [9],
        "nb_reveils": [-9],
    }


def test_decisions(project):
    logic.decide_sentinel(project, "age", 999, True, now=NOW)
    logic.decide_sentinel(project, "sommeil_q3", 9, False, now=NOW)
    age = project.variable("age")
    assert [(d.code, d.missing, d.updated_at) for d in age.sentinel_decisions] == [(999, True, NOW)]
    assert w.confirmed_missing(age) == [999]
    assert w.confirmed_missing(project.variable("sommeil_q3")) == []
    logic.decide_sentinel(project, "age", 999, None, now=NOW)  # taken back
    assert project.variable("age").sentinel_decisions == []
    with pytest.raises(logic.ReviewError, match="pas un code suspect"):
        logic.decide_sentinel(project, "age", 42, True, now=NOW)


def test_what_a_decision_changes(project):
    age = project.variable("age")
    assert "Valeur 999 probablement manquante, non déclarée" in w.entry_anomalies(age)
    logic.decide_sentinel(project, "age", 999, True, now=NOW)
    assert not any("999" in text for text in w.entry_anomalies(age))
    assert "Code 999 confirmé comme manquant par l'équipe" in w.issues(age, "final")
    assert w.missing_codes(age) == "999 (confirmé par l'équipe)"
    logic.decide_sentinel(project, "sommeil_q3", 9, True, now=NOW)
    rows = {
        code: missing
        for code, _label, _count, missing in w.value_rows(project.variable("sommeil_q3"))
    }
    assert rows["9"] is True  # shown as a missing code in the value table
    # The profile describes the file as it is: not computed again.
    assert age.profile.max == 999


def test_the_model_is_told_and_never_asked(project):
    logic.decide_sentinel(project, "age", 999, True, now=NOW)
    sent = {
        v["name"]: v
        for b in build_blocks(project)
        for v in json.loads(b.messages[1]["content"])["variables"]
    }
    assert sent["age"]["declared_missing"] == ["999"]

    def server(request):
        content = json.loads(json.loads(request.content)["messages"][1]["content"])
        answer = {
            "block": content["block"],
            "variables": [
                {"name": v["name"], "label": "L", "description": "D"} for v in content["variables"]
            ],
        }
        return httpx.Response(200, json={"choices": [{"message": {"content": json.dumps(answer)}}]})

    before = [d.model_dump() for d in project.variable("age").sentinel_decisions]
    client = institution("m", server)
    list(enrich(project, client, now=lambda: NOW, block_size=40))
    assert [d.model_dump() for d in project.variable("age").sentinel_decisions] == before
    assert project.variable("nb_reveils").sentinel_decisions == []  # nothing proposed


def test_exports_and_dashboard(project, tmp_path):
    texts = [t.text.format(n=t.count) for t in tasks(project)]
    assert "Codes suspects à confirmer : 3" in texts
    logic.decide_sentinel(project, "age", 999, True, now=NOW)
    logic.decide_sentinel(project, "nb_reveils", -9, False, now=NOW)
    texts = [t.text.format(n=t.count) for t in tasks(project)]
    assert "Codes suspects à confirmer : 1" in texts

    path = tmp_path / "codebook.xlsx"
    write_xlsx(project, path)
    rows = list(load_workbook(path)["Provenance"].iter_rows(values_only=True))
    assert ("age", "Code suspect 999", "code manquant", "humain") in [r[:4] for r in rows]
    assert ("nb_reveils", "Code suspect -9", "vraie valeur", "humain") in [r[:4] for r in rows]
    report = render_markdown(project)
    assert "Code 999 confirmé comme manquant par l'équipe" in report
    assert "Valeur -9 confirmée comme vraie valeur par l'équipe" in report


def test_saved_and_read_back(project, tmp_path):
    logic.decide_sentinel(project, "age", 999, True, now=NOW)
    path = tmp_path / "p.dd.json"
    save_project(project, path)
    assert load_project(path).variable("age").sentinel_decisions[0].missing is True


@pytest.mark.allow_hosts(["127.0.0.1"])
def test_in_the_interface(project, tmp_path):
    path = tmp_path / "p.dd.json"
    save_project(project, path)
    app = AppTest.from_file(APP, default_timeout=30)
    app.session_state["project_path"] = str(path)
    app.session_state["mode"] = "variables"
    app.session_state["filter"] = "all"  # a settled age leaves "Avec signalements"
    app.run()
    app.sidebar.selectbox(key="variable").set_value("age").run()
    assert any("est-ce un code manquant" in x.value for x in app.warning)
    app.button(key="sentinel-age-999-missing").click().run()
    assert not app.exception, app.exception
    assert load_project(path).variable("age").sentinel_decisions[0].missing is True
    assert any("confirmé comme manquant" in x.value for x in app.info)
    app.button(key="sentinel-age-999-undo").click().run()
    assert load_project(path).variable("age").sentinel_decisions == []


def test_filter_of_suspect_codes(project):
    names = lambda: [e.name for e in logic.select(project, "suspects")]  # noqa: E731
    assert names() == ["age", "sommeil_q3", "nb_reveils"]
    logic.decide_sentinel(project, "age", 999, True, now=NOW)
    assert names() == ["sommeil_q3", "nb_reveils"]
    assert [t.filter for t in tasks(project) if t.text.startswith("Codes suspects")] == ["suspects"]
