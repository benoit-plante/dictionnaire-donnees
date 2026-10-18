"""v0.13, step 6: the page "Plan", `ddict plan` and `ddict article` (with `ddict resultats`
as its alias); the same files from the page and from the command line, on the fictional
database of the analyses."""

import json

import pytest
from docx import Document
from streamlit.testing.v1 import AppTest

from ddict.analysis import notebook as nb
from ddict.cli import main
from ddict.project import load_project
from test_notebook_page import APP, _folder

H1 = "Le bien-être au travail est plus élevé dans le groupe d'intervention."
H2 = "Plus les heures de travail sont nombreuses, plus le stress est élevé."
H3 = "Le soutien des collègues prédit la satisfaction au travail."
COMMANDS = [
    ["ajouter", "test_t", "--enonce", H1, "--dependante", "bienetre_total", "--groupe", "groupe",
     "--sens", "premier", "--famille", "F1"],
    ["ajouter", "correlation", "--enonce", H2, "--variables", "heures_travail", "stress",
     "--sens", "positif", "--famille", "F1"],
    ["ajouter", "regression", "--enonce", H3, "--dependante", "satisfaction", "--predicteurs",
     "soutien", "charge", "--predicteur", "soutien", "--sens", "positif"],
    ["regles", "--alpha", "0.05", "--origine", "primaires", "--etat", "avant-analyse"],
]  # fmt: skip
REASON = "Distribution des heures de travail très asymétrique."


def _plan(path, *arguments):
    return main(["plan", arguments[0], str(path), *arguments[1:]])


# --- The command line. ---


def test_the_commands(tmp_path, capsys):
    path = _folder(tmp_path)
    assert _plan(path, "lister") == 0
    out = capsys.readouterr().out
    assert "Seuil α : à taper ; correction : Holm." in out
    assert "Aucune hypothèse : ajoutez-en une avec « ddict plan ajouter »." in out
    for arguments in COMMANDS[:2]:
        assert _plan(path, *arguments) == 0
    assert "Hypothèse ajoutée : H2" in capsys.readouterr().out
    # A regression without its tested effect is refused, with its reason.
    assert _plan(path, *COMMANDS[2][:-4]) == 2
    assert "Choisissez l'effet testé : le modèle ou un prédicteur." in capsys.readouterr().err
    assert _plan(path, *COMMANDS[2]) == 0
    assert _plan(path, "figer") == 2
    err = capsys.readouterr().err
    assert "Tapez le seuil α du plan." in err
    assert "Dites l'origine des données et leur état quand le plan a été écrit." in err
    assert _plan(path, *COMMANDS[3]) == 0
    assert _plan(path, "deplacer", "H3", "1") == 0
    assert "H3 est maintenant au rang 1." in capsys.readouterr().out
    assert _plan(path, "deplacer", "H3", "3") == 0
    assert _plan(path, "figer") == 0
    assert "Plan arrêté le " in capsys.readouterr().out
    notebook = nb.load_notebook(nb.notebook_path(path))
    assert [e.hypothesis for e in notebook.entries] == ["H1", "H2", "H3"]
    # After the freezing, a change needs its reason; `amender` is `modifier`.
    assert _plan(path, "modifier", "H2", "--methode", "spearman") == 2
    assert "donnez la raison de cette modification" in capsys.readouterr().err
    assert _plan(path, "amender", "H2", "--methode", "spearman", "--raison", REASON) == 0
    project = load_project(path)
    assert project.plan.hypotheses[1].options == {"method": "spearman"}
    assert [a.target for a in project.plan.amendments] == ["H2"]
    notebook = nb.load_notebook(nb.notebook_path(path))
    assert [(e.id, e.hypothesis, e.retained) for e in notebook.entries] == [
        ("A1", "H1", True), ("A2", None, False), ("A3", "H3", True), ("A4", "H2", True),
    ]  # fmt: skip
    assert main(["carnet", "lister", str(path)]) == 0
    out = capsys.readouterr().out
    assert "   Face au plan : prévue (H1)" in out
    assert "   Face au plan : remplacée par l'amendement du " in out
    assert _plan(path, "lister") == 0
    out = capsys.readouterr().out
    assert f"H1 (confirmatoire, famille F1) — {H1}" in out
    assert "   Corrélation : relier des variables continues ; méthode : Spearman ;" in out
    assert "- le " in out and "pour H2 (méthode : Spearman au lieu de Pearson) : " in out
    assert _plan(path, "preenregistrement") == 0
    assert path.with_name("etude-analyses-preenregistrement.md").is_file()
    assert main(["article", str(path)]) == 0
    assert path.with_name("etude-analyses-article.docx").is_file()
    assert main(["resultats", str(path)]) == 0
    assert "« ddict resultats » s'appelle maintenant « ddict article »." in (
        capsys.readouterr().err
    )
    assert _plan(path, "effacer") == 2
    assert "ajoutez --confirmer" in capsys.readouterr().err
    assert _plan(path, "effacer", "--confirmer") == 0
    assert load_project(path).plan.hypotheses == []
    notebook = nb.load_notebook(nb.notebook_path(path))
    assert all(e.hypothesis is None and e.superseded is None for e in notebook.entries)


def test_an_unknown_variable_is_refused(tmp_path, capsys):
    path = _folder(tmp_path)
    arguments = ["ajouter", "test_t", "--enonce", H1, "--dependante", "inconnue", "--groupe",
                 "groupe"]  # fmt: skip
    assert _plan(path, *arguments) == 2
    assert "inconnue" in capsys.readouterr().err


# --- The page. ---


def _open(path):
    app = AppTest.from_file(APP, default_timeout=60)
    app.session_state["project_path"] = str(path)
    app.session_state["mode"] = "plan"
    app.run()
    assert not app.exception, app.exception
    return app


def _set(app, key, value):
    widgets = [w for kind in (app.selectbox, app.multiselect, app.text_input, app.text_area,
                              app.radio) for w in kind if w.key == key]  # fmt: skip
    assert widgets, key
    widgets[0].set_value(value).run()
    assert not app.exception, app.exception


def _hypothesis(app, statement, analysis, roles, **fields):
    _set(app, "plan-new-statement", statement)
    _set(app, "plan-new-analysis", analysis)
    for role, value in roles.items():
        _set(app, f"plan-new-role-{analysis}-{role}", value)
    for name, value in fields.items():
        _set(app, f"plan-new-{name}", value)
    app.button(key="plan-form-save").click().run()
    assert not app.exception, app.exception


def _plan_on_the_page(app):
    _hypothesis(app, H1, "t_test", {"outcome": "bienetre_total", "group": "groupe"},
                direction="first_higher", family="F1")  # fmt: skip
    _hypothesis(app, H2, "correlation", {"variables": ["heures_travail", "stress"]},
                direction="positive", family="F1")  # fmt: skip
    _hypothesis(app, H3, "regression", {"outcome": "satisfaction",
                                        "predictors": ["soutien", "charge"]},
                effect="predictor", predictor="soutien", direction="positive")  # fmt: skip
    _set(app, "plan-alpha", "0,05")
    _set(app, "plan-origin", "primary")
    _set(app, "plan-status", "before_analysis")


@pytest.mark.allow_hosts(["127.0.0.1"])
def test_the_plan_page(tmp_path):
    path = _folder(tmp_path)
    app = _open(path)
    assert app.header[0].value == "Plan"
    assert app.button(key="plan-freeze").disabled
    _plan_on_the_page(app)
    project = load_project(path)
    assert [h.id for h in project.plan.hypotheses] == ["H1", "H2", "H3"]
    assert app.text_area(key="plan-new-statement").value == ""  # the form emptied
    assert app.text_input(key="plan-new-family").value == ""
    assert project.plan.hypotheses[0].options == {"method": "welch", "alternative": "greater"}
    assert project.plan.hypotheses[2].effect.kind == "predictor"
    assert (project.plan.alpha, project.plan.data_status) == (0.05, "before_analysis")
    # A refusal says why, and nothing is saved.
    _set(app, "plan-alpha", "1,5")
    assert "strictement entre 0 et 1" in app.error[0].value
    assert load_project(path).plan.alpha == 0.05
    assert app.text_input(key="plan-alpha").value == "0,05"
    assert not app.button(key="plan-freeze").disabled
    app.button(key="plan-freeze").click().run()
    assert not app.exception, app.exception
    assert "Plan figé." in [s.value for s in app.success]
    assert load_project(path).plan.frozen is not None
    # After the freezing: a change without a reason is refused.
    app.button(key="plan-h-H2-edit").click().run()
    _set(app, "plan-new-option-correlation-method", "spearman")
    app.button(key="plan-form-save").click().run()
    assert "donnez la raison" in app.error[0].value
    _set(app, "plan-reason", REASON)
    app.button(key="plan-form-save").click().run()
    assert not app.exception, app.exception
    project = load_project(path)
    assert project.plan.hypotheses[1].options == {"method": "spearman"}
    assert "pour H2 (méthode : Spearman au lieu de Pearson)" in " ".join(
        m.value for m in app.markdown
    )
    app.button(key="plan-prereg").click().run()
    assert any("etude-analyses-preenregistrement.docx" in s.value for s in app.success)
    # Removing the whole plan asks a confirmation on the page.
    app.button(key="plan-clear").click().run()
    assert "Retirer le plan entier efface" in app.warning[-1].value
    app.button(key="plan-clear-cancel").click().run()
    assert load_project(path).plan.frozen is not None
    app.button(key="plan-clear").click().run()
    app.button(key="plan-clear-confirm").click().run()
    assert not app.exception, app.exception
    assert load_project(path).plan.hypotheses == []


@pytest.mark.allow_hosts(["127.0.0.1"])
def test_the_page_gives_the_command_files(tmp_path):
    page_path = _folder(tmp_path, "page")
    app = _open(page_path)
    _plan_on_the_page(app)
    app.button(key="plan-freeze").click().run()
    app.button(key="plan-h-H2-edit").click().run()
    _set(app, "plan-new-option-correlation-method", "spearman")
    _set(app, "plan-reason", REASON)
    app.button(key="plan-form-save").click().run()
    app.button(key="plan-prereg").click().run()
    app.session_state["mode"] = "notebook"
    app.run()
    app.button(key="notebook-write").click().run()
    assert not app.exception, app.exception

    command_path = _folder(tmp_path, "commande")
    for arguments in COMMANDS:
        assert _plan(command_path, *arguments) == 0
    assert _plan(command_path, "figer") == 0
    assert _plan(command_path, "amender", "H2", "--methode", "spearman", "--raison", REASON) == 0
    assert _plan(command_path, "preenregistrement") == 0
    assert main(["article", str(command_path)]) == 0

    def plan(path):
        return json.loads(load_project(path).plan.model_dump_json())

    assert plan(page_path) == plan(command_path)
    name = "etude-analyses-preenregistrement.md"
    assert page_path.with_name(name).read_bytes() == command_path.with_name(name).read_bytes()
    for name in ("etude-analyses-preenregistrement.docx", "etude-analyses-article.docx"):
        page, command = Document(page_path.with_name(name)), Document(command_path.with_name(name))
        assert [p.text for p in page.paragraphs] == [p.text for p in command.paragraphs]
        assert [[c.text for c in t._cells] for t in page.tables] == [
            [c.text for c in t._cells] for t in command.tables
        ]
    name = "etude-analyses-article-R.zip"
    assert page_path.with_name(name).read_bytes() == command_path.with_name(name).read_bytes()
