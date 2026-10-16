"""v0.10, step 5: the APA tables and figures on the Analyses page and with `ddict analyser
--format`: the same files from both, a person's number, title and note in the files, the
numbers never changed (D-086)."""

import io
import shutil
from pathlib import Path

import pytest
from docx import Document
from streamlit.testing.v1 import AppTest

from ddict.analysis.formats import FORMATS
from ddict.cli import main
from ddict.project import save_project
from test_analysis_data import EXAMPLES, KEY, analyses_project

APP = str(Path(__file__).resolve().parent.parent / "src" / "ddict" / "review" / "app.py")
T_TEST = ["--dependante", "bienetre_total", "--groupe", "groupe"]


def _folder(tmp_path, name="projet"):
    folder = tmp_path / name
    folder.mkdir()
    shutil.copy(EXAMPLES / KEY["sav"], folder / KEY["sav"])
    path = folder / "etude-analyses.dd.json"
    save_project(analyses_project(), path)
    return path


def _text(path):
    return FORMATS["docx"].read_text(path.read_bytes())


# --- The command. ---


def test_the_command_writes_the_formats(tmp_path, capsys):
    path = _folder(tmp_path)
    assert main(["analyser", str(path), "test_t", *T_TEST]) == 0
    assert path.with_name("etude-analyses-analyse-1.docx").is_file()
    assert path.with_name("etude-analyses-analyse-1-graphiques.xlsx").is_file()
    assert main(["analyser", str(path), "test_t", *T_TEST, "--format", "docx"]) == 0
    assert path.with_name("etude-analyses-analyse-2.docx").is_file()
    assert not path.with_name("etude-analyses-analyse-2-graphiques.xlsx").exists()
    assert main(["analyser", str(path), "test_t", *T_TEST, "--format", "aucun"]) == 0
    assert not path.with_name("etude-analyses-analyse-3.docx").exists()
    assert path.with_name("etude-analyses-analyse-3.md").is_file()
    capsys.readouterr()
    assert main(["analyser", str(path), "test_t", *T_TEST, "--format", "pdf"]) == 2
    assert "Format inconnu : pdf" in capsys.readouterr().err
    assert not path.with_name("etude-analyses-analyse-4.md").exists()  # refused before


def test_table1_has_no_chart_workbook(tmp_path):
    path = _folder(tmp_path)
    assert main(["analyser", str(path), "tableau1", "--variables", "age", "sexe"]) == 0
    assert path.with_name("etude-analyses-analyse-1.docx").is_file()
    assert not path.with_name("etude-analyses-analyse-1-graphiques.xlsx").exists()


# --- The page. ---


def _page(path):
    app = AppTest.from_file(APP, default_timeout=90)
    app.session_state["project_path"] = str(path)
    app.session_state["mode"] = "analyses"
    app.run()
    app.selectbox(key="analysis-key").set_value("t_test").run()
    app.selectbox(key="analysis-t_test-outcome").set_value("bienetre_total").run()
    app.selectbox(key="analysis-t_test-group").set_value("groupe").run()
    app.button(key="analysis-run").click().run()
    assert not app.exception, app.exception
    return app


@pytest.mark.allow_hosts(["127.0.0.1"])
def test_the_page_gives_the_command_files(tmp_path):
    path = _folder(tmp_path)
    app = _page(path)
    assert any(h.value == "Tableaux et figures APA" for h in app.subheader)
    assert any("Tableau 1" in m.value for m in app.markdown)
    assert any("<table" in m.value and "<i>M</i>" in m.value for m in app.markdown)
    app.button(key="analysis-write").click().run()
    assert not app.exception, app.exception
    other = _folder(tmp_path, "commande")
    main(["analyser", str(other), "test_t", *T_TEST])
    name = "etude-analyses-analyse-1.docx"
    assert _text(path.with_name(name)) == _text(other.with_name(name))
    graphs = "etude-analyses-analyse-1-graphiques.xlsx"
    assert FORMATS["xlsx"].read_text(path.with_name(graphs).read_bytes()) == FORMATS[
        "xlsx"
    ].read_text(other.with_name(graphs).read_bytes())


@pytest.mark.allow_hosts(["127.0.0.1"])
def test_a_persons_number_title_and_note(tmp_path):
    path = _folder(tmp_path)
    app = _page(path)
    numbers = [n for n in app.number_input if n.label == "Numéro"]
    titles = [t for t in app.text_input if t.label == "Titre"]
    notes = [n for n in app.text_area if n.label == "Note"]
    assert len(titles) == 2  # the table and the figure
    numbers[0].set_value(3).run()
    titles[0].input("Bien-être selon le groupe").run()
    [n for n in app.text_area if n.label == "Note"][1].input("Ma note de figure.").run()
    assert not app.exception, app.exception
    assert any("Tableau 3" in m.value and "Bien-être selon le groupe" in m.value
               for m in app.markdown)  # fmt: skip
    app.button(key="analysis-write").click().run()
    word = Document(io.BytesIO(path.with_name("etude-analyses-analyse-1.docx").read_bytes()))
    texts = [p.text for p in word.paragraphs]
    assert "Tableau 3" in texts and "Bien-être selon le groupe" in texts
    assert "Note. Ma note de figure." in texts
    assert word.tables[0].rows[2].cells[1].text == "17,01"  # the numbers never change
    note = next(p for p in word.paragraphs if p.text.startswith("Note. Intervention"))
    assert any(r.italic and r.text == "t" for r in note.runs)  # a kept note keeps its italics
    assert notes  # the fields exist for every item
