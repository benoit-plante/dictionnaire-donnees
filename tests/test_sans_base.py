"""v0.20, step B2: the format 0.9.0 — a project designed before its collection, with no
database yet (D-145), and the structured questionnaire (D-146). A project of v0.19 opens as
it is; what needs the database says so; the rest works without one."""

import io
import json
import shutil

import pytest
from openpyxl import load_workbook
from pydantic import ValidationError

from conftest import ROOT
from ddict import platforms, workflow
from ddict.analysis import data as analysis_data
from ddict.analysis import notebook as nb
from ddict.export import documented
from ddict.export import wording as w
from ddict.export.context import context_package
from ddict.export.markdown import render_markdown
from ddict.export.preparation import render_preparation
from ddict.export.syntax import render_syntax
from ddict.export.xlsx import write_xlsx
from ddict.glossary import propose_from_questionnaire
from ddict.llm.messages import variable_payload
from ddict.overview import row
from ddict.project import dumps, fingerprint, load_project, save_project
from ddict.project.completeness import parts
from ddict.project.model import (
    SCHEMA_VERSION,
    NoDatabase,
    Profile,
    Project,
    Provenance,
    Question,
    QuestionOption,
    SourceInfo,
    VariableEntry,
)
from ddict.questionnaire import link
from ddict.review.logic import codebook_progress
from ddict.workflow import add_questionnaire, profile_database
from projet_concu import project as designed

EXAMPLES = ROOT / "exemples" / "bases-fictives"
# What 0.9.0 adds, all with a default: what a v0.19 file does not have.
NEW_TOP = ("design",)
NEW_VARIABLE = ("planned",)
NEW_QUESTION = ("section", "instruction", "required", "answer", "layout", "provenance")
NEW_ITEM = ("code", "other", "provenance")


def test_the_format_is_0_9_0():
    assert SCHEMA_VERSION == "0.9.0"


# --- A designed project, with no database ----------------------------------------------------


def test_a_designed_project_round_trips_byte_for_byte(tmp_path):
    project = designed()
    assert project.source is None and all(e.profile is None for e in project.variables)
    first, second = tmp_path / "Étude conçue.dd.json", tmp_path / "copie.dd.json"
    save_project(project, first)
    reloaded = load_project(first)
    assert reloaded == project
    save_project(reloaded, second)
    assert first.read_bytes() == second.read_bytes()
    data = json.loads(first.read_text(encoding="utf-8"))
    assert data["source"] is None and data["design"]["created"] == "2026-10-19"
    assert data["questionnaire"]["format"] == "designed"
    assert [q["kind"] for q in data["questionnaire"]["questions"]] == [
        "single", "open", "display", "grid", "checkbox",
    ]  # fmt: skip
    assert data["variables"][0]["planned"] == {"type": "integer"}


def test_the_type_of_a_planned_variable_is_its_planned_type():
    project = designed()
    assert [e.data_type for e in project.variables][:3] == ["integer", "decimal", "integer"]
    assert project.variable("Q5_1").data_type == "text"
    assert project.variable("Q1").suspect_codes == []


def test_what_the_format_refuses():
    project = designed()
    entry = project.variables[0]
    with pytest.raises(ValidationError, match="profiled from the database, or planned"):
        VariableEntry.model_validate(entry.model_dump() | {"planned": None})
    profile = Profile(
        inferred_type="integer", n_rows=0, n_valid=0, n_system_missing=0, n_declared_missing=0
    )
    with pytest.raises(ValidationError, match="without a database, no variable has a profile"):
        Project.model_validate(
            project.model_dump()
            | {"variables": [entry.model_dump() | {"profile": profile.model_dump()}]}
        )
    with pytest.raises(ValidationError, match="names it"):
        Provenance(kind="source")
    with pytest.raises(ValidationError, match="only a text taken from a source"):
        Provenance(kind="team", page=4)
    with pytest.raises(ValidationError, match="only an open question"):
        Question(question_id="Q9", text="?", kind="single", answer="number")
    with pytest.raises(ValidationError, match="list or buttons"):
        Question(question_id="Q9", text="?", kind="open", layout="radio")
    with pytest.raises(ValidationError, match="no answer"):
        Question(
            question_id="Q9",
            text="Merci.",
            kind="display",
            options=[QuestionOption(option_id="Q9.o1", text="Oui")],
        )


def test_what_reads_the_database_works_without_one():
    project = designed()
    assert len(fingerprint.fingerprint(project)) == 64
    assert parts(project) and 0 <= codebook_progress(project) <= 1
    workflow.tasks(project)
    assert workflow.profile_summary(project) == "8 variables prévues ; pas encore de base."
    assert workflow.format_note(project) is None and workflow.platform_notes(project) == []
    assert platforms.of(project) is None
    propose_from_questionnaire(project)
    # A displayed text has no variable: it is never a target to link to.
    assert "Q3" not in {t.target_id for t in link.targets(project.questionnaire)}
    assert link.no_question_reason(project.variable("Q2")) is None
    payload = variable_payload(project.variable("Q1"), [], {}, statistics=True)
    assert payload["type"] == "integer" and "observed" not in payload
    assert "pct_missing" not in payload


def test_the_codebook_of_a_planned_dictionary(tmp_path):
    project = designed()
    text = render_markdown(project)
    assert "# Dictionnaire de données : Étude du sommeil (conception)" in text
    assert w.PLANNED_DICTIONARY in text
    assert "Valides" not in text and "SHA-256`" not in text
    out = io.BytesIO()
    write_xlsx(project, out)
    book = load_workbook(io.BytesIO(out.getvalue()))
    about = {r[0]: r[1] for r in book["Informations"].iter_rows(values_only=True)}
    assert about["Fichier de données"] == w.PLANNED_DICTIONARY
    assert "Lignes" not in about
    header = [c.value for c in next(book["Variables"].iter_rows(max_row=1))]
    first = dict(
        zip(header, next(book["Variables"].iter_rows(min_row=2, values_only=True)), strict=True)
    )
    assert first["Nom"] == "Q1" and first["Valides"] is None


def test_what_needs_the_database_says_so(tmp_path):
    project = designed()
    for render in (
        lambda: render_syntax(project, "spss"),
        lambda: render_syntax(project, "r"),
        lambda: render_preparation(project, "spss"),
        lambda: context_package(project),
    ):
        with pytest.raises(NoDatabase):
            render()
    path = tmp_path / "etude.dd.json"
    with pytest.raises(analysis_data.DataError) as found:
        analysis_data.find_database(project, path)
    assert found.value.reason == "no_base"
    with pytest.raises(documented.CopyError) as copy:
        documented.find_database(project, path)
    assert copy.value.reason == "no_base"
    assert "pas encore de base" in workflow.copy_error_text(copy.value, project)
    with pytest.raises(nb.NotebookError):
        nb.new_notebook(project, path)
    notebook = nb.Notebook(project=nb.ProjectLink(file=path.name, source_sha256="0" * 64))
    assert not nb.belongs(notebook, project)


def test_the_row_of_mes_projets_reads_no_database(tmp_path):
    path = tmp_path / "Études 2026" / "conçue.dd.json"
    path.parent.mkdir()
    save_project(designed(), path)
    found = row(path)
    assert found.source is None and found.base is None


# --- A project of v0.19 opens as it is (D-145) ------------------------------------------------


def _as_v019(project) -> str:
    """The file v0.19 would have written: format 0.8.0, none of the new fields."""
    data = json.loads(dumps(project))
    data["schema_version"] = "0.8.0"
    for key in NEW_TOP:
        del data[key]
    for entry in data["variables"]:
        for key in NEW_VARIABLE:
            del entry[key]
    questionnaire = data["questionnaire"]
    if questionnaire is not None:
        del questionnaire["sections"]
        for question in questionnaire["questions"]:
            for key in NEW_QUESTION:
                del question[key]
            for item in question["options"] + question["rows"]:
                for key in NEW_ITEM:
                    item.pop(key, None)
    return json.dumps(data, ensure_ascii=False, indent=2) + "\n"


@pytest.fixture(scope="module")
def sleep(fictional_files, tmp_path_factory):
    """The fictional sleep study and its questionnaire, as a team documents them."""
    folder = tmp_path_factory.mktemp("Étude du sommeil")
    database = folder / "etude-sommeil.sav"
    shutil.copy(fictional_files["etude-sommeil.sav"], database)
    project, _path = profile_database(database)
    add_questionnaire(project, EXAMPLES / "questionnaire-sommeil.pdf")
    return project


def test_a_project_of_v0_19_opens_as_it_is(sleep, tmp_path):
    assert isinstance(sleep.source, SourceInfo) and sleep.design is None
    path = tmp_path / "ancien.dd.json"
    path.write_text(_as_v019(sleep), encoding="utf-8", newline="\n")
    old = load_project(path)
    assert old.schema_version == "0.9.0"
    assert dumps(old) == dumps(sleep)  # nothing lost, nothing else changed
    assert fingerprint.fingerprint(old) == fingerprint.fingerprint(sleep)
    assert render_markdown(old) == render_markdown(sleep)
    assert all(e.profile is not None and e.planned is None for e in old.variables)
    assert all(q.provenance is None and q.section is None for q in old.questionnaire.questions)
