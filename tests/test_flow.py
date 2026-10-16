"""v0.12, step 4: the flow of participants (D-098), a figure drawn for Word and the table of
its counts in Excel, equal to the answer key of the fictional database of administration;
the common test of the formats (D-083) passed on it."""

import io

import pytest
from openpyxl import load_workbook

from ddict.analysis.apa import ApaDocument, describe_flow, plain
from ddict.analysis.data import apply_decisions
from ddict.analysis.formats import FORMATS, available
from ddict.analysis.formats.figures import _flow_lines, image
from ddict.project import exclusions as ex
from test_exclusions import KEY, _project

FLOW = KEY["exact"]["flow"]


@pytest.fixture(scope="module")
def figure():
    project, table = _project()
    return describe_flow(project, apply_decisions(project, table).flow, language="fr")


def test_the_counts_of_the_answer_key(figure):
    assert figure.kind == "flow" and figure.number == 1
    assert figure.flow_start == FLOW["start"] == 400
    assert [(s.rule, s.excluded) for s in figure.flow] == list(FLOW["excluded"].items())
    assert [s.remaining for s in figure.flow] == [388, 381, 367, 357, 352, 347]
    assert figure.flow[0].reason == "Refus de consentement"
    assert figure.title == "Flux des participants"
    assert plain(figure.note).startswith("Les exclusions sont appliquées dans l'ordre")


def test_without_rules_no_diagram():
    project, table = _project(rules=False)
    assert describe_flow(project, apply_decisions(project, table).flow) is None


def test_in_another_order():
    project, table = _project()
    ex.move(project, "E5", 4)
    found = describe_flow(project, apply_decisions(project, table).flow)
    assert [(s.rule, s.excluded) for s in found.flow] == list(
        KEY["exact"]["flow_e5_before_e4"]["excluded"].items()
    )
    assert found.flow[-1].remaining == 347


def test_the_lines_of_the_box(figure):
    lines = [text for text, _indented in _flow_lines(figure)]
    assert lines[0] == "Exclues ($\\mathit{n}$ = 53)"
    # A long reason is wrapped, its count never cut.
    assert "sur 6 ($\\mathit{n}$ = 14)" in lines[4]
    assert sum("$\\mathit{n}$" in line for line in lines) == 7


def test_the_image_is_the_same_every_time(figure):
    assert image(figure, "fr") == image(figure, "fr")
    odd = figure.__class__(**{**figure.__dict__, "flow": (
        figure.flow[0].__class__("E1", "Gain de 5 $ ou plus", 3, 397),)})  # fmt: skip
    assert image(odd, "fr")  # a "$" in a reason is text, not mathematics


def test_english():
    project, table = _project()
    found = describe_flow(project, apply_decisions(project, table).flow, language="en")
    assert found.title == "Flow of participants" and found.x_label == "Answers received"


@pytest.mark.parametrize("key", list(FORMATS))
def test_every_format_renders_the_flow(figure, key):
    document = ApaDocument(analysis="methode", language="fr", figures=(figure,))
    found = FORMATS[key]
    assert found in available(document)
    text = found.read_text(found.render(document))
    for expected in (figure.title, plain(figure.note), f"Figure {figure.number}"):
        assert expected in text


def test_the_counts_table_in_excel(figure):
    document = ApaDocument(analysis="methode", language="fr", figures=(figure,))
    book = load_workbook(io.BytesIO(FORMATS["xlsx"].render(document)))
    sheet = book["Figure 1"]
    rows = [list(r) for r in sheet.iter_rows(min_row=4, max_row=12, values_only=True)]
    assert rows[0] == ["Étape", "Raison", "Exclues", "Restantes"]
    assert rows[1] == [None, "Réponses reçues", None, 400]
    assert rows[2] == ["E1", "Refus de consentement", 12, 388]
    assert rows[7] == ["E6", "Courriel en double", 5, 347]
    assert rows[8] == [None, "Échantillon analysé", 53, 347]
    assert not sheet._charts  # no chart for a flow diagram (D-098)
