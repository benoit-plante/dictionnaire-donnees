"""v0.15, step 5: the checker of the figures of a manuscript (D-117, D-118;
docs/verification-depot.md, § 2.2 to 2.6): the tool's own draft agrees in full, in French
and in English; a draft altered by hand has each change found and classed; without the
database the notebook is the reference; the report, the commented copy, the command and
the page."""

import hashlib
import shutil

import pytest
from docx import Document

from ddict.analysis import notebook as nb
from ddict.analysis import notebook_export as export
from ddict.analysis import verify
from ddict.analysis.compute import run
from ddict.analysis.data import prepare
from ddict.io import read_database
from ddict.project import build_project, load_project, save_project
from test_recode import EXAMPLES, KEY

ANALYSES = [
    ("paired", {"first": ["anxiete_t1"], "second": ["anxiete_t2"]}),
    ("logistic", {"outcome": ["abandon"], "predictors": ["motivation", "stress"]}),
    ("correlation", {"variables": ["stress", "sommeil", "charge"]}),
    ("t_test", {"outcome": ["stress"], "group": ["abandon"]}),
    ("rm_anova", {"measures": ["anxiete_t1", "anxiete_t2", "anxiete_t3"]}),
]


@pytest.fixture(scope="module")
def study(tmp_path_factory):
    """The follow-up base, five analyses in the notebook, the drafts in French and English."""
    folder = tmp_path_factory.mktemp("verifier")
    shutil.copy(EXAMPLES / KEY["sav"], folder / KEY["sav"])
    path = folder / "etude-suivi.dd.json"
    project = build_project(read_database(folder / KEY["sav"]))
    save_project(project, path)
    prepared = prepare(project, path)
    notebook = nb.new_notebook(project, path)
    for key, chosen in ANALYSES:
        nb.add(notebook, project, run(project, prepared, key, chosen), today="2026-10-18")
    nb.save_notebook(notebook, nb.notebook_path(path))
    drafts = {}
    for language in ("fr", "en"):
        export.write_results(project, prepared, notebook, path, language=language,
                             today="2026-10-18")  # fmt: skip
        drafts[language] = folder / f"article-{language}.docx"
        shutil.copy(export.results_paths(path)["docx"], drafts[language])
    return project, prepared, notebook, path, drafts


def _alter(source, target, changes):
    """The draft with some figures changed by hand, as an author would."""
    document = Document(str(source))
    paragraphs = list(document.paragraphs)
    paragraphs += [p for t in document.tables for row in t.rows for cell in row.cells
                   for p in cell.paragraphs]  # fmt: skip
    for old, new, *context in changes:
        # A context: a text of the paragraph, or "=…" for a table cell's whole text.
        def chosen(p, context=context):
            if not context:
                return True
            if context[0].startswith("="):
                return p.text == context[0][1:]
            return context[0] in p.text

        found = [r for p in paragraphs if chosen(p) for r in p.runs if old in r.text]
        assert found, old
        found[0].text = found[0].text.replace(old, new, 1)
    document.save(str(target))


def _classes(result):
    return [(c.value.text if c.value else c.unread, c.kind, verify.shown(c.expected, c.value),
             c.reference) for c in result.problems()]  # fmt: skip


@pytest.mark.parametrize("language", ["fr", "en"])
def test_the_tool_s_own_draft_agrees_in_full(study, language):
    project, prepared, notebook, _path, drafts = study
    result = verify.verify(project, prepared, notebook, drafts[language])
    counts = result.counts()
    assert result.recomputed and counts["agrees"] > 100
    assert result.problems() == []


def test_each_change_found_and_classed(study, tmp_path):
    project, prepared, notebook, _path, drafts = study
    altered = tmp_path / "manuscrit.docx"
    _alter(drafts["fr"], altered, [
        ("= 8,18", "= 8,19"),            # A1: a t changed by a hundredth
        ("20,29", "20,92"),              # A4: a mean with two digits swapped (the sentence)
        ("6,45", "6,54", "=6,45"),       # A4: its SD, in the table
        ("d = −0,16", "d = 0,16"),       # A4: a d without its sign
        ("28,93", "31,02", "F(2, 404) = 28,93"),  # A5: the F of an older version
        ("p = ,250", "p = ,000"),        # A4: a p written ",000"
        ("t(140,08)", "t(140)"),         # A4: Welch's df rounded: still agrees
        ("Au total, ", "Un autre test donnait t(98) = 2,41. Au total, "),  # invented
    ])  # fmt: skip
    result = verify.verify(project, prepared, notebook, altered)
    found = _classes(result)
    assert ("t(227) = 8,19", "differs", "8,18", "A1") in found
    assert ("M = 20,92", "differs", "20,29", "A4") in found
    assert ("6,54", "differs", "6,45", "A4") in found
    assert ("d = 0,16", "sign", "−0,16", "A4") in found
    assert ("F(2, 404) = 31,02", "differs", "28,93", "A5") in found
    assert ("p = ,000", "differs", ",250", "A4") in found  # written as the author does
    assert ("t(98) = 2,41", "no_origin", "—", "") in found
    assert len(found) == 7  # and nothing else: the rounded df agrees


def test_against_the_notebook_without_the_database(study):
    project, _prepared, notebook, _path, drafts = study
    result = verify.verify(project, None, notebook, drafts["fr"])
    assert not result.recomputed
    first = [c for c in result.checked if c.block.section.startswith("A1")
             and c.block.kind == "sentence"]  # fmt: skip
    assert first and all(c.kind == "agrees" for c in first)
    assert "Base introuvable" in verify.summary_lines(result)[1]


def test_the_figures_of_another_study_have_no_origin(study, tmp_path):
    project, prepared, notebook, _path, _drafts = study
    document = Document()
    document.add_paragraph(
        "Les femmes dormaient autant que les hommes, t(172,26) = 0,18, p = ,861, d = 0,03. "
        "Le sommeil et les réveils n'étaient pas liés, r(198) = ,05, p = ,492."
    )
    other = tmp_path / "autre.docx"
    document.save(str(other))
    result = verify.verify(project, prepared, notebook, other)
    assert result.counts()["agrees"] == 0
    assert {c.kind for c in result.checked} == {"no_origin"}


def test_the_report_and_the_commented_copy(study, tmp_path):
    project, prepared, notebook, _path, drafts = study
    altered = tmp_path / "manuscrit.docx"
    _alter(drafts["fr"], altered, [("= 8,18", "= 8,19"), ("d = −0,16", "d = 0,16")])
    before = hashlib.sha256(altered.read_bytes()).hexdigest()
    result = verify.verify(project, prepared, notebook, altered)
    written = verify.write(result, altered)
    assert [p.name for p in written] == [
        "manuscrit-verification.md", "manuscrit-verification.xlsx", "manuscrit-verifie.docx",
    ]  # fmt: skip
    assert hashlib.sha256(altered.read_bytes()).hexdigest() == before  # never changed
    report = written[0].read_text(encoding="utf-8")
    assert "diffère : 1 ; signe inversé : 1 ; sans origine : 0" in report
    assert "| t(227) = 8,19 | 8,18 | A1 | diffère |" in report
    comments = [c.text for c in Document(str(written[2])).comments]
    assert comments == [
        "Diffère : l'outil écrit 8,18 (A1).",
        "Signe inversé : l'outil écrit −0,16 (A4) ; l'ordre des groupes est peut-être inversé.",
    ]


def test_a_pending_analysis_is_said(study, tmp_path):
    from ddict.project import exclusions as rules

    project, _prepared, notebook, path, drafts = study
    changed = project.model_copy(deep=True)
    rules.add(changed, {"kind": "condition", "variable": "age", "operator": "lt", "value": 25,
                        "reason": "Moins de 25 ans"})  # fmt: skip
    result = verify.verify(changed, prepare(changed, path), notebook, drafts["fr"])
    assert any(c.pending for c in result.checked)
    assert any(c.kind == "differs" for c in result.checked)  # the figures moved
    assert verify.summary_lines(result)[3].startswith("*")


def test_the_command(study, tmp_path, capsys):
    from ddict.cli import main

    _project, _prepared, _notebook, path, drafts = study
    altered = tmp_path / "manuscrit.docx"
    _alter(drafts["fr"], altered, [("= 8,18", "= 8,19")])
    assert main(["verifier", str(path), str(altered)]) == 0
    out = capsys.readouterr().out
    assert "diffère : 1 ;" in out and "manuscrit-verifie.docx" in out
    assert main(["verify", str(path), str(tmp_path / "absent.docx")]) == 2


@pytest.mark.allow_hosts(["127.0.0.1"])
def test_the_page(study, tmp_path):
    from streamlit.testing.v1 import AppTest

    from test_analysis_page import APP

    _project, _prepared, _notebook, path, drafts = study
    altered = tmp_path / "manuscrit.docx"
    _alter(drafts["fr"], altered, [("= 8,18", "= 8,19")])
    app = AppTest.from_file(APP, default_timeout=60)
    app.session_state["project_path"] = str(path)
    app.session_state["mode"] = "verify"
    app.session_state["verify-path"] = str(altered)
    app.run()
    assert not app.exception, app.exception
    app.button(key="verify-run").click().run()
    assert not app.exception, app.exception
    assert any("diffère : 1 ;" in m.value for m in app.markdown)
    assert len(app.dataframe) == 1
    app.button(key="verify-write").click().run()
    assert (tmp_path / "manuscrit-verifie.docx").is_file()
    assert load_project(path).recodes.rules == []  # nothing changed in the project
