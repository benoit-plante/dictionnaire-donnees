"""The English translation, as a spreadsheet for a reviewer (v0.5, D-034), and back.

uv run python scripts/tableur_relecture.py ecrire relecture.xlsx
    writes every string of the interface: where it appears, the French, the English
    proposed, and empty "Correction" and "Comment" columns for the reviewer;
uv run python scripts/tableur_relecture.py lire relecture.xlsx
    puts the reviewer's corrections into the English catalog (then run
    scripts/traductions.py), after checking that each keeps the {placeholders}.
"""

from __future__ import annotations

import argparse
import ast
import re
import sys
from collections import defaultdict
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

sys.path.insert(0, str(Path(__file__).resolve().parent))
import traductions

SOURCE = traductions.SOURCE
GLOSSARY = traductions.ROOT / "docs" / "glossaire-traduction.md"
PLACEHOLDER = re.compile(r"\{[^{}]*\}")

# Where a module's strings appear, in words a reviewer understands.
PLACES = {
    "review/ui/home_view.py": "Home page",
    "review/app.py": "Sidebar and to-do list",
    "review/ui/base_view.py": "Database page",
    "review/ui/questionnaire_view.py": "Questionnaire page",
    "review/ui/glossary_view.py": "Glossary page",
    "review/ui/ai_view.py": "Local AI page",
    "review/ui/bench_view.py": "Local AI page: model bench",
    "review/ui/variables_view.py": "Variables page",
    "review/ui/bulk_view.py": "Variables and Groups pages: grouped decisions",
    "review/ui/groups_view.py": "Groups page",
    "review/ui/export_view.py": "Export page",
    "review/ui/common.py": "Every page",
    "review/ui/dialog.py": "The system's Open window",
    "review/logic.py": "Messages after a decision",
    "export/wording.py": "Codebook (Excel and Markdown)",
    "export/xlsx.py": "Codebook (Excel)",
    "export/markdown.py": "Codebook (Markdown report)",
    "export/syntax.py": "SPSS and R command files (comments)",
    "export/documented.py": "Documented copy",
    "bench.py": "Model bench report",
    "cli.py": "Command line",
    "workflow.py": "Messages of several pages",
    "glossary.py": "Glossary page: messages",
    "questionnaire/labels.py": "Variables page: value labels",
    "questionnaire/link.py": "Variables page: question",
    "project/build.py": "Groups (computed description)",
}


def places() -> dict[str, list[str]]:
    """For each string, the places where it appears."""
    found: dict[str, set[str]] = defaultdict(set)
    for path in sorted(SOURCE.rglob("*.py")):
        module = path.relative_to(SOURCE).as_posix()
        place = PLACES.get(module, module)
        for node in ast.walk(ast.parse(path.read_text(encoding="utf-8"))):
            if (
                isinstance(node, ast.Call)
                and isinstance(node.func, ast.Name)
                and node.func.id in traductions.KEYWORDS
                and node.args
                and isinstance(node.args[0], ast.Constant)
                and isinstance(node.args[0].value, str)
            ):
                found[node.args[0].value].add(place)
    return {text: sorted(where) for text, where in found.items()}


INSTRUCTIONS = [
    "Review of the English translation of dictionnaire-donnees",
    "",
    "The tool helps research teams write the data dictionary (codebook) of a database. It was "
    "written in French; this sheet holds every sentence of its interface, with the English "
    "proposed.",
    "",
    'What to do, sheet "Strings":',
    "- Read the English column. When it is wrong, unclear or unnatural, write your version in "
    '"Correction". Leave "Correction" empty when the English is fine.',
    "- Keep every {placeholder} exactly as it is, with its braces: {n}, {path}, {name}… The "
    "tool replaces them with numbers or names.",
    "- Keep the names of commands and options as they are (--replace, ddict export…).",
    '- Keep the terms of the "Glossary" sheet, so that a notion is always said the same way; '
    'if a term itself should change, say so in "Comment".',
    '- "Where" says on which screen or in which document the sentence appears.',
    '- Use "Comment" for anything else: a question, a doubt, a suggestion.',
    "",
    "The two guides in English are in the project's docs/en folder (installation-windows.md "
    "and installation-macos.md); comments on them are welcome too, in a separate message.",
    "",
    "Thank you!",
]


def write(path: Path) -> None:
    catalog = traductions.read("en")
    where = places()
    book = Workbook()
    sheet = book.active
    sheet.title = "Instructions"
    for line in INSTRUCTIONS:
        sheet.append([line])
    sheet["A1"].font = Font(bold=True, size=14)
    sheet.column_dimensions["A"].width = 120
    for row in sheet.iter_rows():
        row[0].alignment = Alignment(wrap_text=True, vertical="top")

    strings = book.create_sheet("Strings")
    header = ["#", "Where", "French (source)", "English (proposed)", "Correction", "Comment"]
    strings.append(header)
    messages = [m for m in catalog if m.id]
    order = {place: rank for rank, place in enumerate(PLACES.values())}
    messages.sort(
        key=lambda m: (
            min((order.get(p, 99) for p in where.get(m.id, [])), default=99),
            m.id.casefold(),
        )
    )
    for number, message in enumerate(messages, start=1):
        strings.append(
            [number, "; ".join(where.get(message.id, [])), message.id, message.string, "", ""]
        )
    for cell in strings[1]:
        cell.font = Font(bold=True)
    for column, width in zip("ABCDEF", (6, 30, 60, 60, 60, 40), strict=True):
        strings.column_dimensions[column].width = width
    to_fill = PatternFill("solid", fgColor="FFF2CC")
    for row in strings.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        row[4].fill = to_fill
        row[5].fill = to_fill
    strings.freeze_panes = "C2"
    strings.auto_filter.ref = strings.dimensions

    terms = book.create_sheet("Glossary")
    terms.append(["French", "English", "Note"])
    for line in GLOSSARY.read_text(encoding="utf-8").splitlines():
        cells = [c.strip() for c in line.strip().strip("|").split("|")]
        if len(cells) == 3 and cells[0] not in ("Français", "---") and not cells[0].startswith("-"):
            terms.append(cells)
    for cell in terms[1]:
        cell.font = Font(bold=True)
    for column, width in zip("ABC", (40, 40, 50), strict=True):
        terms.column_dimensions[column].width = width
    book.save(path)
    print(f"{path} : {len(messages)} phrases")


def read(path: Path) -> None:
    sheet = load_workbook(path)["Strings"]
    header = [c.value for c in sheet[1]]
    french, correction = header.index("French (source)"), header.index("Correction")
    catalog = traductions.read("en")
    changed, refused = 0, []
    for row in sheet.iter_rows(min_row=2, values_only=True):
        source, new = row[french], row[correction]
        if not source or not new or not str(new).strip():
            continue
        new = str(new).strip()
        if sorted(PLACEHOLDER.findall(source)) != sorted(PLACEHOLDER.findall(new)):
            refused.append(source)
            continue
        if source in catalog and catalog[source].string != new:
            catalog[source].string = new
            changed += 1
    traductions.po_path("en").write_bytes(traductions.po_bytes(catalog))
    print(f"{changed} corrections intégrées.")
    for source in refused:
        print(f"Refusée (champs {{…}} différents de l'original) : {source}")
    print("Ensuite : uv run python scripts/traductions.py")


def main(argv: list[str] | None = None) -> None:
    parser = argparse.ArgumentParser(description="Tableur de relecture de la traduction anglaise.")
    parser.add_argument("action", choices=["ecrire", "lire"])
    parser.add_argument("tableur", type=Path)
    args = parser.parse_args(argv)
    (write if args.action == "ecrire" else read)(args.tableur)


if __name__ == "__main__":
    main()
